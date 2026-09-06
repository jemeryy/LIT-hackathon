"""Regression checks for irrelevant uploads on both sides of a claim."""
import copy
import datetime as dt
import io
import unittest
import zipfile
from unittest.mock import patch

import app
import exports
import rules


def exhibit(eid, facts=(), side="yours", **asset_fields):
    return {"id": eid, "title": eid + ".pdf", "kind": "pdf", "status": "ready",
            "side": side, "spot": "b2" if side == "theirs" else "", "gap": "agreement",
            "assets": [{"id": eid, "kind": "pdf", "filename": eid + ".pdf",
                        "meta": {"author": "third_party", "dated": True, "has_amount": True},
                        "viewer_url": f"/api/viewer?asset_id={eid}&page_index=0", **asset_fields}],
            "facts": [{"asset_id": eid, "category": "payment", "fact": "Deposit paid",
                       "evidence_key": "deposit_paid", "date": "2026-08-18", "fits": True, **f} for f in facts]}


class RelevanceTests(unittest.TestCase):
    def setUp(self):
        self.case = {"claim_type": "tenancy", "exhibits": [exhibit("E1", [{}])],
                     "intake": copy.deepcopy(app.EMPTY_INTAKE)}
        self.question = {"id": "b2", "they": "Repair invoice", "strength": "strong",
                         "when": "yes", "hint": "Check the invoice", "answer": "yes"}

    def test_irrelevant_file_does_not_contaminate_valid_rank_or_timeline(self):
        self.case["exhibits"] += [exhibit("E9", [{"fits": False, "date": "2030-01-01"}]), exhibit("E10", relevance="irrelevant")]
        rows = rules.evidence(self.case)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["strength"], "strong")
        self.assertEqual([s["exhibit_id"] for s in rows[0]["sources"]], ["E1"])
        excluded = rules.unranked_files(self.case)
        self.assertEqual([r["sources"][0]["exhibit_id"] for r in excluded], ["E9", "E10"])
        self.assertTrue(all(r["strength"] == "irrelevant" and r["rank"] is None for r in excluded))
        self.assertFalse(any(e.get("date") == "2030-01-01" for e in rules.timeline(self.case, dt.date(2026, 9, 6))))

    def test_irrelevant_upload_cannot_inherit_invoice_question_strength(self):
        self.case["exhibits"].append(exhibit("E8", side="theirs", relevance="irrelevant"))
        rows = rules.their_evidence([self.question], self.case["exhibits"])
        upload = next(r for r in rows if r["sources"])
        self.assertEqual(upload["strength"], "irrelevant")
        self.assertIsNone(upload["rank"])
        possible = next(r for r in rows if r.get("basis") == "answer")
        self.assertEqual(possible["sources"], [])
        self.assertIsNone(possible["rank"])

    def test_actual_other_side_file_uses_its_own_facts_and_strength(self):
        ex = exhibit("E7", [{"fact": "Repair of sofa cost $300", "evidence_key": "other_side_evidence"}], side="theirs")
        rows = rules.their_evidence([self.question], [ex])
        upload = next(r for r in rows if r.get("basis") == "file")
        self.assertEqual(upload["what"], "Repair of sofa cost $300")
        self.assertEqual(upload["rank"], 1)
        self.assertEqual(upload["strength"], "strong")
        ex["assets"][0]["meta"] = {"author": "claimant"}
        ex["facts"][0]["date"] = None
        rows = rules.their_evidence([self.question], [ex])
        self.assertEqual(next(r for r in rows if r.get("basis") == "file")["strength"], "weak")

    def test_irrelevant_file_stays_visible_when_answer_changes_to_no(self):
        self.question["answer"] = "no"
        rows = rules.their_evidence([self.question], [exhibit("E8", side="theirs", relevance="irrelevant")])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["strength"], "irrelevant")

    def test_null_key_does_not_fill_gather_card(self):
        self.case["exhibits"] = [exhibit("E9", [{"evidence_key": None, "fits": False, "fact": "Unrelated class list"}])]
        self.assertTrue(all(g["have"] == ["Nothing yet"] for g in rules.gaps(self.case)))

    def test_partial_exhibit_assessed_per_file(self):
        ex = self.case["exhibits"][0]
        ex["assets"].append({"id": "E1-2", "kind": "pdf", "filename": "unrelated.pdf"})
        rows = rules.unranked_files(self.case)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["what"], "unrelated.pdf")
        self.assertEqual(rules.evidence(self.case)[0]["strength"], "strong")

    def test_missing_quote_or_unreadable_pdf_needs_review(self):
        ex = exhibit("E9")
        result = {"meta": {}, "facts": [{"evidence_key": "deposit_paid", "fits": True, "quote": "deposit was paid"}]}
        with patch.object(app.extract, "parse", return_value={"text": "Some readable text"}), \
             patch.object(app.extract, "locate", return_value=None), \
             patch.object(app.llm, "extract_facts", return_value=result):
            app.process_asset(self.case, ex, ex["assets"][0])
        self.assertEqual(rules.asset_assessment(ex, ex["assets"][0])["status"], "needs_review")
        with patch.object(app.extract, "parse", return_value={"text": "x" * 13000}), \
             patch.object(app.llm, "extract_facts", return_value={"meta": {}, "facts": []}):
            app.process_asset(self.case, ex, ex["assets"][0])
        self.assertIn("Only part", rules.asset_assessment(ex, ex["assets"][0])["reason"])
        with patch.object(app.extract, "parse", return_value={"text": ""}), \
             patch.object(app.llm, "extract_facts", return_value={"meta": {}, "facts": []}):
            app.process_asset(self.case, ex, ex["assets"][0])
        self.assertEqual(rules.asset_assessment(ex, ex["assets"][0])["status"], "needs_review")

    def test_exports_keep_assessment_but_exclude_irrelevant_claim_files(self):
        self.case["exhibits"] += [exhibit("E9", relevance="irrelevant"), exhibit("E10", review_reason="Cannot read this file")]
        self.case["evidence"] = rules.evidence(self.case)
        with patch.object(exports, "_asset_pdf", return_value=b"%PDF-test"), \
             patch.object(exports, "claim_form_text", return_value="form"), \
             patch.object(exports, "events_text", return_value="events"), \
             patch.object(exports, "written_request", return_value="letter"):
            pack = zipfile.ZipFile(io.BytesIO(exports.claim_pack_zip(self.case)))
        self.assertEqual([n for n in pack.namelist() if n.startswith("exhibits/")], ["exhibits/E1_E1.pdf"])
        self.assertIn("E9 (E9.pdf): excluded, irrelevant", pack.read("manifest.txt").decode())
        self.assertIn("E10 (E10.pdf): excluded, needs review", pack.read("manifest.txt").decode())
        from openpyxl import load_workbook
        sheet = load_workbook(io.BytesIO(exports.evidence_xlsx(self.case))).active
        assessments = [row[5] for row in list(sheet.values)[1:]]
        self.assertIn("Irrelevant", assessments)
        self.assertIn("Needs review", assessments)

    def test_no_extracted_facts_does_not_mean_irrelevant(self):
        ex = exhibit("E9")
        self.assertEqual(rules.asset_assessment(ex, ex["assets"][0])["status"], "needs_review")

    def test_related_context_with_no_rank_key_is_retained(self):
        ex = exhibit("E9", [{"evidence_key": None, "context": True, "fact": "Tenant requested return of deposit"}], relevance="context")
        self.case["exhibits"].append(ex)
        self.assertEqual(rules.asset_assessment(ex, ex["assets"][0])["status"], "context")
        self.assertEqual(len(rules.evidence(self.case)), 1)
        self.assertEqual(rules.unranked_files(self.case)[0]["strength"], "context")
        self.assertTrue(any("Tenant requested" in e.get("detail", "") for e in rules.timeline(self.case, dt.date(2026, 9, 6))))
        with patch.object(exports, "_asset_pdf", return_value=b"%PDF-test"), \
             patch.object(exports, "claim_form_text", return_value="form"), \
             patch.object(exports, "events_text", return_value="events"), \
             patch.object(exports, "written_request", return_value="letter"):
            pack = zipfile.ZipFile(io.BytesIO(exports.claim_pack_zip(self.case)))
        self.assertIn("exhibits/E9_E9.pdf", pack.namelist())

    def test_placeholder_overrides_canned_condition_fact(self):
        ex = exhibit("E9", [{"evidence_key": "moveout_condition", "category": "condition"}], relevance="placeholder")
        self.case["exhibits"] = [ex]
        self.assertEqual(rules.asset_assessment(ex, ex["assets"][0])["status"], "placeholder")
        self.assertEqual(rules.evidence(self.case), [])
        self.assertEqual(rules.unranked_files(self.case)[0]["strength"], "placeholder")
        self.assertTrue(all(g["have"] == ["Nothing yet"] for g in rules.gaps(self.case)))

    def test_example_records_do_not_claim_placeholders_show_real_conditions(self):
        fixtures = rules.content("fixtures")["files"]
        for filename in ("WhatsApp_01.png", "WhatsApp_02.png", "WhatsApp_03.png", "WhatsApp_06.png"):
            self.assertEqual(fixtures[filename]["relevance"], "context")
            self.assertTrue(fixtures[filename]["facts"])
        for filename in ("moveout_walkthrough.mp4", "movein_01.jpg", "movein_02.jpg", "movein_03.jpg", "movein_04.jpg", "movein_05.jpg"):
            self.assertEqual(fixtures[filename]["relevance"], "placeholder")
            self.assertEqual(fixtures[filename]["facts"], [])


if __name__ == "__main__":
    unittest.main()
