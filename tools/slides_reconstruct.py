"""Rebuild a NotebookLM slide-deck PDF (baked raster pages) as an editable PPTX.

Per page: OCR text lines (Windows built-in OCR, upright plus both sideways
turns for rotated labels) -> native text boxes with colour and size sampled
from the pixels; artwork -> one movable picture element per object a vision
model names (see object_boxes; pixel rules alone as the offline fallback),
cropped to its own ink with a transparent surround; everything else -> one
cleaned background image with the text and elements patched out.

Text the illustrator drew AROUND (a callout whose leader line stops at the
words, a label in a gap left in an arrow) is left baked in that background
image, because erasing it leaves a hole nothing can fill correctly. A vision
model decides which lines those are; see diagram_lines().

Usage:
    python tools/slides_reconstruct.py in.pdf out.pptx [--font "Segoe UI"]
        [--font-map map.json]        # {"regex": "Font Name"} per line text
        [--classify-cache c.json]    # default <out>.labels.json

ponytail: aligned text grids stay as positioned text boxes, not native table
objects (identical look, simpler); upgrade to real tables if Jeremy needs to
add rows.
"""
import argparse
import asyncio
import base64
import io
import json
import os
import re
import sys

import fitz
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml

MIN_LINE_H = 8           # px; smaller OCR "lines" are icon noise
MIN_TYPE_PT = 8          # pt; smaller lines stay baked, see rebuild()
# generator watermark: erase, never re-type. Loose, because the OCR reads the
# small grey wordmark differently page to page ("NotebookLM", "Notebook1;M")
NOISE_RE = re.compile(r"^(?:notebook\W?[l1i]\W?m|gemini\s*notebook)$", re.I)
WATERMARK_LOGO = 2.4     # x the watermark's text height, to include its logo mark
PILL_GREY = 35           # |R-G|+|G-B| above which a pixel is the deck's own
                         # coloured ink, never the watermark's grey slab
PILL_PAD = 40            # px window around the watermark's OCR box the pill
                         # may extend into (its padding, corners and logo)
ALNUM_RE = re.compile(r"[0-9A-Za-z]")
# the generator's draft-stamp footers ("SLIDE 03 // PROJECT: ... // STATUS:
# DRAFT // SCALE: 1:1"). Round 9 stripped them as watermarks; round 11
# reversed that: they are part of the deck's bottom design and stay, re-typed
# as editable text with their panels shipped as elements. --strip-stamps
# restores the erase for a deck where they really are unwanted.
STAMP_RE = re.compile(r"SLIDE\s*\d+\s*//|//\s*(?:PROJECT|STATUS)\b", re.I)
STRIP_STAMPS = False
COLOR_DIST = 60          # pixel-to-bg distance that counts as "ink"
ELEM_MIN_AREA = 1200     # px^2; smaller components stay in the background
PAD = 4                  # px padding around boxes


OCR_SCALE = 2            # recognise at this multiple of the page raster
ROW_GAP = 1.2            # x line height; a wider gap on one row is a second column
RULE_RE = re.compile(r"^(?=.*[-‐-―|_~])[^0-9A-Za-z]*[0-9A-Za-z]?[^0-9A-Za-z]*$")
# a step number drawn in a circle, read as a word: bare "1" or "(3)". Only ever
# stripped off the FRONT of a line, because that is where a step number is
# drawn; a TRAILING digit is real text ("TIER 1", "SCALE: 1:1") and stripping
# those was how "TIER 1" ended up on the slide as "TIER". A real list marker
# keeps its full stop ("1. CPF Board"), so it is left alone either way.
CIRCLE_RE = re.compile(r"^\(?\d\)?$")
# ...unlike a real list marker, which is set the same way and sits a similar
# indent from its text, but keeps its full stop or bracket
LIST_RE = re.compile(r"^\d+[.)]$")


# the OCR hands back a letter-spaced "FARA:" as two words, and joining them
# with a space types "FARA :" onto the slide (and fits the line a size small)
PUNCT_RE = re.compile(r"^[:;,.!?)\]}%]+$")


def join_parts(parts):
    out = ""
    for p in parts:
        out = (p["text"] if not out else
               out + p["text"] if PUNCT_RE.match(p["text"]) else
               out + " " + p["text"])
    return out


def strip_rule_glyphs(row):
    """Drop a leader line, tick or numbered circle the OCR read as a word.

    The label "Strict Exclusion:" has a leader line running off it and the OCR
    reads the dash and its end dot as one more word, "-o"; the step number in
    its circle comes back as a leading "2". Left in, they are typed onto the
    slide as text AND they stretch the line's box, so the label is fitted a
    size too large and the erase misses its first letters, which then show
    through under the re-typed line.

    Only the SHAPE of the token decides, never its distance from the line. A
    gap test was tried and cannot work: across this deck a leader line's dot
    sits 0.62-0.71 line heights out and ordinary words sit 0.60-0.87, so the
    two ranges overlap. It was throwing away real text -- the "B" of "SRC B",
    the "//" of the footer -- to catch two dots.
    """
    parts = list(row["parts"])

    def strays(t, lead):
        return (not LIST_RE.match(t)
                and ((lead and CIRCLE_RE.match(t))
                     or (len(t) <= 3 and RULE_RE.match(t))))

    while len(parts) > 1 and strays(parts[-1]["text"], False):
        parts.pop()
    while len(parts) > 1 and strays(parts[0]["text"], True):
        parts.pop(0)
    return {"text": join_parts(parts),
            "words": parts,
            "box": (min(p["box"][0] for p in parts),
                    min(p["box"][1] for p in parts),
                    max(p["box"][2] for p in parts),
                    max(p["box"][3] for p in parts))}


MARKER_MAX = 2           # chars; a longer leading token is a word, not a marker
# ...and these short ones are words too: "An AI research assistant" starts a
# paragraph whose second word happens to sit at the list's indent
MARKER_STOP = {"a", "an", "at", "as", "by", "he", "if", "in", "is", "it", "my",
               "no", "of", "on", "or", "so", "to", "up", "we"}


def strip_markers(rows):
    """Drop a drawn list marker the OCR read as a letter.

    Slide 3 sets its list with a drawn checkbox per line; the OCR reads it as
    "ü" on one line and "Cl" on the next, and both get typed onto the slide in
    front of the text. One line cannot tell a marker from a word, but several
    can: the same short token leading lines whose text all starts at one indent
    is a column of markers. A real "1." list marker is left alone.
    """
    at_indent = {}
    for r in rows:
        p = r["parts"]
        if (len(p) > 1 and len(p[0]["text"]) <= MARKER_MAX
                and p[0]["text"].lower() not in MARKER_STOP
                and not LIST_RE.match(p[0]["text"])):
            at_indent.setdefault(p[1]["box"][0] // SAME_LEFT, []).append(r)
    for group in at_indent.values():
        if len(group) < 2:
            continue
        for r in group:
            r["parts"] = r["parts"][1:]        # the line is rebuilt from parts
    return rows


def merge_rows(frags):
    """Join fragments that sit on one baseline into a single line.

    Windows OCR breaks a line at any wide word gap, so a letter-spaced headline
    comes back as "TRACEABILITY" / "AS" / "THE" / "MOAT" and a paragraph as two
    side-by-side pieces. Rebuilt one text box per fragment, that is exactly the
    broken text Jeremy sees. Merging is by gap RELATIVE to the line height, so
    real columns (2-3 line heights apart in these decks) stay separate.
    """
    rows = []
    # left to right, so a row is always extended in reading order: sorting by
    # y first puts a fragment sitting 1px higher ahead of its own neighbour.
    for f in sorted(frags, key=lambda f: f["box"][0]):
        x0, y0, x1, y1 = f["box"]
        h = y1 - y0
        for parts in rows:
            px1, py0, py1 = (parts[-1]["box"][2],
                             min(p["box"][1] for p in parts),
                             max(p["box"][3] for p in parts))
            # the gap is judged against the TALLER of the two, i.e. the line's
            # own height. A ">" or a "//" has a box a third the height of the
            # words either side, and measuring its gap against its own little
            # box orphaned it: "Right > Fluent" came out as two boxes with the
            # symbol dropped, and the footer lost its "07" and "1:1".
            if (min(y1, py1) - max(y0, py0) > 0.5 * min(h, py1 - py0)
                    and 0 <= x0 - px1 <= ROW_GAP * max(h, py1 - py0)):
                parts.append(f)
                break
        else:
            rows.append([f])

    return [{"text": join_parts(parts),
             "parts": parts,
             "box": (min(p["box"][0] for p in parts),
                     min(p["box"][1] for p in parts),
                     max(p["box"][2] for p in parts),
                     max(p["box"][3] for p in parts))} for parts in rows]


def ocr_lines(img):
    import winocr

    # Recognise upscaled. At the deck's native 1376x768 the OCR silently misses
    # whole lines (page 4's "Official Sources Only:" never came back), which
    # leaves that label baked while its own body text is extracted.
    big = img.resize((img.width * OCR_SCALE, img.height * OCR_SCALE),
                     Image.LANCZOS)

    async def run():
        return await winocr.recognize_pil(big, "en-US")

    result = asyncio.run(run())
    # one fragment per WORD, not per OCR line. Only word boxes can say where a
    # stray glyph welded onto a line actually ends, and the erase has to know
    # that exactly or the line's first letters survive under the re-typed text.
    frags = []
    for l in result.lines:
        for w in l.words:
            r = w.bounding_rect
            frags.append({"text": w.text,
                          "box": (int(r.x) // OCR_SCALE, int(r.y) // OCR_SCALE,
                                  int(r.x + r.width) // OCR_SCALE,
                                  int(r.y + r.height) // OCR_SCALE)})

    lines = []
    for f in (strip_rule_glyphs(r) for r in strip_markers(merge_rows(frags))):
        if f["box"][3] - f["box"][1] < MIN_LINE_H:
            continue
        # a lone glyph is a drawn circle/tick the OCR read as O/o/0/1, not a
        # label. Judged AFTER merging: "[1]," is a lone glyph on its own and a
        # word in its line, and dropping it tore a paragraph in half.
        if len(ALNUM_RE.findall(f["text"])) < 2:
            continue
        box, drop = f["box"], bool(NOISE_RE.match(f["text"]))
        if drop:
            # the generator's watermark, erased and never re-typed. Its logo
            # mark sits just left of the words, so take them together.
            box = (box[0] - int(WATERMARK_LOGO * (box[3] - box[1])),
                   box[1], box[2], box[3])
        text = fix_caps(fix_roman(f["text"]))
        words = [dict(w, text=fix_caps(w["text"])) for w in f["words"]] \
            if text != f["text"] else f["words"]
        lines.append({"text": text, "box": box, "drop": drop, "words": words})
    return lines


GAP_WELD_MAX = 3.5       # x line height; a same-row gap wider than this is a
                         # column gutter, never a missing token (page 3's
                         # "[1]," gap is 2.95x; real gutters here run 8x+)
GAP_CHARSET = "[](){}0123456789.,:;%&+*=/#!?-"
GAP_GLYPH_ERR = 0.5      # blurred shift-searched shape mismatch above which a
                         # gap glyph is no character of the deck's fonts and
                         # the whole weld is refused. Measured on page 3's
                         # "[1],": true glyphs 0.11-0.42, wrong ones 0.46+
GAP_MARGIN = 0.08        # a runner-up DIFFERENT character within this of the
                         # best is an ambiguous read: refuse, never guess
GAP_INK_MIN = 15         # px of ink a gap must hold to be worth reading

_gap_cands = {}


def _gap_alphabet(path, em):
    """Tight anti-aliased glyph masks for the weld charset, rendered at the
    line's own pixel size: warping a 6px-wide bracket to a fixed square
    destroys exactly the hooks that tell ']' from '1'."""
    key = (path, em)
    if key not in _gap_cands:
        from PIL import ImageDraw, ImageFont
        f = ImageFont.truetype(path, em)
        out = []
        for ch in GAP_CHARSET:
            bb = f.getbbox(ch)
            if bb[2] <= bb[0] or bb[3] <= bb[1]:
                continue
            im = Image.new("L", (bb[2] - bb[0] + 4, bb[3] - bb[1] + 4), 255)
            ImageDraw.Draw(im).text((2 - bb[0], 2 - bb[1]), ch, font=f, fill=0)
            g = 1.0 - np.asarray(im, float) / 255
            ys, xs = np.nonzero(g > 0.25)
            if not len(ys):
                continue
            out.append((ch, g[ys.min():ys.max() + 1, xs.min():xs.max() + 1]))
        _gap_cands[key] = out
    return _gap_cands[key]


def _shape_err(a, b):
    """Mismatch of two tight float ink masks, 0 = identical.

    Blurred and shift-searched (+-2px): these are ~20px glyphs of 1-2px
    strokes, where a rigid centre-aligned diff charges a correct match double
    for every pixel of rasteriser disagreement."""
    from scipy import ndimage
    h, w = max(a.shape[0], b.shape[0]) + 4, max(a.shape[1], b.shape[1]) + 4
    pa, pb = np.zeros((h, w)), np.zeros((h, w))
    pa[(h - a.shape[0]) // 2:(h - a.shape[0]) // 2 + a.shape[0],
       (w - a.shape[1]) // 2:(w - a.shape[1]) // 2 + a.shape[1]] = a
    pb[(h - b.shape[0]) // 2:(h - b.shape[0]) // 2 + b.shape[0],
       (w - b.shape[1]) // 2:(w - b.shape[1]) // 2 + b.shape[1]] = b
    pa = ndimage.uniform_filter(pa, size=2)
    pb = ndimage.uniform_filter(pb, size=2)
    return min(
        float(np.abs(pa - np.roll(pb, (dy, dx), (0, 1))).sum()
              / max(pa.sum(), pb.sum(), 1e-6))
        for dy in (-2, -1, 0, 1, 2) for dx in (-2, -1, 0, 1, 2))


def read_gap(img, box, families):
    """Read a short token the OCR missed, from the deck's own fonts.

    winocr returns NOTHING for page 3's "[1]," at any zoom and tesseract
    misreads it, but the ink is crisp vector output: segment it into glyphs
    and match each against the same characters rendered in the deck's own
    faces. Any glyph unmatched refuses the WHOLE token - half a guess would
    type wrong text over right pixels. Returns the token string or None.
    """
    x0, y0, x1, y1 = box
    if x1 <= x0 or y1 <= y0:
        return None
    arr = np.asarray(img.convert("L"), float)[y0:y1, x0:x1]
    bg = float(np.median(arr))
    ink = arr < bg - 60
    if ink.sum() < GAP_INK_MIN:
        return None
    faces = [p for p in (font_file(f) for f in families) if p]
    if not faces:
        return None
    # candidates rendered near the line's own scale: the row box spans
    # roughly bracket-top to descender ("[(g" ink), but OCR row boxes are
    # generous, so a few sizes around the estimate compete and the whole-token
    # fit decides (measured: the true size wins by 2x over one step out)
    base = max((y1 - y0) / ink_ratio(faces[0], "[(g"), 8.0)
    ems = sorted({max(int(round(base * s)), 8)
                  for s in (0.7, 0.8, 0.9, 1.0, 1.1)})
    aa = 1.0 - np.clip(arr / max(bg, 1.0), 0, 1)   # anti-aliased ink weight
    idx = np.flatnonzero(ink.any(axis=0))
    runs = np.split(idx, np.flatnonzero(np.diff(idx) > 1) + 1)
    rows = np.flatnonzero(ink.any(axis=1))
    cell_top, cell = int(rows[0]), max(int(rows[-1]) + 1 - int(rows[0]), 1)
    if len(runs) < 2 or len(runs) > 6:
        return None    # one run is as likely a drawn tick; many is drawing
    out, prev_end = [], None
    for r in runs:
        gx0, gx1 = int(r[0]), int(r[-1]) + 1
        sub = ink[:, gx0:gx1]
        rr = np.flatnonzero(sub.any(axis=1))
        glyph = aa[rr[0]:rr[-1] + 1, gx0:gx1]
        if glyph.size < 4:
            return None
        low = rr[-1] + 1 - cell_top > 0.7 * cell
        scored = sorted(
            (_shape_err(glyph, cm)
             + 0.6 * abs(np.log(glyph.shape[0] / cm.shape[0])), cc)
            for p in faces for em in ems for cc, cm in _gap_alphabet(p, em)
            # a comma's shape is any low blob: it must actually sit low
            if not (cc in ".," and not low))
        if not scored:
            return None            # alphabet empty (degenerate face): refuse
        err, ch = scored[0]
        rival = next((e for e, c in scored if c != ch), None)
        if err > GAP_GLYPH_ERR or (rival is not None
                                   and rival - err < GAP_MARGIN):
            return None
        if prev_end is not None and gx0 - prev_end > 0.45 * (y1 - y0):
            out.append(" ")
        out.append(ch)
        prev_end = gx1
    tok = "".join(out)
    if set(tok) <= set("-=. "):
        return None                # a drawn leader dash, not text
    return tok


def weld_gap_tokens(img, lines, families):
    """Rejoin a row the OCR split around a token it could not read.

    Page 3's rule line comes back as "reference (e.g.," and "[2]). Zero
    uncited claims." with the "[1]," between them never returned at all: the
    row ships as two text boxes and the token stays baked in the panel image
    (Jeremy's round-10 flag). When two same-baseline lines sit a readable-ink
    gap apart, read_gap names the ink; on success the pieces weld into ONE
    line carrying the token, so it is re-typed and erased with its row. An
    unreadable gap keeps both pieces exactly as they are.
    """
    lines = list(lines)
    i = 0
    while i < len(lines):
        a, hit = lines[i], None
        if not a["drop"] and not STAMP_RE.search(a["text"]):
            for j, b in enumerate(lines):
                if (j == i or b["drop"] or STAMP_RE.search(b["text"])):
                    continue
                ax0, ay0, ax1, ay1 = a["box"]
                bx0, by0, bx1, by1 = b["box"]
                ha, hb = ay1 - ay0, by1 - by0
                h = max(ha, hb)
                if min(ay1, by1) - max(ay0, by0) < 0.6 * min(ha, hb):
                    continue                  # not one baseline
                gap = bx0 - ax1
                if not ROW_GAP * h < gap <= GAP_WELD_MAX * h:
                    continue                  # already merged, or a gutter
                gbox = (ax1 + 2, min(ay0, by0), bx0 - 2, max(ay1, by1))
                # the gap must hold no OTHER OCR line: reading a third
                # fragment's ink as the token would type it twice (once
                # welded, once as its own line)
                if any(c["box"][0] < gbox[2] and gbox[0] < c["box"][2]
                       and c["box"][1] < gbox[3] and gbox[1] < c["box"][3]
                       for n, c in enumerate(lines) if n != i and n != j):
                    continue
                tok = read_gap(img, gbox, families)
                if tok:
                    hit = (j, tok, gbox)
                    break
        if hit is None:
            i += 1
            continue
        j, tok, gbox = hit
        b = lines[j]
        words = a["words"] + [{"text": tok, "box": gbox}] + b["words"]
        a["words"], a["text"] = words, join_parts(words)
        if b.get("sure"):
            a["sure"] = True       # a vision-verified fragment keeps its
            # small-type exemption through the weld (same OR as merge at
            # the footer path)
        a["box"] = (min(a["box"][0], b["box"][0]),
                    min(a["box"][1], b["box"][1]),
                    max(a["box"][2], b["box"][2]),
                    max(a["box"][3], b["box"][3]))
        del lines[j]                   # re-scan i: the row may reach further
        if j < i:
            i -= 1
    return lines


# ponytail: OCR reads leading roman numerals I/II/III as 1/11/111 and IV/IX as "1 V"
# / "1 X". Only rewrite when the rest of the line is uppercase, so a real
# "11. Something" list item survives.
ROMAN_RE = re.compile(r"^(1{1,3})\.\s+(?=[^a-z]*$)")
ROMAN_IX_RE = re.compile(r"^1\s*([VX])\.\s+(?=[^a-z]*$)")


def fix_roman(text):
    text = ROMAN_IX_RE.sub(lambda m: "I" + m.group(1) + ". ", text)
    return ROMAN_RE.sub(lambda m: "I" * len(m.group(1)) + ". ", text)


# ...and it reads a capital I inside a caps word as a 1, which is how the deck's
# own subject came out as "AN A1 BUILT FOR ACCOUNTABILITY". Only a token that
# already holds a capital is rewritten, so "SLIDE 01" and "TIER 1" keep theirs.
CAPS_I_RE = re.compile(r"\b(?=[A-Z1]*[A-Z])[A-Z]*1[A-Z1]*\b")


def fix_caps(text):
    if re.search(r"[a-z]", text):
        return text
    return CAPS_I_RE.sub(lambda m: m.group(0).replace("1", "I"), text)


def em_factor(text):
    """Point size / OCR bbox height. The bbox is tight to the ink, so what the
    line contains decides how much of the em it covers."""
    if re.search(r"[a-z]", text):
        return 0.88          # mixed case: OCR box ≈ ascender..descender
    return 1.38              # ALL CAPS: cap height only ≈ 0.7em


_font_index = None


def font_file(family):
    """Path of an installed regular TTF/OTF for a font family name, or None."""
    global _font_index
    if _font_index is None:
        import glob
        from PIL import ImageFont
        _font_index = {}
        dirs = [r"C:\Windows\Fonts",
                os.path.expandvars(r"%LOCALAPPDATA%\Microsoft\Windows\Fonts")]
        for d in dirs:
            for p in sorted(glob.glob(os.path.join(d, "*.ttf"))
                            + glob.glob(os.path.join(d, "*.otf"))):
                try:
                    fam, style = ImageFont.truetype(p, 12).getname()
                except Exception:
                    continue
                # keep the plainest style seen for a family
                key = fam.lower()
                if key not in _font_index or "regular" in style.lower():
                    _font_index[key] = p
    return _font_index.get(family.lower())


_ratio_cache = {}


def ink_ratio(path, text):
    """Ink height of `text` in this font, as a fraction of the em size.

    em_factor's constants are a guess that only holds for one font: a Didone
    and a slab have different cap heights, so a fixed factor sizes one of them
    wrong. Measuring the actual glyphs removes the guess.
    """
    key = (path, re.sub(r"[a-z]", "a", re.sub(r"[A-Z]", "A", text))[:40])
    if key not in _ratio_cache:
        from PIL import ImageFont
        f = ImageFont.truetype(path, 100)
        top, bottom = f.getbbox(text)[1], f.getbbox(text)[3]
        _ratio_cache[key] = max((bottom - top) / 100.0, 0.1)
    return _ratio_cache[key]


# em; where PowerPoint puts the first line's font origin relative to the frame's
# top edge. Fitted, not derived: PowerPoint's own line box is not the ascender.
ORIGIN_ABOVE_TOP = 0.10
# Same question for a paragraph carrying an explicit line height, which line 0
# does whenever its block has a second line. Least squares over 36 rendered
# probes (3 sizes x 4 heights x 3 ink tops): first ink row sits
# GAP_TO_INK*height + (ink top - INK_BASE)*size below the frame's top edge,
# within 2.4px. Re-fit with scratch probe4.py if the renderer changes.
GAP_TO_INK, INK_BASE = 0.7055, 0.9824


def ink_top(path, text):
    """Distance from the font origin down to the first inked row, in ems.

    Text-dependent, and that is the point: ALL-CAPS ink starts ~0.04em lower
    than lowercase with ascenders. A single constant for both put every capped
    heading ~2px low against the original.
    """
    from PIL import ImageFont
    # keyed on the real text, NOT ink_ratio's case-normalised form: "n" and "h"
    # are one height class but their ink starts a fifth of an em apart
    key = ("top", path, text[:40])
    if key not in _ratio_cache:
        _ratio_cache[key] = ImageFont.truetype(path, 100).getbbox(text)[1] / 100.0
    return _ratio_cache[key]


def ink_width(path, text, pt):
    """Width of the drawn glyphs, not the advance.

    The OCR box is tight to the ink, so fitting against the advance width
    (which carries the trailing sidebearing) stops a size early: every line
    came out ~5% small.

    Measured at the REAL size, fractions and all. Pillow renders a fractional
    size and a:sz carries one, so rounding to a whole point here fitted against
    a width the slide never draws: 17.5pt is 3% wider than 17pt, enough to
    steer the bisect in fit_size a size off.
    """
    from PIL import ImageFont
    b = ImageFont.truetype(path, max(pt, 1.0)).getbbox(text)
    return b[2] - b[0]


def pick_font(words, families):
    """Which family a line was set in, judged by the WORD widths.

    A deck mixes a mono and a proportional face, and re-typing a proportional
    line in the mono one is the ragged, too-wide look. Total width cannot tell
    them apart (a tracked-out mono line is wide too), but the widths of the
    words RELATIVE to each other can: in a mono they are exactly proportional
    to character count, in a proportional face "Tax" and "Loan" are not.
    Least squares through the origin, so tracking and size fall out of it.
    """
    from PIL import ImageFont
    m = np.array([w["box"][2] - w["box"][0] for w in words], float)
    if len(m) < 3 or m.min() <= 0:
        return families[0]
    # a display line is tracked out, which adds width per LETTER GAP on top of
    # the glyphs. Left out of the model, that extra width looks like the widths
    # of a proportional face and flips a mono headline to the wrong family.
    gaps = np.array([max(len(w["text"]) - 1, 0) for w in words], float)
    best = None
    for fam in families:
        path = font_file(fam)
        if not path:
            continue
        f = ImageFont.truetype(path, 100)
        p = np.array([max(f.getbbox(w["text"])[2] - f.getbbox(w["text"])[0], 1)
                      for w in words], float)
        A = np.column_stack([p, gaps]) if len(m) > 3 else p[:, None]
        res = float(np.linalg.lstsq(A, m, rcond=None)[1][0] / np.sum(m ** 2))
        if best is None or res < best[0]:
            best = (res, fam)
    return best[1] if best else families[0]


TRACK_MIN = 0.02         # x size; smaller letter-spacing is measurement noise
TIGHT_MAX = 0.05         # x size; tightening further makes the glyphs collide


def fit_size(text, family, box_w, box_h):
    """(px em size, px letter-spacing) that reproduces the original line.

    The width is always reproduced; the size comes from the HEIGHT whenever the
    letter-spacing that takes is one the glyphs survive. Fitting on width alone
    reproduces the line's length but not its size unless the substitute font is
    the original's width: this deck's mono is ~25% narrower than any installed
    mono, so every body line used to come out a fifth short. A font that IS the
    right width lands at ~0 spacing, and the display titles, which really are
    tracked out, fall out of the same rule with a positive one.

    Known ceiling: tightening past TIGHT_MAX collides the glyphs ("complex"
    renders as "cormlex"), so a font this much too wide still gives up some
    height -- less than before, not none.
    """
    path = font_file(family)
    if not path or len(text) < 2:
        return box_h * em_factor(text), 0.0
    gaps = len(text) - 1
    h_size = box_h / ink_ratio(path, text)

    def drawn(size):     # width with the most tightening the glyphs survive
        return ink_width(path, text, size) - TIGHT_MAX * size * gaps

    lo, hi = 4.0, max(h_size, 8.0)
    if drawn(lo) > box_w:
        return lo, 0.0
    for _ in range(12):                      # bisect on rendered width
        mid = (lo + hi) / 2
        if drawn(mid) <= box_w:
            lo = mid
        else:
            hi = mid
    # lo also bounds an OCR box that swallowed a taller neighbour: the height
    # can never size the line past what its width could explain
    size = min(h_size, lo)
    spacing = (box_w - ink_width(path, text, size)) / gaps
    return size, (spacing if abs(spacing) > TRACK_MIN * size else 0.0)


def ring_median(arr, box, pad=PAD):
    """Median colour of the ring just outside box - the local background."""
    x0, y0, x1, y1 = box
    h, w = arr.shape[:2]
    X0, Y0 = max(x0 - 3 * pad, 0), max(y0 - 3 * pad, 0)
    X1, Y1 = min(x1 + 3 * pad, w), min(y1 + 3 * pad, h)
    outer = arr[Y0:Y1, X0:X1].reshape(-1, 3)
    inner = arr[max(y0 - pad, 0):y1 + pad, max(x0 - pad, 0):x1 + pad].reshape(-1, 3)
    # cheap ring: median of outer block is dominated by bg even including inner
    del inner
    return np.median(outer, axis=0)


def band(arr, y, x0, x1, k=3):
    """Median colour per column over k rows starting at y (clamped)."""
    h = arr.shape[0]
    y = min(max(y, 0), h - k)
    return np.median(arr[y:y + k, x0:x1].astype(float), axis=0)


def text_color(arr, box):
    x0, y0, x1, y1 = box
    px = arr[y0:y1, x0:x1].reshape(-1, 3).astype(int)
    bg = ring_median(arr, box)
    dist = np.abs(px - bg).sum(axis=1)
    ink = px[dist > COLOR_DIST]
    if len(ink) == 0:
        return (245, 245, 245)
    # antialiased edge pixels outnumber solid ones, so a plain median comes out
    # washed towards the background. Take the core: the quarter furthest from it.
    d = np.abs(ink - bg).sum(axis=1)
    core = ink[d >= np.quantile(d, 0.75)]
    return tuple(int(v) for v in np.median(core, axis=0))


def free_row(y, x0, x1, others, step, limit=240):
    """Walk from y in direction step until the row clears every box in others.

    Sampling the background right above a line lands inside the line above it,
    which smears that line's ink down the patch as vertical streaks.
    """
    for _ in range(limit):
        if not any(b[1] - PAD <= y < b[3] + PAD and b[0] - PAD < x1 and b[2] + PAD > x0
                   for b in others):
            break
        y += step
    return y


INVENT_MAX = 100         # strip disagreement above which a blend invents ink
TILE_SEARCH = 90         # px of page above a region to mine for its pattern
TILE_MIN = 24            # px; a narrower clean window cannot hold one period
MASK_PAD = 12            # px of true background around a masked patch that
                         # scores the fill candidates: at PAD=4 the ring holds
                         # no whole grid cell, so a flat fill scores level with
                         # the grid-continuing one and page 9's panels came
                         # back as white slabs
GRID_PRIOR = 0.9         # weight of |fill - synthesized grid| INSIDE a masked
                         # region; just under the ring's weight, so a
                         # candidate the surviving ring fully endorses beats
                         # the grid on what would otherwise be a dead tie
                         # (an element sitting on a locally tinted ground).
                         # Exists because the ring cannot see a phantom line
                         # a tile candidate copies into the interior (page
                         # 9's patch reprinted the panel above's border)


def _tile(arr, X0, X1, Y0, Y1, rest):
    """Continue the pattern sitting above the region downward: fill row i
    copies the row one whole period up.

    Blending cannot reproduce a drawn grid (it smears the crossing lines into
    gradients); repeating it can. The period is the lag at which the rows
    above best repeat. The window shrinks past any box sitting in the way (a
    neighbouring element, a text line), because rows under those may already
    be patched and must not be copied. Scored like any other candidate, so
    where the background does not actually repeat this fill simply loses.
    """
    # clamped to the page edge, not failed at it: page 10's disc reaches to
    # within 77px of the page bottom, and the from-below tile that continues
    # the mesh under it lives entirely in that margin
    lo = max(Y0 - TILE_SEARCH, 0)
    for b in rest:
        if b[0] - PAD < X1 and b[2] + PAD > X0 and b[1] - PAD < Y0 \
                and b[3] + PAD > lo:
            lo = max(lo, b[3] + PAD)
    if Y0 - lo < TILE_MIN:
        return None
    above = arr[lo:Y0, X0:X1].astype(float)
    ref = above[-1]
    lag = min((float(np.abs(above[-1 - k] - ref).mean()), k)
              for k in range(4, above.shape[0] - 1))[1]
    return above[[above.shape[0] - lag + (i % lag) for i in range(Y1 - Y0)]]


def _diffuse(actual, m):
    """The masked pixels relaxed to a smooth membrane anchored on the pixels
    around them (a Laplace fill, solved coarse then refined).

    The glow behind a dark slide's title is a smooth dome, and no strip
    blend or tile can rebuild a dome: interpolating straight across the
    erased words fills them with the darkness either side of the glow,
    which is the black band page 10's disc carried behind its re-typed
    title. Diffusion rebuilds exactly the smooth part, and on a flat
    ground it converges to that same flat and merely ties.
    """
    from scipy import ndimage
    if not m.any() or m.all():
        return None
    f = actual.astype(float).copy()
    f[m] = f[~m].mean(axis=0)
    h, w = m.shape
    s = max(1, min(h, w) // 40)
    fs, ms = f[::s, ::s].copy(), m[::s, ::s]
    for _ in range(200):
        fs[ms] = ndimage.uniform_filter(fs, size=(3, 3, 1))[ms]
    if s > 1:
        up = ndimage.zoom(fs, (h / fs.shape[0], w / fs.shape[1], 1), order=1)
        f[m] = up[:h, :w][m]
    for _ in range(30):
        f[m] = ndimage.uniform_filter(f, size=(3, 3, 1))[m]
    return f


def blend(arr, X0, X1, Y0, Y1, rest):
    """Fill a region by interpolating the strip above it into the strip below,
    per column.

    A flat fill leaves a visible rectangle wherever the background is not flat
    (the glow behind the titles on the dark slides), so interpolate instead:
    constant backgrounds come out constant, gradients stay smooth.
    """
    top = band(arr, free_row(Y0 - 3, X0, X1, rest, -1), X0, X1)
    bot = band(arr, free_row(Y1, X0, X1, rest, 1), X0, X1)
    t = np.linspace(0, 1, Y1 - Y0)[:, None, None]
    # how far the two source strips disagree, ignoring the few worst lines. A
    # rule that really crosses the region shows in BOTH strips; a leader line
    # that only touches one side does not, and interpolating it paints a stroke
    # straight through the erased words.
    return (top * (1 - t) + bot * t,
            float(np.quantile(np.abs(top - bot).sum(axis=1), 0.95)))


def fill_error(fill, actual, ink=None, drop=None):
    """How badly a fill reproduces what was really behind the erased glyphs.

    Scored against the region's OWN pixels minus the ink, so structure sitting
    inside the region (a rule crossing the middle of a line) counts. Comparing
    the strips above and below cannot see that: both come out clean and the
    fill silently wipes the rule.

    `drop` is a bool mask of pixels that are being erased (an element's own
    ink): those cannot judge the fill, only the background surviving around
    them can.
    """
    a = actual.astype(int)
    if drop is not None:
        keep = ~drop
    elif ink is None:
        keep = np.ones(a.shape[:2], bool)
    else:
        from scipy import ndimage
        # antialiased glyph edges sit between ink and background, so a colour
        # test alone leaves a halo that swamps the score. Grow the mask instead.
        near = np.abs(a - np.array(ink)).sum(axis=2) <= COLOR_DIST
        keep = ~ndimage.binary_dilation(near, iterations=2)
    if keep.sum() < 10:
        return 0.0
    return float(np.abs(fill - a).sum(axis=2)[keep].mean())


def _steer_to_grid(raw, gcut, actual, ink, drop):
    """Re-score fill candidates toward the synthesized empty page, where it
    models this ground.

    The plain score rewards smoothness: a pale ruled mesh occupies ~15% of
    the pixels at ~20 levels, so a membrane or blend that misses it entirely
    still beats the grid's phase noise, and every erased line left a flat
    slab in page 7's mesh. The prior charges a candidate for disagreeing
    with the grid exactly where content is being invented (under the erased
    ink), which no surviving pixel can judge. Gated by GRID_MISFIT on the
    pixels that CAN judge, so a synth that fails this ground (page 10's
    mesh) steers nothing.
    """
    gerr = fill_error(gcut, actual, ink, drop)
    if gerr > GRID_MISFIT * (min(e for e, _ in raw) + 1.0):
        return raw
    if drop is None:
        if ink is None:
            return raw
        from scipy import ndimage
        drop = ndimage.binary_dilation(
            np.abs(actual.astype(int) - np.array(ink)).sum(axis=2)
            <= COLOR_DIST, iterations=2)
    if not drop.any():
        return raw
    if int((~drop).sum()) < 10:
        # no surviving pixel can judge (fill_error's own floor): the gate
        # is meaningless and the prior would steer to the grid on no
        # evidence at all (page 10's comb on a fully masked crop). Counted
        # directly, because a legitimate exact-zero fit with many judges
        # must still steer.
        return raw
    return [(e + GRID_PRIOR * float(np.abs(np.asarray(f, float) - gcut)
                                    .sum(axis=2)[drop].mean()), f)
            for e, f in raw]


GRID_LINE_MIN = 6        # colour-sum deviation of a whole row/column median
                         # that reads as a ruled grid line
GRID_TINT = 6            # per-channel spread above which that deviation is a
                         # colour wash, not a grey rule
PALE_LINE = 60           # colour-sum amplitude up to which a line seen only
                         # by the FULL medians is trusted (page 7's regional
                         # grid); darker needs the cut medians' confirmation
                         # (page 8's axes and graph paper are 66+)
GRID_PITCH_MIN = 12      # px; full-median-only lines pitched tighter than
                         # this are an artwork texture, not page ruling
GRID_GAP = 6             # px outside an excluded box within which a line must
                         # show ink for the synth to continue it UNDER the box
GRID_EVID = 0.12         # x the covered span; less total visible line ink
                         # than this and the "line" is a stub touching the box
                         # (page 6's stub + register mark total 0.06x), not
                         # page ruling (page 8's margins 0.19x, page 9's 0.6x)
GRID_VIS = 0.5           # fraction of an excluded span whose pixels must sit
                         # within 2*TINT_MATCH of the line colour for the line
                         # to be VISIBLE through the box (a frame rule an
                         # oversized box merely overlaps: page 8's right frame
                         # reads 1.0, its phantom mid-line 0.11)
GRID_FAMILY = 4          # sibling lines (other flagged strokes whose eligible
                         # spans overlap this one) at which interrupted ruling
                         # is trusted to continue INVISIBLY under artwork:
                         # page 9's ruling has 12+ everywhere, page 8's
                         # annotation line and page 4's leader line have <= 3
GRID_MISFIT = 2.0        # x the best candidate's ring error at which the
                         # synthesized grid has clearly failed to model the
                         # ground around a masked patch, and its prior is
                         # dropped for that patch: page 10's mesh ring reads
                         # the comb-like synth at 3.5x the mesh tile's error,
                         # page 9's ruled panels keep theirs at ~1x
TEX_W = 1.0              # weight of the texture-energy prior on masked
                         # patches: how hard a candidate is charged, per
                         # unit of high-frequency amplitude, for filling a
                         # hole smoother than the ground around it really is
MESH_FAMILY = 8          # parallel lines at a steady pitch at which the
                         # page is RULED as a sheet (page 7's mesh, page 9's
                         # ruling) and every member line paints across the
                         # family's whole evidence hull: the per-line gates
                         # judge each interruption alone and leave holes in
                         # a mesh that plainly covers the page (page 7's
                         # right column lost its verticals, and every patch
                         # there painted a horizontal-only comb)
MESH_JITTER = 0.35       # x the median pitch; a family whose pitches vary
                         # more is decoration, not sheet ruling


def page_grid(arr, exclude=(), claimed=None):
    """A synthesized EMPTY page: background plus the page's own ruled grid.

    `claimed` (optional bool mask) marks pixels that leave the page when the
    deck is rebuilt: element ink a crop carries away, text that is re-typed.
    Those pixels cannot testify that a line lives on the EMPTY page - a leader
    line's own claimed ink was exactly the evidence that kept page 4's phantom
    row alive, and every patch under an element then repainted the stroke the
    element had carried away, which showed the moment the element moved.

    Column and row medians see through sparse content: a grid line that runs
    the page keeps its colour, text and drawings vanish. Positions are
    absolute, so a patch filled from this can never drift out of phase - the
    integer-lag tile does, and a big patched area came back flat where the
    page clearly ruled on (pages 8 and 9). None when the page is not ruled.

    `exclude` boxes (the vision model's artwork objects) are cut out of the
    medians: page 8's chart covers over half its rows and columns, so its
    own axes and graph paper passed the median test and the "empty" page
    carried a skeleton of the chart, which every patch then reprinted. But
    a page can ALSO rule only the region its artwork sits on (page 7's
    graph paper), and there cutting the box blinds the synth to real grid
    lines. Both medians together resolve it: a line the cut medians see is
    page pattern by construction; a line only the FULL medians see is kept
    when it is PALE (worst case a faint phantom under an element) and
    dropped when it is dark (the chart's own axes, a panel border).
    """
    a = arr.astype(float)
    bg = np.median(a.reshape(-1, 3), axis=0)
    cols = np.median(a, axis=0)
    rows = np.median(a, axis=1)

    def lines(med):
        # a ruled line shifts the background near-uniformly across channels;
        # a COLOURED band whose median clears GRID_LINE_MIN is content (the
        # tint wash down page 7's funnel), and baking it into the reference
        # would repaint it in every patch
        d = med - bg
        uniform = np.abs(d - d.mean(axis=1, keepdims=True)).max(axis=1) \
            <= GRID_TINT
        return (np.abs(d).sum(axis=1) > GRID_LINE_MIN) & uniform

    vc, hr = lines(cols), lines(rows)
    if exclude:
        import warnings
        cut = a.copy()
        for x0, y0, x1, y1 in exclude:
            # clamp stops to >=0: a box wholly above/left of the page gives a
            # negative x1/y1, and cut[..:-n] would NaN most of the page
            cut[max(y0, 0):max(y1, 0), max(x0, 0):max(x1, 0)] = np.nan
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")     # all-NaN slices: filled below
            cols_x = np.nanmedian(cut, axis=0)
            rows_x = np.nanmedian(cut, axis=1)
        cols_x = np.where(np.isnan(cols_x), bg, cols_x)
        rows_x = np.where(np.isnan(rows_x), bg, rows_x)
        vx, hx = lines(cols_x), lines(rows_x)
        pale_c = np.abs(cols - bg).sum(axis=1) <= PALE_LINE
        pale_r = np.abs(rows - bg).sum(axis=1) <= PALE_LINE
        # a full-median-only line too DARK for the pale gate can still be
        # the page's ruling when the px that make it are UNCLAIMED (round
        # 14: page 5's band grid lives only inside the artwork box, at
        # amp ~87). Claimed ink leaves the page with its element, so a
        # line that survives with the claimed px removed is page pattern;
        # an element's own axis is claimed and vanishes (page 8).
        if claimed is not None and claimed.any():
            cut2 = a.copy()
            cut2[claimed] = np.nan
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                cols_u = np.nanmedian(cut2, axis=0)
                rows_u = np.nanmedian(cut2, axis=1)
            cols_u = np.where(np.isnan(cols_u), bg, cols_u)
            rows_u = np.where(np.isnan(rows_u), bg, rows_u)
            pale_c |= lines(cols_u)
            pale_r |= lines(rows_u)

        def ruled(extra):
            # page ruling is sparse (~19px pitch on these decks); a denser
            # line set is an artwork's own texture (page 8's ~8px graph
            # paper survives the amplitude test because the median dilutes
            # it) and must stay out of the empty page
            prev = np.concatenate([[False], extra[:-1]])
            idx = np.flatnonzero(extra & ~prev)
            return len(idx) < 3 \
                or float(np.median(np.diff(idx))) >= GRID_PITCH_MIN
        ex_c, ex_r = vc & pale_c & ~vx, hr & pale_r & ~hx
        vc = vx | (ex_c if ruled(ex_c) else np.zeros_like(vc))
        hr = hx | (ex_r if ruled(ex_r) else np.zeros_like(hr))
        # the union must never sink a valid cut-median grid: if the extras
        # push an axis past the texture abort below, drop the extras
        if vc.mean() > 0.5:
            vc = vx
        if hr.mean() > 0.5:
            hr = hx
        # the cut medians' value is cleaner wherever they see the line
        cols = np.where(vx[:, None], cols_x, cols)
        rows = np.where(hx[:, None], rows_x, rows)
    if not (vc.any() and hr.any()) or vc.mean() > 0.5 or hr.mean() > 0.5:
        return None                    # no grid, or a texture/gradient

    def span(px, L, ivals, cl=None):
        """Where along a flagged line the ink is actually seen, plus the
        excluded (artwork) intervals that pass the evidence gates.

        A cut-median line can be REAL yet PARTIAL: page 1's half-width ruler
        and the mid-height register dashes median as full lines, and painting
        them page-wide printed a phantom stroke through every patch behind
        the artwork. Gaps under a pitch are ticks and dashes.
        """
        from scipy import ndimage
        obs = (np.abs(px - L).sum(axis=1) < np.abs(px - bg).sum(axis=1)) \
            & (np.abs(px - bg).sum(axis=1) > GRID_LINE_MIN)
        if cl is not None:
            obs &= ~cl                 # claimed ink is not the empty page's
            # ...but a claimed run is an interval like an artwork box:
            # ruling with edge evidence and a family continues under it
            # (page 9's divider lies ON a ruled row, and without this the
            # row went dark along the divider's whole span), while an
            # annotation that owes its existence to its own claimed ink
            # finds no family and dies (page 4's leader row)
            on = np.flatnonzero(cl)
            if len(on):
                runs = np.split(on, np.flatnonzero(np.diff(on) > 1) + 1)
                ivals = list(ivals) + [(int(r[0]), int(r[-1]) + 1)
                                       for r in runs]
        # inside a box the artwork's own ink can sit closer to the line
        # colour than to the background (page 1's mint cylinder vs its grey
        # ruler), so raw obs is meaningless there: only edge evidence may
        # continue a line under a box
        for lo, hi in ivals:
            obs[max(lo, 0):max(hi, 0)] = False   # clamp: a box wholly above
            # the page top gives hi<0, and obs[..:-n] would erase most of it
        # OR keeps the original: closing is only meant to bridge tick gaps,
        # but scipy's zero border erodes a run touching the array edge
        obs |= ndimage.binary_closing(obs, structure=np.ones(GRID_PITCH_MIN))
        # neighbouring grown boxes overlap or nearly touch (page 9's panel
        # row), leaving no room for evidence between them: judge the run of
        # boxes as one covered span
        merged = []
        for lo, hi in sorted((max(lo, 0), min(hi, len(obs)))
                             for lo, hi in ivals):
            if merged and lo - merged[-1][1] <= GRID_GAP:
                merged[-1][1] = max(merged[-1][1], hi)
            else:
                merged.append([lo, hi])
        # ONE edge suffices (page 9's panels run nearly to the frame, so the
        # far side shows no ink); the evidence fraction is what separates
        # real ruling from a stub at a box edge
        seen = int(obs.sum())
        elig, weak = [], []
        for lo, hi in merged:
            if hi <= lo:
                continue
            left = bool(obs[max(lo - GRID_GAP, 0):lo].any())
            right = bool(obs[hi:hi + GRID_GAP].any())
            if (left or right) and seen >= GRID_EVID * (hi - lo):
                elig.append((lo, hi))
            elif left and right and seen >= 2 * GRID_PITCH_MIN:
                # under-inked because ONE artwork box covers nearly the whole
                # span (page 5's grid sits behind a page-wide diagram, so
                # every horizontal fails the evidence fraction at once and no
                # line was left to vouch for its family). Both edges seen is
                # the stronger claim that stands in for the missing ink; the
                # family gate in paint() still decides whether it fills. The
                # absolute obs floor keeps a claimed TEXT row with a few
                # stray px at its box edges out: real ruling always shows a
                # run of line ink somewhere, stray specks never do.
                weak.append((lo, hi))
        return obs, elig, weak

    def family_span(flags):
        """(lo, hi) index range of a big steady-pitch line family, or None.
        Adjacent flagged indices are one physical stroke."""
        idx = np.flatnonzero(flags)
        pos = [int(i) for k, i in enumerate(idx)
               if k == 0 or i - idx[k - 1] > 2]
        if len(pos) < MESH_FAMILY:
            return None
        pitch = np.diff(pos)
        med = float(np.median(pitch))
        if med <= 0 or float(np.median(np.abs(pitch - med))) \
                > MESH_JITTER * med:
            return None
        return pos[0], pos[-1] + 1

    def paint(flag, pxs, Ls, iv_of, cl_of=lambda i: None, clamp=None):
        """obs per flagged line, eligible intervals filled only when earned.

        Edge evidence and ink fractions cannot separate an ANNOTATION line
        that stops at the artwork on purpose (page 8's chart-to-panel
        connector, page 4's leader: both end in a dot or arrowhead at the
        box edge) from RULING interrupted by artwork laid over it (page 9).
        What does separate them, measured across this deck: a line may
        continue under a box only if it is literally visible through it
        (GRID_VIS - a frame rule an oversized box merely overlaps), or it
        belongs to a pitched family of parallel lines interrupted the same
        way (GRID_FAMILY - ruling never comes alone, an annotation does).
        """
        info = [(i, *span(pxs(i), Ls(i), iv_of(i), cl_of(i)))
                for i in np.flatnonzero(flag)]
        raw_obs = [info[k][1].copy() for k in range(len(info))]
        # adjacent flagged indices are one physical stroke, not a family
        cl, spans_of = [], []
        for k, (i, _, elig, weak) in enumerate(info):
            if cl and i - info[cl[-1][-1]][0] <= 2:
                cl[-1].append(k)
                spans_of[-1] += elig + weak
            else:
                cl.append([k])
                spans_of.append(list(elig) + list(weak))
        for ci, c in enumerate(cl):
            for k in c:
                i, obs, elig, weak = info[k]
                for lo, hi in elig + weak:
                    close = np.abs(pxs(i)[lo:hi] - Ls(i)).sum(axis=1) \
                        <= 2 * TINT_MATCH
                    cm = cl_of(i)
                    if cm is not None:
                        close &= ~cm[lo:hi]    # claimed ink is not "visible
                        # through the box": the element carries it away
                    vis = float(close.mean())
                    fam = sum(1 for cj in range(len(cl)) if cj != ci
                              and any(l2 < hi and lo < h2
                                      for l2, h2 in spans_of[cj]))
                    if vis >= GRID_VIS or fam >= GRID_FAMILY:
                        obs[lo:hi] = True
        # a big steady-pitch family is sheet ruling: each member paints
        # across its OWN evidence hull (first to last RAW sighting, before
        # the interval fills above - a filled interval is inference, not a
        # sighting), clamped to the span the PERPENDICULAR family occupies:
        # a mesh is ruled both ways over one region, so page 5's verticals
        # stop where its horizontal ruling stops, and the stray sightings
        # that walked them below the band (a ruler tick, text fringe) fall
        # outside the clamp. The per-interval gates judge each interruption
        # alone, which leaves holes in a mesh that plainly rules the page
        # (page 7's right column lost its verticals and every patch there
        # painted a horizontal-only comb). Members with next to no evidence
        # of their own (a flagged noise line) stay as they are.
        pos = [info[c[0]][0] for c in cl]
        if len(pos) >= MESH_FAMILY and clamp is not None:
            pitch = np.diff(pos)
            med = float(np.median(pitch))
            if med > 0 and float(np.median(np.abs(pitch - med))) \
                    <= MESH_JITTER * med:
                for c in cl:
                    for k in c:
                        on = np.flatnonzero(raw_obs[k])
                        if len(on) < 2 * GRID_PITCH_MIN:
                            continue
                        lo = max(int(on[0]), clamp[0])
                        hi = min(int(on[-1]) + 1, clamp[1])
                        if hi <= lo:
                            continue
                        # only where a KNOWN occluder hides the line
                        # (an artwork box, claimed ink) or across small
                        # detection dropouts: a member's hull must never
                        # bridge an INTENTIONAL blank gap drawn in the
                        # open (a two-block ruler's designed break). An
                        # isolated 1-2px sighting is a PERPENDICULAR
                        # line's crossing, not this line's ink: counted,
                        # it split such a blank into sub-dropout runs
                        # that all bridged
                        i = info[k][0]
                        on2 = np.flatnonzero(raw_obs[k])
                        allow = np.zeros_like(raw_obs[k])
                        for r in np.split(on2, np.flatnonzero(
                                np.diff(on2) > 1) + 1):
                            if len(r) >= 3:
                                allow[r[0]:r[-1] + 1] = True
                        for l2, h2 in iv_of(i):
                            allow[max(l2, 0):max(h2, 0)] = True
                        cm = cl_of(i)
                        if cm is not None:
                            allow |= cm
                        off = np.flatnonzero(~allow[lo:hi])
                        for r in np.split(off, np.flatnonzero(
                                np.diff(off) > 1) + 1):
                            if len(r) and len(r) <= 2 * GRID_PITCH_MIN:
                                allow[lo + r[0]:lo + r[-1] + 1] = True
                        info[k][1][lo:hi] |= allow[lo:hi]
        return info

    vlay = np.broadcast_to(bg, arr.shape).copy()
    hlay = np.broadcast_to(bg, arr.shape).copy()
    if exclude or claimed is not None:
        for c, m, *_ in paint(vc, lambda c: a[:, c], lambda c: cols[c],
                             lambda c: [(b[1], b[3]) for b in exclude
                                        if b[0] <= c < b[2]],
                             (lambda c: claimed[:, c])
                             if claimed is not None else lambda c: None,
                             clamp=family_span(hr)):
            vlay[m, c] = cols[c]
        for r, m, *_ in paint(hr, lambda r: a[r], lambda r: rows[r],
                             lambda r: [(b[0], b[2]) for b in exclude
                                        if b[1] <= r < b[3]],
                             (lambda r: claimed[r])
                             if claimed is not None else lambda r: None,
                             clamp=family_span(vc)):
            hlay[r, m] = rows[r]
    else:
        vlay[:, vc] = cols[vc]
        hlay[hr] = rows[hr][:, None]
    dark = np.concatenate([cols[vc], rows[hr]]).mean() < bg.mean()
    return np.minimum(vlay, hlay) if dark else np.maximum(vlay, hlay)


def plan_patch(arr, box, pad, others, ink=None, mask=None, grid=None):
    """Best fill for a region, and how wrong it still is.

    Erasing a line means guessing what sits behind it. Blending down the page
    works on a flat panel, but where a horizontal rule or the edge of a filled
    box runs through the line, that blend smears the rule into a gradient. Such
    a rule is constant SIDEWAYS, so the same blend run left-to-right restores
    it exactly. Try both directions, keep whichever reproduces the real pixels.

    With a `mask` (erasing an element, not a line) two more candidates join:
    the strip of page directly above and directly below, copied in whole. On a
    patterned background (slide 9's grid) a copy that happens to align beats
    any blend, and the surviving background around the mask is what scores it,
    so a misaligned or artwork-carrying copy loses honestly.
    """
    x0, y0, x1, y1 = box
    h, w = arr.shape[:2]
    X0, X1 = max(x0 - pad, 0), min(x1 + pad, w)
    Y0, Y1 = max(y0 - pad, 0), min(y1 + pad, h)
    if X1 <= X0 or Y1 <= Y0:
        return None
    rest = [b for b in others if b != box]
    fv, dv = blend(arr, X0, X1, Y0, Y1, rest)
    fh, dh = blend(arr.transpose(1, 0, 2), Y0, Y1, X0, X1,
                   [(b[1], b[0], b[3], b[2]) for b in rest])
    fh = fh.transpose(1, 0, 2)
    actual = arr[Y0:Y1, X0:X1]
    drop = mask[Y0:Y1, X0:X1] if mask is not None else None
    gcut = grid[Y0:Y1, X0:X1] if grid is not None else None
    # diffusion serves the TEXT patches only (the smooth glow behind page
    # 10's re-typed title). On a MASKED patch it keeps every judge pixel
    # verbatim, scores a dishonest ~0, and wins while erasing whatever
    # crossed the hole (page 1's ruler vanished under the watermark pill).
    # The erased words must not anchor it: only the window's rim.
    dif = None
    if drop is None:
        dmask = np.ones(actual.shape[:2], bool)
        dmask[:2] = dmask[-2:] = False
        dmask[:, :2] = dmask[:, -2:] = False
        dif = _diffuse(actual, dmask)
    if drop is None and min(dv, dh) > INVENT_MAX:
        # both directions would carry structure in from one side only (the
        # leader line that stops AT a label). Flat local background instead:
        # it cannot reproduce a rule, but it cannot draw one through the words
        # (the page's own synthesized grid may, and beats flat where it fits).
        flat = np.broadcast_to(ring_median(arr, box), actual.shape).copy()
        cands = [flat] + ([gcut] if gcut is not None else []) \
            + ([dif] if dif is not None else [])
        scored = [(fill_error(f, actual, ink), f) for f in cands]
        if gcut is not None:
            scored = _steer_to_grid(scored, gcut, actual, ink, None)
        err, best = min(scored, key=lambda t: t[0])
        return (X0, X1, Y0, Y1), best, err
    tiles = []
    t = _tile(arr, X0, X1, Y0, Y1, rest)
    if t is not None:
        tiles.append(t)
    flipped = [(b[0], h - b[3], b[2], h - b[1]) for b in rest]
    t = _tile(arr[::-1], X0, X1, h - Y1, h - Y0, flipped)
    if t is not None:
        tiles.append(t[::-1])
    # texture from the tile, light from the blend: a plain tile cannot follow
    # a lit gradient (page 10's mesh darkens down the page and a copied strip
    # leaves a bright rectangle), but its detail rides on the blend's
    # luminance fine. Hybrids lead the masked list, so where too little true
    # background survives to score the fills, texture wins over smoothness.
    from scipy import ndimage as ndi
    hybrids = [fv + t - ndi.uniform_filter(t, size=(31, 31, 1)) for t in tiles]
    # grid + texture: the synth carries the coarse ruling at absolute
    # positions but is blind to a fine jittered sub-grid (page 5's paper
    # never survives the medians); the tile carries the fine texture but
    # its own coarse lines land off-phase. Ride the tile's high-frequency
    # detail on the synth (round 14).
    if gcut is not None:
        hybrids += [gcut + t - ndi.uniform_filter(t, size=(9, 9, 1))
                    for t in tiles]
    # a GRID is both blends at once: the vertical blend continues the vertical
    # lines, the horizontal one the horizontal lines, and per-pixel min (dark
    # lines on light ground) or max (light on dark) keeps both sets where a
    # single direction smears one of them away.
    crossed = [np.minimum(fv, fh), np.maximum(fv, fh)]
    if gcut is not None:
        crossed.append(gcut)
    if drop is None:
        cands = [fv, fh] + tiles + hybrids + crossed
    else:
        cands = hybrids + tiles + crossed + \
            [fv, fh, np.broadcast_to(ring_median(arr, box),
                                     actual.shape).astype(float)]
    if dif is not None:
        cands.append(dif)

    raw = [(fill_error(f, actual, ink, drop), f) for f in cands]
    if gcut is not None:
        # the prior guards against a tile smuggling a phantom line into the
        # interior, but it may only steer where the synth actually models
        # this ground: on page 10's mesh the synth is a comb of verticals,
        # its ring error dwarfs the mesh tile's, and steering toward it
        # painted that comb behind the ring
        raw = _steer_to_grid(raw, gcut, actual, ink, drop)
    if drop is not None and drop.any() and int((~drop).sum()) >= 10:
        # the numeric score rewards smoothness, the eye notices texture
        # vanishing: a fine jittered grid is invisible to every median
        # yet its absence reads as a plateau (page 5, round 14). Charge a
        # candidate for carrying less high-frequency energy into the hole
        # than the judges' own ground shows.
        def hp(x):
            x = np.asarray(x, float)
            return np.abs(x - ndi.uniform_filter(x, size=(7, 7, 1))) \
                .sum(axis=2)
        eref = float(hp(actual)[~drop].mean())
        raw = [(e + TEX_W * max(0.0, eref - float(hp(f)[drop].mean())), f)
               for e, f in raw]
    err, best = min(raw, key=lambda t: t[0])
    return (X0, X1, Y0, Y1), best, err


def patch(arr, box, pad=PAD, others=(), ink=None, mask=None, grid=None):
    """Erase a region with whichever blend direction reproduces it better.

    With a `mask`, only the masked pixels are replaced: the background around
    an element (the page's own grid, a neighbouring drawing) stays untouched.
    Each connected piece of the mask is filled from ITS own surroundings: one
    blend spanning a page-sized element box drifts off the page's real
    gradient by a few levels, which reprints a dark page's drawing as a
    faint ghost in exactly the erased shape.
    """
    if mask is not None:
        from scipy import ndimage
        lab, n = ndimage.label(mask)
        if n > 1:
            for i, sl in enumerate(ndimage.find_objects(lab), 1):
                sub = np.zeros_like(mask)
                sub[sl] = lab[sl] == i
                patch(arr, (sl[1].start, sl[0].start, sl[1].stop, sl[0].stop),
                      pad, others, ink, sub, grid)
            return
    p = plan_patch(arr, box, pad, others, ink, mask, grid)
    if p is None:
        return
    (X0, X1, Y0, Y1), fill, _ = p
    # hybrids are unbounded floats: without the clip, 260 wraps to 4 and a
    # black speck lands in the patch
    fill = np.clip(np.round(fill), 0, 255).astype(arr.dtype)
    if mask is None:
        arr[Y0:Y1, X0:X1] = fill
    else:
        m = mask[Y0:Y1, X0:X1]
        arr[Y0:Y1, X0:X1][m] = fill[m]


WASH_MIN = 300           # px; a smaller near-background run in a crop keeps
                         # its original pixels (antialiasing, small marks)
WASH_FIT = 24            # summed-channel residual (95th pct) a translucency
                         # fit must reach; worse means the region carries its
                         # OWN structure (page 8's graph paper) and flattening
                         # it to one colour would destroy that, so it stays
                         # opaque
RES_DONATE = 800         # px of a residue component that may fall outside its
                         # owning element and still be donated to it (page 8's
                         # crosshair tails run 43% past the chart's extent)
TEX_MIN = 5000           # px; an unclaimed component this big is the page's
                         # own repeating pattern (page 10's mesh is 55k), and
                         # its colour vetoes residue donation. Real dark marks
                         # outside the boxes (page 5's ticks, frame furniture)
                         # come in hundreds of pixels and stay out of the
                         # palette, so their donations survive
CONTAIN = 0.9            # fraction of an element's box inside another element's
                         # box at which the two ship as ONE picture (page 4's
                         # REJECT scribble over its funnel)
RULE_THIN = 14           # px; an element this thin on either axis is a drawn
                         # rule (page 10's title underline), never merged into
                         # the artwork whose box it happens to sit in
STAMP_PANEL_MAX = 8      # x the stamp text box's area; a bigger element that
                         # overlaps the stamp is page artwork, not the stamp's
                         # own panel, and must not be dropped with it
PANEL_SWEEP = 24         # px around a dropped stamp panel swept for the frame
                         # specks its mask ran short of
                         # ponytail: proximity stands in for ownership here; a
                         # deck that draws a legit sub-60px icon inside the wm
                         # pill rect within 24px of a draft panel would lose
                         # it - add a colour-match to the panel frame if that
                         # deck ever shows up
PANEL_SPECK = 60         # px; bigger leftovers near a dropped panel are page
                         # furniture (a crosshair), not the panel's debris


def _wash_solve(P, B):
    """(alpha, wash colour) explaining pixels P as a translucent wash over
    ground B, or None. Two honest failure modes return None: a fit that
    cannot reproduce the pixels (the region carries its own structure, like
    page 8's graph paper), and a flat background (any (a, W) pair fits, so
    the alpha is unidentifiable). A fit pinned at the alpha floor is a
    region that "explains itself away" against a background resembling it,
    not a real wash: the deck's true washes sit at 0.3-0.7. A best fit at
    ~1.0 alpha means the region is opaque; the 0.91 ceiling absorbs float
    noise. The most opaque reading that fits wins: over-transparent washes
    shift colour visibly the moment the element moves to different ground."""
    if B.std(axis=0).max() < 3:
        return None
    fits = []
    for a in np.arange(0.15, 1.0001, 0.05):
        W = np.clip((P - (1 - a) * B).mean(axis=0) / a, 0, 255)
        r = np.abs(a * W + (1 - a) * B - P).sum(axis=1)
        fits.append((float(np.quantile(r, 0.95)), a, W))
    good = [f for f in fits if f[0] <= WASH_FIT]
    if not good:
        return None
    qmin = min(f[0] for f in good)
    q, a, W = max((f for f in good if f[0] <= qmin + 0.5),
                  key=lambda f: f[1])
    if a > 0.91 or a < 0.2:
        return None
    return a, W


def translucent(tile, a8, mm, bg):
    """Give a crop's wash pixels their real colour and alpha, in place.

    A tinted wash (page 7's funnel and bucket fills) is drawn TRANSLUCENT
    over the page, so its pixels carry the page's own grid: cropped opaque
    they bake that grid into the element, and skipped they leave the fill
    behind. Solve the compositing instead: over the PATCHED background `bg`,
    a wash component is P = a*W + (1-a)*bg for one (a, W), so the crop can
    carry W at alpha a and the page's grid shows through wherever the
    element sits. The honest failure modes stay opaque (see _wash_solve).
    """
    from scipy import ndimage
    d = np.abs(tile.astype(int) - bg.astype(int)).sum(axis=2)
    solid = mm & (d > ELEM_DIST)
    # solid strokes and the 2px antialiased skirt around them stay opaque
    wash = mm & ~ndimage.binary_dilation(solid, iterations=2)
    lab, _ = ndimage.label(wash)
    for i, sl in enumerate(ndimage.find_objects(lab), 1):
        comp = wash[sl] & (lab[sl] == i)
        if comp.sum() < WASH_MIN:
            continue
        got = _wash_solve(tile[sl][comp].astype(float),
                          bg[sl][comp].astype(float))
        if got is None:
            continue
        a, W = got
        tile[sl][comp] = W.round().astype(np.uint8)
        a8[sl][comp] = int(round(a * 255))


CLOSE_R = 10             # px iterations; morphological closing that bridges
                         # the design gaps in an element's outline before its
                         # interior is solidified (page 10's ring breaks at
                         # the tick crossings)
KEY_LO, KEY_HI = 10, 40  # summed-channel distance to the patched ground: a
                         # mask-rim pixel below KEY_LO is copied background
                         # (alpha 0), above KEY_HI the element's own ink
                         # (opaque); between, alpha ramps
RIM = 3                  # px; only this outer band of the ink hull is
                         # edge-keyed, interior fill is never keyed away
MARGIN_INK = 100         # summed-channel distance below which a pixel outside
                         # the ink hull is copied page (a pale grid stub) and
                         # is dropped; a real pale stroke sits well above it
MARGIN_BG = 30           # component-median distance to the patched ground
                         # below which a margin component is the page's own
                         # content the mask swept in, not the element's paint
CORE_INK = 150           # summed-channel distance that counts as a STRONG
                         # stroke when building the hull: a pale texture
                         # mispredict (~60) must never seed the hull, or the
                         # hull grows diamonds around grid crossings and the
                         # true fill between them is keyed away (page 9's
                         # panel interior)
HOLE_WASH_MIN = 50       # px; smallest hole component the wash fit judges
                         # (translucent()'s WASH_MIN guards open regions,
                         # where a small fit is noise; an ENCLOSED hole is
                         # pre-qualified by its own boundary)
CELL_SHELL = 2           # px; the ring of mask ink around a hole that
                         # testifies to what encloses it: pale furniture (a
                         # ruler's teeth) or a label's dark glyphs
CELL_BOUND = 0.65        # fraction of a hole's boundary ring the mask must
                         # cover for it to read as a compartment: a nook the
                         # closing sealed at a connector's elbow is walled on
                         # two sides only and stays open
CELL_MAX = 150           # px; the largest core an erosion(2) may leave in a
                         # page-alike hole for it to read as furniture cells
                         # (page 8's comb cells core out at ~120 even chained
                         # into thousands of px): designed open space (page
                         # 6's cloud interior), line-art corridors (page 2's
                         # maze) and glass faces (page 3's cube) all keep a
                         # bigger core and stay open
SPECK_MAX = 48           # px; a transparent pocket this small, enclosed and
                         # page-alike, fills as dust whatever walls it (page
                         # 8's comb speckle); designed open gaps all read far
                         # bigger
CONTOUR_SIGMA = 2.0      # px; scale of the sticker die-cut silhouette
                         # smoothing - only page-alike px move on either
                         # side of the smoothed contour, so the trim/fill
                         # is bounded to a few px and ink is never touched
DARK_FILL = 200          # summed-channel median below which a crop is dark:
                         # it ships strokes-only (its light drawing), and its
                         # ground-alike px return to the page as original
                         # pixels (page 10's ring; round 14)


def solidify(tile, a8, mm, bg, avoid=None):
    """Solid fills, page-free crops, clean edges - in place. Returns the
    opaque mask. `avoid` marks pixels that belong to OTHERS - sibling
    elements' claims and every text line's box (re-typed and baked): no
    cell fill may cover them (a plate under a word) and no restore may
    bake them into the page (doubled text or sibling ink on drag).

    1. Interior holes are judged per component against the patched page
       beneath (`bg`): a hole holding the PAGE (a diagram's arms enclosing
       plain ground, a glyph counter) stays transparent; a hole that is a
       translucent wash over the page carries the wash's own colour and
       alpha; a hole punched through a real fill by a sibling's claim
       becomes opaque, coloured from that ground - every other element's
       ink is already erased there, so the ground IS the fill's
       continuation (page 10's disc band). A DARK crop fills every hole:
       its fill legitimately matches the ground and the page tests above
       would read it as page.
    2. Every opaque pixel outside the strong-ink hull is alpha-matted
       against that same ground (fully keyed when it IS the ground - a
       swept-in grid stub, dust - and ramped across the antialiased skirt),
       plus a 1px skirt of the source's own antialiasing is carried at
       matted alpha, so a dragged element shows a clean smooth edge and no
       copied page.
    3. Subtle mottle inside the fill (shadows of re-typed text) is
       flattened to the fill's own large-scale colour.
    """
    from scipy import ndimage

    def closed(m, r=CLOSE_R):
        # closing on a zero-padded copy: scipy's unpadded closing erodes
        # with a zero border, which chews r off every crop edge (tight
        # crops have mask ON the edge) and empties any crop thinner than
        # 2r; a border_value=1 erosion instead grows false opaque bays
        # along the border. Padding gives the honest closing.
        p = np.pad(m, r)
        return ndimage.binary_closing(p, iterations=r)[r:-r, r:-r]

    dark = bool(np.median(tile[mm].sum(axis=1)) < DARK_FILL) if mm.any() \
        else False
    if avoid is None:
        avoid = np.zeros(mm.shape, bool)
    d = np.abs(tile.astype(int) - bg.astype(int)).sum(axis=2)
    ramp = np.clip((d - KEY_LO) * 255 // max(KEY_HI - KEY_LO, 1),
                   0, 255).astype(np.uint8)
    forced = np.zeros(mm.shape, bool)  # hole fills the matting must not undo:
    # a fill copies the ground, so its distance to the ground is ~0 and the
    # keying below would take it straight back out

    def _cells(s, c):
        # furniture cells: compartments walled by PALE ink on most of
        # their boundary whose erosion cores stay tiny however long the
        # chain runs. Dark-glyph pockets (a plate in the making), corner
        # nooks the closing sealed, and anything with body all fail; so
        # does anything touching a text line or a sibling's claim (a fill
        # there is a plate under a word), or the crop border (an edge bay
        # is the world, and the clipped ring would read it as walled).
        if avoid[s][c].any() \
                or (s[0].start == 0 and c[0].any()) \
                or (s[0].stop == mm.shape[0] and c[-1].any()) \
                or (s[1].start == 0 and c[:, 0].any()) \
                or (s[1].stop == mm.shape[1] and c[:, -1].any()):
            return False
        gy0 = max(s[0].start - CELL_SHELL, 0)
        gy1 = min(s[0].stop + CELL_SHELL, mm.shape[0])
        gx0 = max(s[1].start - CELL_SHELL, 0)
        gx1 = min(s[1].stop + CELL_SHELL, mm.shape[1])
        gsl = (slice(gy0, gy1), slice(gx0, gx1))
        big = np.zeros((gy1 - gy0, gx1 - gx0), bool)
        big[s[0].start - gy0:s[0].stop - gy0,
            s[1].start - gx0:s[1].stop - gx0] = c
        # boundary coverage on a 1px ring: the comb's teeth are 1px wide,
        # so a wider ring reaches past them into the next cell and reads
        # a fully-walled cell as half-open
        ring = ndimage.binary_dilation(big, iterations=1) & ~big
        wall = a8[gsl] == 255      # mask ink, and every fill made so far:
        shell = ndimage.binary_dilation(big, iterations=CELL_SHELL) \
            & ~big & wall          # a filled cell walls its own slivers
        if (ring & wall).sum() < CELL_BOUND * ring.sum() \
                or not shell.any() \
                or float(np.median(d[gsl][shell])) > CORE_INK:
            return False
        er2 = ndimage.binary_erosion(c, iterations=2)
        if not er2.any():
            # hairline: invisible either way IF genuinely small - a 4px
            # slot between two long rules also empties under erosion, and
            # filling it would weld the rules across their designed gap.
            # A big hairline WEB is different (round 14): page 8's fine
            # sub-grid band chains thousands of 4px cells into one comp.
            # Every px of a comb web has a wall within ~3px along BOTH
            # axes; a designed slot runs open along its long axis.
            if float(c.sum()) <= CELL_MAX:
                return True
            if not ndimage.binary_erosion(c, iterations=1).any():
                return True   # a 1-2px sliver of any length: invisible
                # either way, unlike the 4px designed slot (which keeps
                # a core under erosion(1) and stays size-capped)
            return not ndimage.binary_erosion(
                c, structure=np.ones((1, 7), bool)).any() \
                and not ndimage.binary_erosion(
                    c, structure=np.ones((7, 1), bool)).any()
        il, n = ndimage.label(er2)
        return float(ndimage.sum(er2, il, range(1, n + 1)).max()) <= CELL_MAX
    def _speck(s, c):
        # pinhole speckle: a tiny enclosed page-alike pocket amid fills
        # reads as dust however it is walled (round 14: page 8's comb kept
        # hundreds of sub-cell pockets the _cells wall tests rejected).
        # Same border and avoid guards as _cells; size is the whole test.
        if c.sum() > SPECK_MAX or avoid[s][c].any():
            return False
        return not (s[0].start == 0 and c[0].any()) \
            and not (s[0].stop == mm.shape[0] and c[-1].any()) \
            and not (s[1].start == 0 and c[:, 0].any()) \
            and not (s[1].stop == mm.shape[1] and c[:, -1].any())
    if dark:
        # strokes-only (round 14): a dark crop's ground-alike "fill" IS the
        # page (page 10's disc was never a disc, just dark ground inside a
        # ring), and shipping it flat clashed with the page texture the
        # moment the element moved. The element is its LIGHT drawing:
        # strong ink plus a 2px antialias skirt, matted against the
        # ground; every other pixel returns to the page as its ORIGINAL
        # pixels (no patch candidate can redraw the mesh), so in place the
        # page stays pixel-perfect and dragged the drawing is clean line
        # art. Near-black markers ON the ring key away with the ground -
        # accepted, invisible-dark either way.
        core = mm & (d > CORE_INK)
        if not core.any():
            return mm          # a lone dark wash: nothing to key against
        strokes = ndimage.binary_dilation(core, iterations=2)
        # translucent() ran first and may have carried a wash here: those px
        # hold its fitted colour W (not the original page) at a partial alpha.
        # Leave them be - keying them would drop the wash AND baking tile back
        # into bg would stamp the deconvolved W onto the page, not the page's
        # own pixel (round 14 review).
        keepw = (a8 > 0) & (a8 < 255)
        a8[mm & ~strokes & ~keepw] = 0
        sk = mm & strokes & ~keepw
        a8[sk] = np.minimum(a8[sk], ramp[sk])
        gone = mm & (a8 == 0) & ~avoid
        bg[gone] = tile[gone]
        return mm & (a8 > 0)
    # no closing on the SILHOUETTE: closing welded a label's glyphs into
    # one slab and an outline's design gaps shut, and the fill then
    # baked the page between the glyphs into the crop (page 5's arms
    # enclosed half the page; "QUERY" shipped on a white plate)
    fh = ndimage.binary_fill_holes(mm)
    solid = fh.copy()
    retry = []
    hlab, _ = ndimage.label(fh & ~mm)
    for hi, hsl in enumerate(ndimage.find_objects(hlab), 1):
        comp = hlab[hsl] == hi
        P = tile[hsl][comp].astype(float)
        B = bg[hsl][comp].astype(float)
        # a lower size floor than translucent()'s: a small tinted hole
        # judged only by the d test below would fill opaque from the
        # ground and lose its tint, and an enclosed hole's boundary is
        # already strong evidence it belongs to the drawing
        got = _wash_solve(P, B) if comp.sum() >= HOLE_WASH_MIN else None
        if got is not None:        # a wash drawn over the page (page
            a, W = got             # 7's funnel band): carry it as one
            tile[hsl][comp] = W.round().astype(np.uint8)
            a8[hsl][comp] = int(round(a * 255))
        elif np.median(d[hsl][comp]) < MARGIN_BG \
                and not _cells(hsl, comp):
            solid[hsl][comp] = False   # the page itself: stays open
            retry.append((hsl, comp))  # ...unless a later cell fill
                                       # walls it in (see below)
        else:                      # a sibling's claim punched through
            tile[hsl][comp] = bg[hsl][comp]     # a real fill, or the
            a8[hsl][comp] = 255                 # page inside furniture
            forced[hsl][comp] = True            # cells (page 8's comb)
    # welded holes: cells that only enclose once a closing seals their
    # drain (page 8's comb drains through 1-2px antialiased channels,
    # and through its 10-20px registration-mark rail breaks). The seal
    # also welds pockets between a label's glyphs and encloses page
    # bays everywhere, so a welded hole fills ONLY as furniture cells:
    # anything else stays exactly as round 12 shipped it.
    domA = ndimage.binary_fill_holes(closed(mm, 2))
    for dom in (domA & ~mm & ~fh,
                ndimage.binary_fill_holes(closed(mm)) & ~mm & ~domA):
        blab, _ = ndimage.label(dom)
        for bi, bsl in enumerate(ndimage.find_objects(blab), 1):
            comp = blab[bsl] == bi
            if np.median(d[bsl][comp]) >= MARGIN_BG \
                    or not _cells(bsl, comp):
                retry.append((bsl, comp))
                continue
            tile[bsl][comp] = bg[bsl][comp]
            a8[bsl][comp] = 255
            forced[bsl][comp] = True
            solid[bsl][comp] = True
    # slivers between a tooth and a cell fail the wall test until the
    # cell itself fills: one retry with the filled cells counting as
    # walls picks them up, whichever loop saw them first. Tiny enclosed
    # pockets go as specks even where the wall tests keep failing.
    for rsl, comp in retry:
        comp = comp & ~forced[rsl]
        if not comp.any() or np.median(d[rsl][comp]) >= MARGIN_BG \
                or not (_cells(rsl, comp) or _speck(rsl, comp)):
            continue
        tile[rsl][comp] = bg[rsl][comp]
        a8[rsl][comp] = 255
        forced[rsl][comp] = True
        solid[rsl][comp] = True
    # furniture-band fill (round 14): page 8's fine sub-grid band DRAINS
    # to the open page, so no enclosure verdict can ever reach it. Inside
    # a densely-inked region, an open page-alike piece whose every px has
    # a wall within ~3px along BOTH axes is furniture cells however it
    # drains; a maze corridor or a glass face runs open along an axis and
    # survives the axis erosions.
    dens = ndimage.uniform_filter(mm.astype(float), size=9) > 0.3
    bandr = ndimage.binary_closing(np.pad(dens, 3), iterations=3)[3:-3, 3:-3]

    def _sided(m):
        # walls within 3px on all four sides: an open px boxed that
        # tightly is a furniture cell's interior; a corridor or a glass
        # face has at least one open side and stays
        h, w = m.shape
        L = np.zeros_like(m)
        R = np.zeros_like(m)
        U = np.zeros_like(m)
        D = np.zeros_like(m)
        for k in (1, 2, 3):
            L[:, k:] |= m[:, :-k]
            R[:, :-k] |= m[:, k:]
            U[k:, :] |= m[:-k, :]
            D[:-k, :] |= m[k:, :]
        return L & R & U & D
    walls = mm | forced
    fillpx = bandr & (a8 == 0) & (d < MARGIN_BG) & ~avoid & _sided(walls)
    tile[fillpx] = bg[fillpx]
    a8[fillpx] = 255
    forced |= fillpx
    solid |= fillpx
    # the element's real extent is the hull of its strong ink; opaque
    # pixels beyond it (a neighbouring grid run the claims swept in, page
    # 9's footer stubs, the antialiased skirt) are matted against the
    # patched ground.
    core = solid & (d > CORE_INK)
    hull = ndimage.binary_fill_holes(closed(core))
    hull = ndimage.binary_dilation(hull, iterations=2)
    if hull.sum() < 0.05 * solid.sum():
        hull = solid    # an all-pale element (a lone wash): nothing to key
    # matting domain: fully opaque pixels only - translucent()'s washes
    # already carry their own fitted alpha and must not be re-keyed
    # (their recoloured W sits further from the ground than the drawn
    # mix ever did)
    mat = solid & ~hull & (a8 == 255) & ~forced
    # a component the patch mostly REPRODUCES (median distance near
    # zero) is copied page - a swept-in grid run, page 9's footer
    # stubs, interior dust - and is keyed out whole, mispredicted
    # pixels included (a real pale stroke sits above MARGIN_INK).
    # Everything else gets the per-pixel ramp: ground-alike pixels
    # vanish, the antialiased skirt fades smoothly, pale strokes stay.
    mlab, _ = ndimage.label(mat)
    for mi, msl in enumerate(ndimage.find_objects(mlab), 1):
        mcomp = (mlab[msl] == mi)
        if np.median(d[msl][mcomp]) < MARGIN_BG:
            a8[msl][mcomp & (d[msl] < MARGIN_INK)] = 0
    a8[mat] = np.minimum(a8[mat], ramp[mat])
    # ...and a 1px skirt of the source's own antialiasing joins at
    # matted alpha: the claims cut through the skirt at arbitrary
    # pixels, which is the ragged staircase edge. The erase already
    # outruns the claim by FRINGE, so these pixels left the page.
    # ~solid: pixels INSIDE solid with a8 == 0 are ones the keying
    # above just removed, and the ramp would resurrect any of them
    # sitting past KEY_LO (a swept-in stub's mispredicted grid pixels)
    skirt = ndimage.binary_dilation(solid, iterations=1) \
        & ~solid & (a8 == 0)
    a8[skirt] = ramp[skirt]
    solid = solid | (skirt & (ramp > 0))
    # a pixel keyed fully away never left the page: hand the background
    # its ORIGINAL pixel back in place of the synthetic patch, so the
    # page stays pixel-perfect once the crop no longer covers it.
    # Not on `avoid` px: a keyed pixel there may be a sibling's ink or
    # an erased line, and the restore would duplicate it on the page.
    # Only px the patch all but reproduces anyway (d < MARGIN_BG): a
    # thin real stroke at d~90 inside a page-alike component would
    # otherwise be baked into the page and left behind on drag.
    gone = mat & (a8 == 0) & ~avoid & (d < MARGIN_BG)
    bg[gone] = tile[gone]
    # pinholes: page-alike mask fragments the keying removed read as
    # speckle once the fills around them close; a transparent
    # compartment fully walled by opaque pixels fills like any cell,
    # and a tiny enclosed pocket fills as a speck (round 14)
    elab, _ = ndimage.label(a8 == 0)
    for ei, esl in enumerate(ndimage.find_objects(elab), 1):
        comp = elab[esl] == ei
        if np.median(d[esl][comp]) >= MARGIN_BG \
                or not (_cells(esl, comp) or _speck(esl, comp)):
            continue
        tile[esl][comp] = bg[esl][comp]
        a8[esl][comp] = 255
        forced[esl][comp] = True
        solid[esl][comp] = True
    er = ndimage.binary_erosion(hull, iterations=RIM, border_value=1)
    rim = solid & hull & ~er & ~forced
    if rim.any():
        a8[rim] = np.minimum(a8[rim], ramp[rim])
    # Photoshop-style edge cleanup (round 14, after Select-and-Mask):
    # Smooth + Shift Edge + Decontaminate Colors, applied to the final
    # silhouette so a dragged element shows a clean outline.
    # (a) the sticker die-cut: a dragged element's white matte is fine
    # WHEN its contour is smooth - what reads as dirt is the ragged
    # staircase the claims cut. Smooth the silhouette at CONTOUR_SIGMA:
    # page-alike nicks INSIDE the smoothed contour fill (their px are the
    # original page, already in the tile), page-alike jags OUTSIDE it are
    # keyed away and returned to the page. Real ink is never touched on
    # either side of the contour, however thin.
    op = a8 > 0
    g = ndimage.gaussian_filter(op.astype(float), CONTOUR_SIGMA)
    nick = ~op & (g > 0.5) & (d < MARGIN_BG) & ~avoid
    a8[nick] = 255
    forced |= nick
    jag = op & (g < 0.5) & (d < MARGIN_BG) & ~forced
    a8[jag] = 0
    gone = jag & ~avoid
    bg[gone] = tile[gone]
    # ...and anti-alias the smoothed boundary: page-alike edge px take
    # the contour's own coverage as alpha, so the cut fades over 1-2px
    # instead of stepping
    op = a8 > 0
    edge = op & ~ndimage.binary_erosion(op, border_value=1) \
        & (d < MARGIN_BG) & (a8 == 255)
    a8[edge] = np.clip(((g[edge] - 0.35) / 0.3 * 255), 0, 255) \
        .astype(np.uint8)
    # (c) decontaminate: an edge pixel fading out still carries the
    # page's colour mixed into its own; pull its colour from the
    # nearest fully-opaque pixel instead, so no pale halo rides the
    # fade when the element sits on darker ground
    solid = a8 > 0
    part = solid & (a8 < 255) \
        & ndimage.binary_dilation(a8 == 0, iterations=1)
    if part.any() and (a8 == 255).any():
        iy, ix = ndimage.distance_transform_edt(
            a8 != 255, return_indices=True)[1]
        tile[part] = tile[iy[part], ix[part]]
    return solid


CLASSIFY_MODEL = "claude-sonnet-4-6"
MODEL_MISSES = []        # one entry per page the vision model could not serve


def _load_env():
    """The key lives in the repo's .env; a copy of this tool in a sibling repo
    (stock pitch) has none, so fall back to the personal-assistant one. Without
    a key every vision pass silently skips and the deck ships OCR-only."""
    from dotenv import load_dotenv
    here = os.path.dirname(os.path.abspath(__file__))
    for env in (os.path.join(here, os.pardir, ".env"),
                os.path.join(os.path.expanduser("~"), "code",
                             "personal assistant", ".env")):
        if os.path.exists(env):
            load_dotenv(env)
            if os.environ.get("ANTHROPIC_API_KEY"):
                return


def _client():
    import anthropic
    _load_env()
    return anthropic.Anthropic()
CLASSIFY_PROMPT = """This slide was flattened to a picture. I am rebuilding it \
as editable text, which means erasing a line from the picture and re-typing it.

That only works for text SET ON the slide: titles, body paragraphs, panel \
headings, captions, standalone labels sitting on empty background.

It does not work for text drawn INSIDE the artwork, where the drawing was made \
around the words: a callout whose leader line stops at them, a label in a gap \
left in an arrow or a rule, text inside a shape, anything a grid or stroke \
breaks for. Erasing those leaves a hole nothing can fill correctly.

Here are the lines found on this slide, numbered:
{lines}

Reply with ONLY the numbers of the lines drawn INSIDE the artwork, comma \
separated. Reply NONE if every line is set on the slide."""


def diagram_lines(img, lines, cache_path=None, page=0):
    """Indices of lines drawn inside the artwork, decided by a vision model.

    Every pixel statistic tried here (how far the strips above and below
    disagree, how well a fill reproduces the region, how much drawing surrounds
    the line) puts real diagram labels and ordinary body paragraphs in the same
    range. The distinction is authorial, not photometric: did the illustrator
    draw around these words? So ask something that can see the slide.

    No key, no network, or a bad reply -> nothing is flagged and every line is
    extracted, which is the old behaviour.
    """
    if not lines:
        return set()
    texts = [l["text"] for l in lines]
    if cache_path and os.path.exists(cache_path):
        hit = json.load(open(cache_path, encoding="utf-8")).get(str(page))
        # the answer is a list of line INDICES, so it is only valid for the
        # exact line list it was asked about. A change to the OCR renumbers
        # them and a stale cache would bake the wrong lines, silently.
        if isinstance(hit, dict) and hit.get("lines") == texts:
            return set(hit["drawn"])
    try:
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        listing = "\n".join(f"{n}. {l['text']}" for n, l in enumerate(lines))
        msg = _client().messages.create(
            model=CLASSIFY_MODEL, max_tokens=300,
            messages=[{"role": "user", "content": [
                {"type": "image", "source": {
                    "type": "base64", "media_type": "image/png",
                    "data": base64.b64encode(buf.getvalue()).decode()}},
                {"type": "text",
                 "text": CLASSIFY_PROMPT.format(lines=listing)}]}])
        reply = msg.content[0].text
        got = {int(n) for n in re.findall(r"\d+", reply) if int(n) < len(lines)}
    except Exception as e:                       # offline, no key, quota, junk
        print(f"  (classifier unavailable: {e}; keeping all text)", file=sys.stderr)
        MODEL_MISSES.append(f"classifier p{page}: {e}")
        return set()
    if cache_path:
        seen = (json.load(open(cache_path, encoding="utf-8"))
                if os.path.exists(cache_path) else {})
        seen[str(page)] = {"lines": texts, "drawn": sorted(got)}
        json.dump(seen, open(cache_path, "w", encoding="utf-8"))
    return got


SMALL_H = 16             # px line height; shorter lines are what the OCR garbles
                         # (page 6's footer is 15px and came out "SCALE: 1 1")
SMALL_PROMPT = """This slide was flattened to a picture, {w}x{h} pixels. I \
re-typed its text from OCR, but the smallest lines come back garbled: this \
deck's slashed zeros read as 8 ("SLIDE 01" as "SLIDE 81") and its smallest \
type as junk ("SCALE: 1:1" as "smug. 1 -1"). Here are the small lines, \
numbered, with their pixel boxes:
{lines}

Read the ACTUAL pixels at each box and give the exact text drawn there. \
Type a slashed zero as 0. If a line is partially hidden under the NotebookLM \
watermark, COMPLETE the hidden part: copy the identical text elsewhere on \
the page, or infer it from the deck's repeating footer pattern. A confident \
completion beats leaving the line cut off.

Also list any OTHER short text runs the list misses entirely (a lone "1:1" \
or "07" near a footer, a stray tick label), each with a tight pixel box. Do \
not list the watermark itself, or text that is part of a diagram.

Reply ONLY JSON:
{{"fixed": {{"<number>": "corrected text", ...}},
  "extra": [{{"text": "...", "box": [x0, y0, x1, y1]}}, ...]}}
Put under "fixed" only lines whose text needs a change."""


PROOF_PROMPT = """This is one slide of a NotebookLM deck, {w}x{h} px. An OCR \
pass read its text lines, numbered below. The deck's condensed bold face \
makes the OCR misread glyphs: "~" as "Nl" or "N", "%" as "0/0", "|" as "I", \
"DAU" as "DAIJ", "FY" as "r-Y", "$" as a junk symbol, and it sometimes \
types a word that is not on the slide at all.

{lines}

Read the actual pixels of each line and reply ONLY JSON, mapping a line \
number to its corrected text, for lines whose OCR text is wrong:
{{"fixed": {{"<number>": "exact text drawn", ...}}}}
Correct glyphs and words only. Never reorder, shorten, expand, or reword a \
line, never merge two lines, and leave a correct line out."""


def proofread_text(img, lines, cache_path=None, page=0):
    """Vision-correct EVERY extracted line's text (glyph-level only).

    fix_small_text only looks at the footer band; body text that OCRs as
    "(Nl 5% of revenue)" or "pal 10/0" used to ship as typed. A fix is
    accepted only when it stays close to the OCR length, so the model cannot
    rewrite content, and boxes are never moved. No model -> unchanged."""
    idx = [n for n, l in enumerate(lines)
           if not l["drop"] and not l.get("sure") and l["text"].strip()]
    if not idx:
        return lines
    texts = [lines[n]["text"] for n in idx]
    key = f"proof1:{page}"
    got = None
    if cache_path and os.path.exists(cache_path):
        hit = json.load(open(cache_path, encoding="utf-8")).get(key)
        if isinstance(hit, dict) and hit.get("lines") == texts:
            got = hit["reply"]
    if got is None:
        try:
            buf = io.BytesIO()
            img.save(buf, format="PNG")
            listing = "\n".join(f'{n}. "{t}"' for n, t in enumerate(texts))
            msg = _client().messages.create(
                model=CLASSIFY_MODEL, max_tokens=2500,
                messages=[{"role": "user", "content": [
                    {"type": "image", "source": {
                        "type": "base64", "media_type": "image/png",
                        "data": base64.b64encode(buf.getvalue()).decode()}},
                    {"type": "text",
                     "text": PROOF_PROMPT.format(w=img.width, h=img.height,
                                                 lines=listing)}]}])
            got = json.loads(re.search(r"\{.*\}", msg.content[0].text,
                                       re.S).group(0))
        except Exception as e:
            print(f"  (proofreader unavailable: {e})", file=sys.stderr)
            MODEL_MISSES.append(f"proofread p{page}: {e}")
            return lines
        if cache_path:
            seen = (json.load(open(cache_path, encoding="utf-8"))
                    if os.path.exists(cache_path) else {})
            seen[key] = {"lines": texts, "reply": got}
            json.dump(seen, open(cache_path, "w", encoding="utf-8"))
    n_fixed = 0
    for k, txt in (got.get("fixed") or {}).items():
        if not (str(k).isdigit() and int(k) < len(idx)):
            continue
        txt = str(txt).strip()
        l = lines[idx[int(k)]]
        old = l["text"]
        # glyph fixes only: a reply that drifts far from the OCR length is
        # the model rewording the slide, and the erase box would not fit it
        if not txt or txt == old or abs(len(txt) - len(old)) > max(4, len(old) // 4):
            continue
        ws = txt.split()
        l["words"] = ([dict(w, text=t) for w, t in zip(l["words"], ws)]
                      if len(ws) == len(l["words"])
                      else [{"text": txt, "box": l["box"]}])
        l["text"] = txt
        n_fixed += 1
    if n_fixed:
        print(f"  proofread: {n_fixed} line(s) corrected", file=sys.stderr)
    return lines


def _extend_right(arr, box, stops=(), limit=200, ignore=None):
    """Grow a line's erase box over contiguous ink to its right: the visible
    stub of a footer the watermark cut off ("SCA...") must be erased along
    with the re-typed full line, or the slide shows both. Stops at anything
    taller than the line (a panel border is structure, not text) and before
    any box in `stops` (the NEXT fragment of the footer is its own line).
    Pixels in `ignore` (the watermark pill) count as background, so a tail
    the pill half-covers ("1:1" over its top edge) is still walked over."""
    x0, y0, x1, y1 = box
    h, w = arr.shape[:2]
    lh, stop = y1 - y0, min(x1 + limit, w)
    for s in stops:
        if s[1] < y1 and y0 < s[3] and s[0] >= x1:
            stop = min(stop, s[0] - 1)
    while x1 < stop:
        nx24 = min(x1 + 24, stop)
        strip = arr[y0:y1, x1:nx24].astype(int)
        bg = np.median(strip.reshape(-1, 3), axis=0)
        ink = np.abs(strip - bg).sum(axis=2) > COLOR_DIST
        if ignore is not None:
            ink &= ~ignore[y0:y1, x1:nx24]
        cols = np.flatnonzero(ink.any(axis=0))
        if not len(cols):
            break
        nx = x1 + int(cols[-1]) + 1
        ty0, ty1 = max(y0 - lh, 0), min(y1 + lh, h)
        tall = arr[ty0:ty1, x1 + int(cols[0]):nx]
        tink = np.abs(tall.astype(int) - bg).sum(axis=2) > COLOR_DIST
        if ignore is not None:
            tink &= ~ignore[ty0:ty1, x1 + int(cols[0]):nx]
        rows = np.flatnonzero(tink.any(axis=1))
        # judged by ink OUTSIDE the line's own band: a border fills the
        # window, a stray speckle above the footer does not
        if len(rows[(rows < lh) | (rows >= 2 * lh)]) > lh:
            break                      # a border or drawing, not the line
        x1 = nx
    return (x0, y0, x1, y1)


def _snap_ink(arr, box, grow=6):
    """Tighten a sloppy model box onto the ink actually under it."""
    h, w = arr.shape[:2]
    x0, y0 = max(box[0] - grow, 0), max(box[1] - grow, 0)
    x1, y1 = min(box[2] + grow, w), min(box[3] + grow, h)
    if x1 <= x0 or y1 <= y0:
        return None
    sub = arr[y0:y1, x0:x1]
    bg = np.median(sub.reshape(-1, 3), axis=0)
    ink = np.abs(sub.astype(int) - bg).sum(axis=2) > COLOR_DIST
    t = _tight(ink)
    if t is None:
        return None
    return (x0 + t[0], y0 + t[1], x0 + t[2], y0 + t[3])


def fix_small_text(img, lines, cache_path=None, page=0, pill=None):
    """Vision-correct the smallest lines (footers, tick labels).

    The OCR misreads this deck's smallest type at any scale (measured: 4x
    re-reads recover nothing), but a vision model reads it fine. Corrected
    and added lines are marked `sure`, which exempts them from the
    MIN_TYPE_PT bake. No model -> lines pass through unchanged."""
    arr = np.asarray(img)
    band = 0.8 * img.height            # footers and their fragments live in
    # the bottom band; body text is the same height and must never be sent
    # (the model once "corrected" a body line and glued 1:1 into it)
    small = [n for n, l in enumerate(lines)
             if l["box"][3] - l["box"][1] <= SMALL_H and not l["drop"]
             and l["box"][1] > band]
    texts = [lines[n]["text"] for n in small]
    key = f"small2:{page}"   # bumped with the completion wording
    got = None
    if cache_path and os.path.exists(cache_path):
        hit = json.load(open(cache_path, encoding="utf-8")).get(key)
        if isinstance(hit, dict) and hit.get("lines") == texts:
            got = hit["reply"]
    if got is None:
        try:
            buf = io.BytesIO()
            img.save(buf, format="PNG")
            listing = "\n".join(
                f'{n}. "{lines[i]["text"]}" at {list(lines[i]["box"])}'
                for n, i in enumerate(small)) or "(none)"
            msg = _client().messages.create(
                model=CLASSIFY_MODEL, max_tokens=1500,
                messages=[{"role": "user", "content": [
                    {"type": "image", "source": {
                        "type": "base64", "media_type": "image/png",
                        "data": base64.b64encode(buf.getvalue()).decode()}},
                    {"type": "text",
                     "text": SMALL_PROMPT.format(w=img.width, h=img.height,
                                                 lines=listing)}]}])
            got = json.loads(re.search(r"\{.*\}", msg.content[0].text,
                                       re.S).group(0))
        except Exception as e:
            print(f"  (small-text reader unavailable: {e})", file=sys.stderr)
            MODEL_MISSES.append(f"small-text p{page}: {e}")
            return lines
        if cache_path:
            seen = (json.load(open(cache_path, encoding="utf-8"))
                    if os.path.exists(cache_path) else {})
            seen[key] = {"lines": texts, "reply": got}
            json.dump(seen, open(cache_path, "w", encoding="utf-8"))
    for k, txt in (got.get("fixed") or {}).items():
        if not (k.isdigit() and int(k) < len(small) and txt.strip()):
            continue
        l = lines[small[int(k)]]
        if len(txt) > len(l["text"]) + 2:      # reconstructed a hidden tail:
            l["box"] = _extend_right(          # erase the visible stub too
                arr, l["box"], [o["box"] for o in lines if o is not l],
                ignore=pill)
        ws = txt.split()
        l["words"] = ([dict(w, text=t) for w, t in zip(l["words"], ws)]
                      if len(ws) == len(l["words"])
                      else [{"text": txt, "box": l["box"]}])
        l["text"], l["sure"] = txt, True
    for e in got.get("extra") or []:
        try:
            txt, box = str(e["text"]).strip(), [int(v) for v in e["box"]]
        except (KeyError, TypeError, ValueError):
            continue
        if not txt or NOISE_RE.match(txt):
            continue
        snapped = _snap_ink(arr, box)
        if snapped is not None and (snapped[2] - snapped[0] < 8
                                    or snapped[3] - snapped[1] < 6):
            continue               # a sliver of a construction line, not text
        if snapped is None or any(
                snapped[0] < l["box"][2] and l["box"][0] < snapped[2]
                and snapped[1] < l["box"][3] and l["box"][1] < snapped[3]
                for l in lines):
            continue                   # nothing there, or already covered
        if snapped[3] - snapped[1] > 2 * (snapped[2] - snapped[0]):
            continue                   # portrait: that is a ROTATED label,
            # and typing it upright both garbles it and blocks the sideways
            # pass that reads it properly (page 8's TRACE axis)
        if snapped[1] <= band:
            continue                   # extras exist for the footer band only
        lines.append({"text": txt, "box": snapped, "drop": False,
                      "sure": True,
                      "words": [{"text": txt, "box": snapped}]})
    # a footer often OCRs as several fragments on one baseline; once the
    # sure line carries the full text, a right-neighbour fragment whose text
    # it contains is the same words twice. Absorb its span, drop the line.
    for a in lines:
        if not a.get("sure"):
            continue
        for b in lines:
            if b is a or b["drop"] or not b["text"].strip():
                continue
            ax0, ay0, ax1, ay1 = a["box"]
            bx0, by0, bx1, by1 = b["box"]
            if (min(ay1, by1) - max(ay0, by0) > 0.5 * (ay1 - ay0)
                    and 0 <= bx0 - ax1 <= 60
                    and b["text"].strip() in a["text"]
                    and len(b["text"]) < len(a["text"])):
                a["box"] = (ax0, min(ay0, by0), max(ax1, bx1), max(ay1, by1))
                b["text"] = ""
    lines = [l for l in lines if l["drop"] or l["text"].strip()]
    # ...and adjacent surviving fragments on one baseline are ONE line: the
    # footer "SLIDE 07" + "// PROJECT ..." must come back as one textbox
    merged = True
    while merged:
        merged = False
        for a in lines:
            for b in lines:
                if (a is b or a["drop"] or b["drop"]
                        or a["box"][3] - a["box"][1] > SMALL_H
                        or b["box"][3] - b["box"][1] > SMALL_H
                        or a["box"][1] <= band or b["box"][1] <= band):
                    continue           # footers only, never body text
                ax0, ay0, ax1, ay1 = a["box"]
                bx0, by0, bx1, by1 = b["box"]
                if (min(ay1, by1) - max(ay0, by0) > 0.5 * (ay1 - ay0)
                        and 0 <= bx0 - ax1 <= 40):
                    a["text"] = a["text"].rstrip() + " " + b["text"].lstrip()
                    a["box"] = (ax0, min(ay0, by0), max(ax1, bx1),
                                max(ay1, by1))
                    a["words"] = a["words"] + b["words"]
                    a["sure"] = a.get("sure", False) or b.get("sure", False)
                    lines.remove(b)
                    merged = True
                    break
            if merged:
                break
    # a reconstructed footer now CONTAINS the stray fragment the OCR read as
    # its own line ("1:1" beside "... SCALE:"): typed both, the slide shows
    # "SCALE: 1:1 1:1". A line whose box a longer sure line's erase already
    # covers, and whose text that line already carries, is that fragment.
    def swallowed(a):
        return any(b is not a and b.get("sure") and a["text"] in b["text"]
                   and a["box"][0] >= b["box"][0] - PAD
                   and a["box"][1] >= b["box"][1] - PAD
                   and a["box"][2] <= b["box"][2] + PAD
                   and a["box"][3] <= b["box"][3] + PAD
                   for b in lines)
    return [l for l in lines if l["drop"] or not swallowed(l)]


OBJECT_PROMPT = """This slide was flattened to one picture, {w}x{h} pixels. I \
am rebuilding it as an editable deck: text is re-typed separately, and each \
distinct piece of ARTWORK becomes its own movable picture element.

List the artwork objects a designer would want to move as one unit:

- FIRST decide: does the artwork read as ONE integrated diagram (a flowchart, \
a machine, a funnel, a chart, a pipeline of stages feeding into each other)? \
Then it is ONE object -- its blocks, its arrows, its stages above and below, \
and its internal labels all together. NEVER box an integrated diagram's \
parts, stages or internal connectors separately.
- Only independent illustrations sitting apart (a row of separate pictures or \
icon panels) are separate objects, and then a connector arrow floating \
between two of them is its own object.
- Small standalone marks OUTSIDE the artwork count too: a circled step \
number, a leader line pointing at the diagram, a divider bar or underline, an \
icon. A divider bar or underline near text is its own object even when it \
sits inside another object's area. Every leader/pointer line gets its own \
box; never leave one out.
- The slide's footer or status box in the margin (e.g. "SLIDE 06 // ...") is \
ALWAYS its own object, never part of a diagram, even when a diagram's area \
overlaps it.
- The page's furniture is NOT an object: background grids, margin tick marks, \
frame rules, corner crosses, watermark textures, anything that spans the slide \
and cannot be moved on its own -- and the deck's DRAFT MARKS anywhere on the \
page, even between text columns or near a footer: small plus/cross marks, \
thin grey construction or guide lines (often ending in a small circle), \
dimension bars and their tiny measurement labels, rulers, sawtooth strips. \
Never list those. A coloured divider bar or underline near text IS an object; \
a grey draft mark is not. A full-slide centrepiece illustration (a title \
page's emblem or ring) is still an object, not furniture.
- Plain text lines sitting directly on the page are not objects; they are \
re-typed separately. But a PANEL IS an object: any tinted or outlined box \
that text sits in or on -- a callout card, a labelled status box, a footer \
box in the margin. Box the panel generously (its frame and fill); ignore \
the words inside it.
- When unsure whether a decorative mark is furniture, it is furniture.

Reply with ONLY a JSON array, one item per object:
[{{"name": "short description", "box": [x0, y0, x1, y1]}}]
Boxes are in pixels of this image, generous enough to contain the whole \
object. Reply [] if there is no artwork."""


FURNITURE_RE = re.compile(r"crosshair|cross-hair|cross mark|plus mark"
                          r"|registration|notebook\s*lm|watermark|corner mark"
                          r"|tick mark|frame rule|construction|guide line"
                          r"|dimension|ruler|sawtooth|margin mark", re.I)
FURN_MAX_AREA = 20000    # px^2; a bigger interior object is real artwork even
                         # when its DESCRIPTION mentions a furniture word (page
                         # 8's "chart with crosshair and bullseye" is the chart,
                         # page 10's "reticle with crosshair ticks" is the ring)


def _furniture(o, w, h):
    """Furniture by name AND by geometry: small, or strip-thin.
    Name alone dropped page 8's chart because its description mentioned the
    crosshair drawn inside it. Corner marks and the watermark are small; a
    sawtooth ruler or frame rule is a long thin strip; a real chart, ring
    or centrepiece is both big and thick, wherever its box sits."""
    if not FURNITURE_RE.search(o["name"]):
        return False
    x0, y0, x1, y1 = o["box"]
    return ((x1 - x0) * (y1 - y0) < FURN_MAX_AREA
            or min(x1 - x0, y1 - y0) <= 2 * EDGE)


def object_boxes(img, cache_path=None, page=0):
    """The slide's artwork objects, named and boxed by a vision model, or None
    when no model is reachable (the caller falls back to the pixel rules).

    Pixel rules cannot know that a maze, a cross and a box wired together are
    ONE schematic while six panels sitting apart are SIX: that is authorial
    intent, not photometry, so ask something that can see the slide. A sloppy
    box is fine: it only CLAIMS ink components, the ink decides the cut.
    Objects the model names as page furniture anyway (corner crosshairs, the
    watermark) are dropped: extracted, they turn into junk crumb elements.
    """
    key = f"obj4:{page}"   # bumped with the draft-mark furniture rules
    if cache_path and os.path.exists(cache_path):
        hit = json.load(open(cache_path, encoding="utf-8")).get(key)
        if isinstance(hit, dict) and hit.get("wh") == list(img.size):
            return [o for o in hit["objects"]
                    if not _furniture(o, img.width, img.height)]
    try:
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        msg = _client().messages.create(
            model=CLASSIFY_MODEL, max_tokens=2000,
            messages=[{"role": "user", "content": [
                {"type": "image", "source": {
                    "type": "base64", "media_type": "image/png",
                    "data": base64.b64encode(buf.getvalue()).decode()}},
                {"type": "text",
                 "text": OBJECT_PROMPT.format(w=img.width, h=img.height)}]}])
        objects = []
        for o in json.loads(re.search(r"\[.*\]", msg.content[0].text, re.S)
                            .group(0)):
            x0, y0, x1, y1 = (int(v) for v in o["box"])
            x0, y0 = max(x0, 0), max(y0, 0)
            x1, y1 = min(x1, img.width), min(y1, img.height)
            if x1 > x0 and y1 > y0:
                objects.append({"name": str(o.get("name", "")),
                                "box": [x0, y0, x1, y1]})
    except Exception as e:                       # offline, no key, junk reply
        print(f"  (object model unavailable: {e}; pixel rules only)",
              file=sys.stderr)
        MODEL_MISSES.append(f"objects p{page}: {e}")
        return None
    if cache_path:
        seen = (json.load(open(cache_path, encoding="utf-8"))
                if os.path.exists(cache_path) else {})
        seen[key] = {"wh": list(img.size), "objects": objects}
        json.dump(seen, open(cache_path, "w", encoding="utf-8"))
    return [o for o in objects if not _furniture(o, img.width, img.height)]


OBJ_MIN_INK = 40         # ink px; a smaller component is dust, not artwork
OBJ_GROW = 12            # px; the model's boxes sit a few pixels off the ink
STRAY_FRAC = 0.15        # overlap below which a component cannot be grabbed
NEAR = 15                # px; how close a straggler must sit to a claimed mask
                         # (page 4's REJECT label hangs 11px under its blob)
PALE_DIST = 30           # pixel-to-bg distance of the palest texture an
                         # object's own extent still claims
RIDE_PAD = 4             # px of slack on that extent for a riding component
TINT_MATCH = 12          # colour distance to the page's own texture below
                         # which a contained pale slab is background, not shading
HOLE_MIN = 400           # px^2; a smaller hole in a drawing is a pin-hole and
                         # always rides (only a real window shows background)
HOLE_CONTENT = 0.03      # fraction of hole pixels far from every page colour
                         # above which the hole holds drawing, not background
TEX_RATE_MIN = 0.02      # page-interior texm rate below which the texture-
                         # hiding test is meaningless (page 6 sits at 0.015,
                         # the meshed page 10 at 0.12, ruled page 9 at 0.055)
HOLE_TEX = 0.25          # x the page's texm rate; a hole showing less of the
                         # page's texture than this HIDES it and rides (page
                         # 10's ring interior suppresses the mesh to 0.002
                         # against the page's 0.12 - the mesh analogue of the
                         # grid-hiding test, for pages with no grid synth)
HOLE_CLOSE = 12          # px; a ring drawn with tick gaps (page 10's reticle,
                         # ~24px where the arcs part around a tick) never
                         # closes for fill_holes. Dilating by this bridges
                         # gaps up to ~2x it, ONLY to find such an interior;
                         # it rides solely on the hiding tests, so an open
                         # C-shape stays open
RING_MIN = 4000          # px^2; a gap-bridged interior smaller than this is a
                         # closure artifact (a concave corner), never a disc
ABSORB_MAX = 5000        # px^2 bbox; a bigger unclaimed component is its own
                         # structure (a footer box), never a straggler
THIN_TRIM = 8            # px; rows of a leader mask's hanging piece up to
                         # this wide are stroke-plus-skirt (page 4's cone
                         # tip reads 5-8 with its antialiasing), never the
                         # label's own circle or bracket serif; the same
                         # number is the least ROWS a trimmed run may have,
                         # so a dot's 1-3 row crown is never nibbled
LEAD_GROW = 18           # extra px of slack around a leader/pointer box
LEADER_RE = re.compile(r"leader|pointer", re.I)
# an object that IS an arrow/connector ("dashed arrow from noise blob into
# funnel"), not a diagram whose description mentions its arrows: the word
# must sit in the name's first three tokens. The \b before the keyword is a
# whole-word guard - without it "narrow", "sparrow" and "disconnector" match
# on the substring and wrongly weld their masks.
ARROW_RE = re.compile(r"(?:\S+\s+){0,2}\S*\b(?:arrow|connector)", re.I)
# a chart-like object owns its grey axes, ruler and hatching; the (?<![A-Za-z])
# / (?![A-Za-z]) whole-word guards keep "flowchart", "graphic", "chartreuse"
# and "axisymmetric" from qualifying (the optional s keeps the plurals in)
CHART_RE = re.compile(
    r"(?<![A-Za-z])(?:charts?|graphs?|plots?|ax[ei]s)(?![A-Za-z])", re.I)
PROTRUDE = 40            # px a claimed component may poke beyond its box; more
                         # means the box never meant it (page 4's chunk bracket
                         # pokes 82; page 10's own reticle tick only 28)
THIN_BOX = 20            # px; a box no wider than this is a stroke box (a
                         # divider rule, a leader line), claimed under the
                         # stricter thin-box rules below
THIN_SPAN = 0.3          # of the stroke box's long axis a component must run:
                         # a divider box placed sloppily over a diagram's
                         # dashed tips must not tear the dashes off it
GREY_SPREAD = 15         # summed channel spread at or below which a stroke's
                         # ink is GREY: a stroke box claiming grey ink has
                         # boxed the page's draft furniture (the construction
                         # lines between page 2's text columns), never the
                         # deck's coloured dividers and leaders
GREY_INK = 300           # summed-channel depth below the reference past which
                         # a grey residue is solid drawn ink (page 5's black
                         # tick marks, ~660 deep), never a pale page-pattern
                         # mismatch (the grid remnants sit at 40-90)
HALO = 4                 # px around claimed ink that still belongs to it
FRINGE = 2               # px the ERASE outruns the claim by: the antialiased
                         # skirt sits under every ink threshold and ghosts the
                         # drawing's outline if it stays in the background
DIM_DIST = 12            # ink threshold against a FLAT background, where even
                         # this dim a stroke reads clearly (above JPEG noise)
MERGE_NEAR = 12          # px; two big masses this close are one drawing
                         # (page 9's separate panels sit 16px apart and stay)
MERGE_MIN = 12000        # px^2 of bbox below which a mass is an annotation


def vision_elements(arr, keep_out, objects, baked=(), grid=None):
    """(box, mask) per artwork object, ink components assigned to the model's
    boxes.

    Each connected ink component goes whole to the box holding most of it (the
    smaller box on a tie, so an annotation sitting on the main diagram's box
    still comes off on its own); ink no box claims is page furniture and stays
    in the background. The mask is the component's own pixels, so the crop and
    the erase agree exactly: nothing is duplicated, nothing is left behind.

    The model's boxes are approximate and ELEM_DIST is blind to pale texture,
    so three repair passes follow. An EMPTY box grabs the best component still
    touching it (page 4's circled 2 sat 37px below its box). A straggler that
    overlaps a box and sits within NEAR px of that box's claimed ink joins it
    (the REJECT label under its blob). And faint ink is added where it sits
    wholly inside a claimed object's own extent (page 8's hatched band, the
    card fills) or, for a still-empty box, wholly inside the box (page 1's
    divider bar, fainter than ELEM_DIST ever sees). Page furniture survives
    every pass because it always runs beyond any one box.
    """
    from scipy import ndimage
    h, w = arr.shape[:2]
    raw = [o["box"] for o in objects]
    # a leader/pointer box is thin and the model places it sloppily (page 4's
    # step-2 leader box sat 25px under the real line, which left the line to
    # be straggler-absorbed into the diagram): grow those further, the extra
    # slack only ever reaches the stroke the box was drawn for
    grown = []
    for o in objects:
        b = o["box"]
        g = OBJ_GROW + (LEAD_GROW if LEADER_RE.search(o["name"]) else 0)
        grown.append((max(b[0] - g, 0), max(b[1] - g, 0),
                      min(b[2] + g, w), min(b[3] + g, h)))
    areas = [(b[2] - b[0]) * (b[3] - b[1]) for b in raw]
    masks = [np.zeros((h, w), bool) for _ in objects]

    strong = _ink_mask(arr, keep_out)
    labels, _ = ndimage.label(ndimage.binary_dilation(strong,
                                                      iterations=ELEM_DILATE))
    comps = []               # (ys, xs, bbox, share per object, centroid)
    for i, sl in enumerate(ndimage.find_objects(labels), 1):
        part = strong[sl] & (labels[sl] == i)
        if part.sum() < OBJ_MIN_INK:
            continue
        ys, xs = np.nonzero(part)
        ys, xs = ys + sl[0].start, xs + sl[1].start
        comps.append([ys, xs,
                      (int(xs.min()), int(ys.min()),
                       int(xs.max()) + 1, int(ys.max()) + 1),
                      [float(((xs >= b[0]) & (xs < b[2])
                              & (ys >= b[1]) & (ys < b[3])).mean())
                       for b in grown],
                      (float(xs.mean()), float(ys.mean()))])

    def best_of(share):
        return max((((f, -areas[j]), j) for j, f in enumerate(share)),
                   default=((0.0, 0), -1))

    def interior(b):
        return (b[0] > EDGE and b[1] > EDGE
                and b[2] < w - EDGE and b[3] < h - EDGE)

    def anchored(j, c):
        # centroid inside the box, or the component's body covering a real
        # share of it: the model's box often sits sloppily off its panel
        bx = raw[j]
        if bx[0] <= c[4][0] < bx[2] and bx[1] <= c[4][1] < bx[3]:
            return True
        ox = min(bx[2], c[2][2]) - max(bx[0], c[2][0])
        oy = min(bx[3], c[2][3]) - max(bx[1], c[2][1])
        return (ox > 0 and oy > 0
                and ox * oy >= 0.1 * (bx[2] - bx[0]) * (bx[3] - bx[1]))

    def edge_junk(bx):
        # rule-shaped or tiny: in the edge band that is the frame's own
        # hairline under a footer, or a stray cross mark, never a panel
        x0, y0, x1, y1 = bx
        hh, ww = y1 - y0, x1 - x0
        return (ww * hh < ELEM_MIN_AREA
                or (hh <= RULE_MAX_H and ww > RULE_ASPECT * hh)
                or (ww <= RULE_MAX_H and hh > RULE_ASPECT * ww))

    def grey_ink(ys, xs):
        px = arr[ys, xs].astype(int)
        d = np.abs(px - np.median(arr.reshape(-1, 3), axis=0)).sum(axis=1)
        core = px[d >= np.quantile(d, 0.75)] if len(px) > 4 else px
        m = np.median(core, axis=0)
        return abs(m[0] - m[1]) + abs(m[1] - m[2]) <= GREY_SPREAD

    def thin_ok(j, c):
        # a stroke box (divider rule, leader line) claims only a component
        # that runs a real fraction of it AND is drawn in the deck's own
        # coloured ink: page 7's divider box sat over the diagram's dashed
        # tips, and page 2's boxed the grey draft line between its columns
        b = raw[j]
        bw, bh = b[2] - b[0], b[3] - b[1]
        if min(bw, bh) > THIN_BOX:
            return True
        # a non-leader stroke box must actually TOUCH its component: page
        # 7's divider box sat wholly above the diagram's dashed tips and
        # claimed their antialiased fringe as a phantom element. (A title
        # rule whose box misses it entirely still comes back through the
        # unclaimed-rule pass.) Leader boxes are exempt: LEAD_GROW exists
        # because the model places those beside their stroke.
        if not LEADER_RE.search(objects[j]["name"]) \
                and (min(b[2], c[2][2]) <= max(b[0], c[2][0])
                     or min(b[3], c[2][3]) <= max(b[1], c[2][1])):
            return False
        ext = c[2][2] - c[2][0] if bw >= bh else c[2][3] - c[2][1]
        return (ext >= THIN_SPAN * max(bw, bh)
                and not grey_ink(c[0], c[1]))

    def is_baked(c):
        # the component IS a baked text line (page 4's REJECT caption): it
        # belongs WITH its drawing. Never box-claimed, box-grabbed or turned
        # into a rule; the straggler/force-claim passes hand it to the
        # nearest drawing.
        return any(c[2][0] >= b[0] - PAD and c[2][1] >= b[1] - PAD
                   and c[2][2] <= b[2] + PAD and c[2][3] <= b[3] + PAD
                   for b in baked)

    taken = [False] * len(comps)
    claimed_to = {}
    leads = [j for j, o in enumerate(objects) if LEADER_RE.search(o["name"])]
    for ci, c in enumerate(comps):     # most of it inside a box: claimed.
        if is_baked(c):
            continue
        (frac, _), j = best_of(c[3])
        # a stroke-shaped component that a leader box also wants belongs to
        # the label, not to whatever artwork box happened to swallow more of
        # it: page 4's leader 2 ran into the noise blob's box, the blob
        # claimed line and dot, and dragging the blob took the label's
        # leader with it. The model places leader boxes sloppily, so the
        # label's own share of its line can sit well under the usual half.
        need = 0.5
        if min(c[2][2] - c[2][0], c[2][3] - c[2][1]) <= THIN_BOX:
            lead = [(c[3][jj], -areas[jj], jj) for jj in leads
                    if c[3][jj] >= 2 * STRAY_FRAC and thin_ok(jj, c)]
            if lead:
                frac, _, j = max(lead)
                need = 2 * STRAY_FRAC
        # A component in the page's edge band is suspect furniture: a
        # generous box grown past the frame catches strips running just
        # outside it (page 1's footer panel got the frame's sawtooth). Those
        # must be anchored to the box AND shaped like content (page 2's
        # footer hairline is anchored yet still the frame's own rule).
        # Interior components are never furniture.
        if frac >= need and thin_ok(j, c) \
                and (interior(c[2])
                     or (anchored(j, c) and not edge_junk(c[2]))):
            masks[j][c[0], c[1]] = True
            taken[ci] = True
            claimed_to[ci] = j
    # a thin annotation the model drew no box for welds into a neighbour's
    # claim when most of its ink falls inside that box (page 4's chunk
    # bracket sat 63% inside the diagram's box, so the diagram took it and
    # the carve later tore its circled 4 off). A claimed component poking
    # far beyond the box, small next to the box's real mass, was never that
    # box's: it becomes its own element.
    annot = []
    for ci, j in claimed_to.items():
        c = comps[ci]
        gx0, gy0, gx1, gy1 = grown[j]
        prot = max(gx0 - c[2][0], c[2][2] - gx1, gy0 - c[2][1], c[2][3] - gy1)
        # ...unless the component carries text left baked in the artwork
        # (page 6's PIPELINE label pokes past the diagram's box): baked text
        # must ride with its drawing, never split off
        if any(b[0] < c[2][2] and c[2][0] < b[2]
               and b[1] < c[2][3] and c[2][1] < b[3] for b in baked):
            continue
        if prot > PROTRUDE and len(c[0]) < 0.3 * int(masks[j].sum()):
            m2 = np.zeros((h, w), bool)
            m2[c[0], c[1]] = True
            masks[j] &= ~m2
            annot.append(m2)
    # an empty box grabs its best untaken component - and carves back any of
    # its own ink another box's claim swallowed (page 4's circled 4 came
    # welded to the chunk grid's bracket, which the diagram claimed whole;
    # page 7's footer panel had a corner strip inside the diagram's mask).
    # Furniture is never grabbed: it hugs a page edge (page 4's corner mark)
    # without covering the box (page 1's frame sawtooth vs the footer panel).
    for j, m in enumerate(masks):
        if m.any():
            continue
        cand = [(c[3][j], ci) for ci, c in enumerate(comps)
                if not taken[ci] and c[3][j] >= STRAY_FRAC and thin_ok(j, c)
                and not is_baked(c)
                and (interior(c[2])
                     or (anchored(j, c) and not edge_junk(c[2])))
                and ((c[2][2] - c[2][0]) * (c[2][3] - c[2][1])
                     <= 6 * (grown[j][2] - grown[j][0])
                     * (grown[j][3] - grown[j][1]))]
        if cand:
            ci = max(cand)[1]
            m[comps[ci][0], comps[ci][1]] = True
            taken[ci] = True
        # a stroke box never carves: its grown rectangle bites a strip out
        # of whatever drawing runs past it (page 7's empty divider box took
        # a 10-row slice off the diagram's funnel tips; a leader's LEAD_GROW
        # rectangle would bite a 30px slab)
        b = raw[j]
        if min(b[2] - b[0], b[3] - b[1]) <= THIN_BOX:
            continue
        gx0, gy0, gx1, gy1 = grown[j]
        for m2 in masks:
            if m2 is m:
                continue
            part = m2[gy0:gy1, gx0:gx1]
            n = int(part.sum())
            if n >= OBJ_MIN_INK and n < 0.5 * int(m2.sum()):
                m[gy0:gy1, gx0:gx1] |= part
                m2[gy0:gy1, gx0:gx1] = False
                break
    # an unclaimed divider rule away from the page edge is an element in its
    # own right: every page of a deck carries one under its title, and the
    # model does not reliably box them. Page furniture (frame rules, margin
    # marks) hugs the edges and is excluded by EDGE. Decided BEFORE the
    # straggler absorb below, or a rule near a drawing melts into it.
    for ci, c in enumerate(comps):
        x0, y0, x1, y1 = c[2]
        hgt, wid = y1 - y0, x1 - x0
        # grey vetoed: an unboxed grey hairline is the page's own draft
        # furniture, not a divider (the deck's real dividers are coloured);
        # baked text is never a rule either
        if (not taken[ci] and hgt <= RULE_MAX_H and wid > RULE_ASPECT * hgt
                and wid >= RULE_MIN_W and x0 > EDGE and y0 > EDGE
                and x1 < w - EDGE and y1 < h - EDGE
                and not grey_ink(c[0], c[1]) and not is_baked(c)):
            m2 = np.zeros((h, w), bool)
            m2[c[0], c[1]] = True
            masks.append(m2)
            taken[ci] = True
    # ...and the split-out annotations join the pool here, AFTER the per-box
    # repair passes (whose indexing follows `objects`) and exempt from the
    # big-merge below, or the bracket melts straight back into the diagram
    # it protrudes from.
    annot_idx = set()
    for m2 in annot:
        annot_idx.add(len(masks))
        masks.append(m2)
    # the footer/status box is always its own object (the model is told so);
    # a leader/circled-number mask is an ANNOTATION whatever its size. Used
    # by the straggler, big-merge and touch-join passes below.
    feet = {i for i, o in enumerate(objects)
            if re.search(r"footer|status box|leader|pointer|circled",
                         o["name"], re.I)} | annot_idx
    # stragglers join nearby claimed ink, iteratively: page 4's REJECT label
    # hangs 11px under its blob, and the model's box cut page 1's diagram at
    # its bottom tip - a CHAIN of dashes, each too far from the box to overlap
    # it but within NEAR of the dash before, so one sweep is not enough. Page
    # furniture never joins: the ruler, the frame and the corner marks all
    # hug a page edge.
    changed = True
    while changed:
        changed = False
        for ci, c in enumerate(comps):
            if taken[ci]:
                continue
            x0, y0, x1, y1 = c[2]
            if x0 <= EDGE or y0 <= EDGE or x1 >= w - EDGE or y1 >= h - EDGE:
                continue
            if (x1 - x0) * (y1 - y0) > ABSORB_MAX:
                continue               # a footer box, a panel: its own thing
            win = (slice(max(y0 - NEAR, 0), y1 + NEAR),
                   slice(max(x0 - NEAR, 0), x1 + NEAR))
            # joins whichever mask has the most ink nearby, not the first in
            # list order: a circled label's leftover arc must rejoin the
            # label, not the diagram whose bracket also runs close
            near = max(masks, key=lambda m: int(m[win].sum()), default=None)
            if near is not None and near[win].any():
                near[c[0], c[1]] = True
                taken[ci] = True
                changed = True

    faint = _ink_mask(arr, keep_out, COLOR_DIST) & ~strong
    # a box still empty holds a mark fainter than ELEM_DIST ever sees (page
    # 1's divider bar): rescue whole faint components sitting wholly inside
    # it. The page's faint furniture always runs beyond any one box.
    empties = [j for j in range(len(objects)) if not masks[j].any()]
    if empties:
        flab, _ = ndimage.label(faint)
        empties.sort(key=lambda j: (grown[j][2] - grown[j][0])
                     * (grown[j][3] - grown[j][1]))
        for i, sl in enumerate(ndimage.find_objects(flab), 1):
            by0, by1, bx0, bx1 = (sl[0].start, sl[0].stop,
                                  sl[1].start, sl[1].stop)
            for j in empties:
                gx0, gy0, gx1, gy1 = grown[j]
                if bx0 >= gx0 and by0 >= gy0 and bx1 <= gx1 and by1 <= gy1:
                    # an edge-band faint strip inside a box grown past the
                    # frame is the frame's own furniture (page 1's sawtooth
                    # ruler under the footer panel, the hairline over page
                    # 2's footer): it must centre INSIDE the model's raw box
                    # AND be shaped like content, and a stroke box only ever
                    # rescues its own coloured stroke
                    bb = (bx0, by0, bx1, by1)
                    cx, cy = (bx0 + bx1) / 2, (by0 + by1) / 2
                    rb = raw[j]
                    if not (interior(bb)
                            or (rb[0] <= cx < rb[2] and rb[1] <= cy < rb[3]
                                and not edge_junk(bb))):
                        continue
                    part = faint[sl] & (flab[sl] == i)
                    if part.sum() < 20:
                        break
                    ys2, xs2 = np.nonzero(part)
                    ys2, xs2 = ys2 + sl[0].start, xs2 + sl[1].start
                    # an edge-band faint GREY strip is the frame's own
                    # hairline furniture (page 4's footer sat between two
                    # frame rules and the rescue shipped them as an element);
                    # the deck's real rescued marks are coloured
                    if not interior(bb) and grey_ink(ys2, xs2):
                        continue
                    if not thin_ok(j, [ys2, xs2, bb]):
                        continue
                    masks[j][sl][part] = True
                    break
    # the model sometimes splits one connected drawing across two boxes (page
    # 4's source row vs the funnel it feeds): two BIG claimed masses nearly
    # touching are one drawing, so a designer can move it whole. Annotations
    # stay separate: a leader line, a circled number or a divider rule is
    # small or flat and never qualifies.
    def big(t):
        return (t and t[3] - t[1] > RULE_MAX_H + 4
                and (t[2] - t[0]) * (t[3] - t[1]) >= MERGE_MIN)

    # the footer/status box is always its own object (the model is told so);
    # a diagram reaching down to it must not weld it in. A leader/circled-
    # number mask is an ANNOTATION whatever its size: page 4's leader box
    # claimed the chunk bracket whole (making it big) and the merge welded
    # bracket, leader and circled 4 straight back into the diagram.
    merging = True
    while merging:
        merging = False
        for a in range(len(masks)):
            if a in feet:
                continue
            ta = _tight(masks[a])
            if not big(ta):
                continue
            for b in range(a + 1, len(masks)):
                if b in feet:
                    continue
                tb = _tight(masks[b])
                if not big(tb):
                    continue
                if not (ta[0] - MERGE_NEAR < tb[2] and tb[0] - MERGE_NEAR < ta[2]
                        and ta[1] - MERGE_NEAR < tb[3]
                        and tb[1] - MERGE_NEAR < ta[3]):
                    continue
                x0, y0 = min(ta[0], tb[0]), min(ta[1], tb[1])
                x1, y1 = max(ta[2], tb[2]), max(ta[3], tb[3])
                win = (slice(max(y0 - 1, 0), y1 + 1), slice(max(x0 - 1, 0), x1 + 1))
                if (ndimage.binary_dilation(masks[a][win],
                                            iterations=MERGE_NEAR)
                        & masks[b][win]).any():
                    masks[a] |= masks[b]
                    masks[b][:] = False
                    ta = _tight(masks[a])
                    merging = True

    # text the classifier left baked in the artwork belongs to SOME object:
    # force-claim its ink into the nearest mask, past every furniture guard
    # (page 2's VERIFY sits inside the border band the absorb pass refuses).
    tight_now = [(t, m) for m in masks for t in [_tight(m)] if t]
    for b in baked:
        bx0, by0, bx1, by1 = b
        if not tight_now:
            break
        blob = strong[by0:by1, bx0:bx1].copy()
        for m in masks:
            blob &= ~m[by0:by1, bx0:bx1]
        if not blob.any():
            continue
        cx, cy = (bx0 + bx1) / 2, (by0 + by1) / 2
        _, m = min(tight_now,
                   key=lambda p: (max(p[0][0] - cx, 0, cx - p[0][2]) ** 2
                                  + max(p[0][1] - cy, 0, cy - p[0][3]) ** 2))
        m[by0:by1, bx0:bx1] |= blob

    # pale texture and unclaimed dust ride with the object whose extent wholly
    # contains them: page 8's shaded band sits under COLOR_DIST where no
    # component test can see it, and the OBJ_MIN_INK floor drops the specks of
    # a dashed outline. WHOLE components only, never a bbox crop of the pale
    # field: the page's own tinted grid panels cross any big object's extent,
    # and cropping them in is how a crop ends up carrying slabs of background.
    # Runs AFTER the merge so halo slack never counts toward the merge
    # distance (page 9's panels sit 16px apart and must stay separate).
    loose = _ink_mask(arr, keep_out, PALE_DIST)
    for m in masks:
        loose &= ~m
    tights = sorted(((t, m) for m in masks for t in [_tight(m)] if t),
                    key=lambda p: (p[0][2] - p[0][0]) * (p[0][3] - p[0][1]))
    # a contained component can still be the page's background showing through
    # a CLOSED shape (the strokes seal a slab of grid into its own component).
    # Colour tells it from the drawing's own shading: page texture repeats
    # outside the objects, a shading band does not.
    outside = loose.copy()
    for t, _ in tights:
        outside[max(t[1] - RIDE_PAD, 0):t[3] + RIDE_PAD,
                max(t[0] - RIDE_PAD, 0):t[2] + RIDE_PAD] = False
    olab, _ = ndimage.label(outside)
    outmeds, texture = [], False
    for i, osl in enumerate(ndimage.find_objects(olab), 1):
        opart = outside[osl] & (olab[osl] == i)
        if opart.sum() >= 100:
            outmeds.append(np.median(arr[osl][opart].reshape(-1, 3), axis=0))
            # does this texture live in the page INTERIOR (a grid, a mesh) or
            # only in the border band (the ruler's glow, the frame's)?
            if (min(osl[1].stop, w - EDGE) - max(osl[1].start, EDGE) > 0
                    and min(osl[0].stop, h - EDGE) - max(osl[0].start, EDGE) > 0):
                texture = True
    # per-pixel map of the page's own texture colours (the ruled grid): the
    # halo and dilation rings below must never carry these, or every element
    # cut from a ruled page shows slabs of grid around its borders
    texm = np.zeros((h, w), bool)
    for om in outmeds:
        # at TINT_MATCH exactly: looser catches more border grid but starts
        # rejecting a rule's own pale fringe, which then ghosts behind it
        texm |= np.abs(arr.astype(int) - om).sum(axis=2) <= TINT_MATCH
    llab, _ = ndimage.label(loose)
    for i, sl in enumerate(ndimage.find_objects(llab), 1):
        for t, m in tights:
            if (sl[1].start >= t[0] - RIDE_PAD and sl[0].start >= t[1] - RIDE_PAD
                    and sl[1].stop <= t[2] + RIDE_PAD
                    and sl[0].stop <= t[3] + RIDE_PAD):
                part = loose[sl] & (llab[sl] == i)
                med = np.median(arr[sl][part].reshape(-1, 3), axis=0)
                if any(np.abs(med - om).sum() <= TINT_MATCH for om in outmeds):
                    break              # page texture: stays background
                m[sl] |= part
                loose[sl] &= ~part
                break
    # ...and loose ink hugging a mask joins it regardless: an antialiased halo
    # can touch page furniture and label as one run-beyond component, but the
    # strip right around claimed ink is the element's own edge, and left
    # behind it reprints the drawing as a ghost ring once the element moves.
    for _, m in tights:
        ring = ndimage.binary_dilation(m, iterations=HALO) & loose & ~texm
        m |= ring
        loose &= ~ring

    # on a PLAIN ground (no texture in the page interior) ink even dimmer
    # than PALE_DIST still reads clearly, and left behind it reprints the
    # drawing once the element moves (page 1's dim lattice arms sit ~12 off
    # the dark page). Reach it by connectivity from the claimed ink, bridging
    # dashed-stroke gaps; the border band is zeroed so the flood can never
    # swallow the frame, the ruler or the corner marks. A textured interior
    # (a grid, page 10's mesh gradient) makes this threshold meaningless, and
    # texture hides dim residue anyway: skip.
    if not texture:
        dim = _ink_mask(arr, keep_out, DIM_DIST)
        for om in outmeds:
            dim &= np.abs(arr.astype(int) - om).sum(axis=2) > TINT_MATCH
        for m in masks:
            dim &= ~m
        dim = ndimage.binary_dilation(dim, iterations=2)
        dim[:EDGE] = dim[-EDGE:] = False
        dim[:, :EDGE] = dim[:, -EDGE:] = False
        for _, m in tights:
            add = ndimage.binary_propagation(m, mask=dim | m) & dim
            m |= add
            dim &= ~add

    # an object the model NAMES an arrow or connector bridges drawings: the
    # noise blob, the dashed arrow and the funnel it feeds are one
    # composition to a designer (page 4), even though each got its own box,
    # and dragging the diagram must take all three. Weld the connector and
    # every non-annotation mask its ink touches. Name-anchored to the first
    # words, because a big diagram's DESCRIPTION mentions its arrows too.
    for j, o in enumerate(objects):
        if not ARROW_RE.match(o["name"]) or j in feet or not masks[j].any():
            continue
        linked = [k for k, m2 in enumerate(masks)
                  if k != j and k not in feet and m2.any()
                  and (ndimage.binary_dilation(masks[j],
                                               iterations=MERGE_NEAR)
                       & m2).any()]
        if linked:
            tgt = max(linked + [j], key=lambda k: int(masks[k].sum()))
            for k in linked + [j]:
                if k != tgt:
                    masks[tgt] |= masks[k]
                    masks[k][:] = False

    # an annotation cut in two by rival claims re-unites: a small mask whose
    # ink sits within touching distance of another mask's ink is one drawing
    # with it (page 4's circled 4 was carved off the bracket its stub curves
    # into). Only a SMALL mask may initiate: independent panels sit further
    # apart than HALO and are never dragged in.
    joining = True
    while joining:
        joining = False
        for mi, m in enumerate(masks):
            t = _tight(m)
            if not t or (t[2] - t[0]) * (t[3] - t[1]) > ABSORB_MAX:
                continue
            win = (slice(max(t[1] - HALO, 0), t[3] + HALO),
                   slice(max(t[0] - HALO, 0), t[2] + HALO))
            for m2i, m2 in enumerate(masks):
                if m2 is m or not m2[win].any():
                    continue
                # reuniting is for annotation pieces cut apart by rival
                # claims: a label's leader must never melt into the artwork
                # it points at, nor artwork into a label (page 4's leader 2
                # sat within a halo of the noise blob it points at)
                if (mi in feet) != (m2i in feet):
                    continue
                if (ndimage.binary_dilation(m[win], iterations=HALO)
                        & m2[win]).any():
                    m2 |= m
                    m[:] = False
                    joining = True
                    break

    # a label's leader accretes scraps of neighbouring artwork on the way
    # here (page 4's funnel cone tip reached the leader dot through the
    # straggler and dim passes). Against the mask's own line rows, a piece
    # hanging outside that band which TAPERS OFF at its far end -- a run of
    # rows thinner than any dot -- is such a scrap: the label's own hanging
    # shapes (a dot's arc, a circled number, a bracket's serif) are all wide
    # at their far row, and a run consuming a whole piece is a bare stroke
    # of the label's own (a bracket's plain bar). Scraps go to the nearest
    # artwork mask; with no artwork nearby they stay put rather than orphan.
    # Horizontal leaders only: this deck draws no vertical ones.
    for j in leads:
        if j >= len(masks) or not masks[j].any():
            continue
        sub = masks[j]
        wds = sub.sum(axis=1)
        long_rows = np.flatnonzero(wds >= max(0.5 * wds.max(), THIN_BOX))
        if not long_rows.size:         # no row wide enough to be the leader's
            continue                   # own horizontal stroke (a vertical or
            # sub-THIN_BOX leader, or a mask whittled to a dot by an earlier
            # pass): there is no line band to trim scraps against
        band0, band1 = int(long_rows.min()), int(long_rows.max()) + 1
        hang = sub.copy()
        hang[band0:band1] = False
        lab2, _ = ndimage.label(hang)
        for i2, sl2 in enumerate(ndimage.find_objects(lab2), 1):
            piece = hang[sl2] & (lab2[sl2] == i2)
            # a piece that never reaches the line's band is not hanging off
            # the label at all -- it is disconnected scrap (the cone tip
            # arrives as a chain of broken fragments) and moves out whole.
            # A piece touching the band (the dot, the circle, a bracket
            # bar) only ever loses a long thin far-end taper, and a taper
            # that IS the whole piece is the label's own bare stroke.
            if sl2[0].stop == band0 or sl2[0].start == band1:
                widths = piece.sum(axis=1)
                rows = (range(len(widths)) if sl2[0].stop <= band0
                        else range(len(widths) - 1, -1, -1))
                trim = []
                for r in rows:
                    if widths[r] > THIN_TRIM:
                        break
                    trim.append(r)
                if len(trim) < THIN_TRIM \
                        or len(trim) == int((widths > 0).sum()):
                    continue
                cut = np.zeros_like(piece)
                cut[trim] = piece[trim]
            else:
                cut = piece
            t2 = _tight(cut)
            win = (slice(max(sl2[0].start + t2[1] - NEAR, 0),
                         sl2[0].start + t2[3] + NEAR),
                   slice(max(sl2[1].start + t2[0] - NEAR, 0),
                         sl2[1].start + t2[2] + NEAR))
            art = [(int(m2[win].sum()), k2) for k2, m2 in enumerate(masks)
                   if k2 not in feet and k2 != j and m2[win].any()]
            if not art:
                continue
            full = np.zeros((h, w), bool)
            full[sl2] = cut
            masks[max(art)[1]] |= full
            sub &= ~full

    # a hole in a drawing is either its own fill (a card's white, a tower's
    # dark, page 9's pale radar arcs: rides, or the crop turns see-through)
    # or the page showing through a closed stroke (a slab of grid: STAYS, or
    # the crop carries background). Colour and content decide: page texture
    # repeats outside the objects, a drawing's own content does not. And a
    # hole holding ANOTHER element's ink is always punched open, or this crop
    # carries a picture of that element on top of the element itself (page
    # 4's circled 4 rode twice).
    pagecols = outmeds + [np.median(arr.reshape(-1, 3), axis=0)]
    # the 2px growth covers antialiased edges; growth pixels that ARE the
    # page's texture colour are the grid showing at the border, not edge
    parts = []
    for m in masks:
        if not m.any():
            continue
        for p in _split_rules(m):
            g = ndimage.binary_dilation(p, iterations=2)
            g &= p | ~texm
            parts.append(g)
    allp = np.zeros((h, w), bool)
    for p in parts:
        allp |= p
    # the page's texture rate in the interior: the reference for "this hole
    # HIDES the page's texture" on an unruled textured page (page 10's mesh),
    # where no grid synth exists to say what showing-through must look like
    inr = texm[EDGE:h - EDGE, EDGE:w - EDGE]
    tex_rate = float(inr.mean()) if texture and inr.size else 0.0
    out = []
    texrode = []           # per part: an interior rode by TEXTURE-hiding.
    # Under such an interior the page's ground is unknowable (the texture is
    # hidden and no grid models it), so the caller fills it smooth instead
    # of letting a ring-scored candidate paint a comb through it (page 10).
    for part in parts:
        tr = False
        others = allp & ~part
        others_d = ndimage.binary_dilation(others, iterations=2)
        filled = ndimage.binary_fill_holes(part)
        holes, _ = ndimage.label(filled & ~part)
        for hi, hsl in enumerate(ndimage.find_objects(holes), 1):
            hole = holes[hsl] == hi
            # every verdict below reads the hole MINUS other elements' ink:
            # page 4's circled 4 sat in the bracket's hole and its own dark
            # pixels made the hole read as "drawing", so it rode twice
            own = hole & ~others_d[hsl]
            opened = False
            if hole.sum() >= HOLE_MIN and own.any():
                # a pin-hole in a dense drawing (page 8's graph-paper cells)
                # never opens: nothing shows THROUGH that
                rides = False
                if grid is not None:
                    # on a ruled page the synth says what showing-through
                    # MUST look like: the grid's lines crossing the hole. A
                    # hole that hides them is the drawing's own opaque fill
                    # (page 9's panel interiors read as page-white by
                    # colour, but no grid runs through them) and rides.
                    g = grid[hsl][own]
                    onl = np.abs(g - np.median(grid.reshape(-1, 3), axis=0)) \
                        .sum(axis=1) > GRID_LINE_MIN
                    rides = onl.sum() >= 40 and float(
                        np.abs(arr[hsl][own].astype(float) - g)
                        .sum(axis=1)[onl].mean()) > 2 * TINT_MATCH
                if not rides and tex_rate >= TEX_RATE_MIN:
                    # ...and on a meshed page the texture itself says it: a
                    # hole showing almost none of a texture the page repeats
                    # everywhere is an opaque fill over it (page 10's ring
                    # interior suppresses the mesh) and rides
                    rides = float(texm[hsl][own].mean()) < HOLE_TEX * tex_rate
                    tr |= rides and hole.sum() >= RING_MIN
                if not rides:
                    px = arr[hsl][own].astype(int)
                    dist = np.min([np.abs(px - pc).sum(axis=1)
                                   for pc in pagecols], axis=0)
                    if (dist > 2 * TINT_MATCH).mean() <= HOLE_CONTENT:
                        # no drawing of its own (page 9's arcs, a bucket's
                        # fill would ride): page colour means page
                        med = np.median(px, axis=0)
                        opened = any(np.abs(med - pc).sum() <= TINT_MATCH
                                     for pc in pagecols)
            if opened:
                filled[hsl] &= ~hole
            else:
                # riding - but another element's own pixels stay out of this
                # crop, or it carries a picture of that element on top of
                # the element itself (page 10's divider slots out of the
                # ring's disc and covers its slot exactly when in place)
                ov = hole & others_d[hsl]
                if ov.any():
                    filled[hsl] &= ~ov
        # a ring drawn with tick gaps never closes for fill_holes, yet its
        # interior is an opaque disc in the source (page 10 suppresses the
        # mesh inside it). Bridge the gaps by DILATION just to find such an
        # interior (closing fails: erosion cuts the thin arc-gap bridges),
        # then grow the found core back out to the stroke. It rides only on
        # the strength of the hiding tests, so an open shape whose inside
        # really shows the page stays open.
        dil = ndimage.binary_dilation(part, iterations=HOLE_CLOSE)
        dfil = ndimage.binary_fill_holes(dil)
        xholes, _ = ndimage.label(dfil & ~dil)
        for hi, hsl in enumerate(ndimage.find_objects(xholes), 1):
            core = xholes[hsl] == hi
            if core.sum() < RING_MIN:
                continue
            grown = np.zeros((h, w), bool)
            grown[hsl] = core
            grown = ndimage.binary_dilation(grown, iterations=HOLE_CLOSE) \
                & dfil & ~filled
            t = _tight(grown)
            if t is None:
                continue
            hsl = (slice(t[1], t[3]), slice(t[0], t[2]))
            hole = grown[hsl]
            own = hole & ~others_d[hsl]
            if not own.any():
                continue
            rides = False
            if grid is not None:
                g = grid[hsl][own]
                onl = np.abs(g - np.median(grid.reshape(-1, 3), axis=0)) \
                    .sum(axis=1) > GRID_LINE_MIN
                rides = onl.sum() >= 40 and float(
                    np.abs(arr[hsl][own].astype(float) - g)
                    .sum(axis=1)[onl].mean()) > 2 * TINT_MATCH
            if not rides and tex_rate >= TEX_RATE_MIN:
                rides = float(texm[hsl][own].mean()) < HOLE_TEX * tex_rate
                tr |= rides
            if rides:
                filled[hsl] |= own
        out.append((_tight(filled), filled))
        texrode.append(tr)
    # exposed as a function attribute so no caller's unpacking changes: the
    # flags line up with the returned parts, and only rebuild() reads them
    vision_elements.texrode = texrode
    return out


RULE_MAX_H = 8           # px; a taller component is a drawing, not a rule
RULE_ASPECT = 15         # w/h a rule must exceed
RULE_SPAN = 0.5          # of the object's own width a rule must cross
RULE_MIN_W = 150         # px; shorter unclaimed strokes stay background
EDGE = 50                # px; anything this close to the page edge is furniture


def _split_rules(m):
    """Split a standalone divider rule out of an object's mask.

    Page 10's divider bar sits inside the ring's box, so the claim welds them;
    a designer moves a rule and a ring separately. Only a component that is
    rule-shaped AND crosses at least half the object's width comes out: a dash
    of a dashed border is rule-shaped too, and it must stay with its drawing.
    """
    from scipy import ndimage
    labels, n = ndimage.label(m)
    if n < 2:
        return [m]
    x0, _, x1, _ = _tight(m)
    rest, out = np.zeros_like(m), []
    for i, sl in enumerate(ndimage.find_objects(labels), 1):
        h, w = sl[0].stop - sl[0].start, sl[1].stop - sl[1].start
        part = labels == i
        if h <= RULE_MAX_H and w > RULE_ASPECT * h and w >= RULE_SPAN * (x1 - x0):
            out.append(part)
        else:
            rest |= part
    return ([rest] if rest.any() else []) + out


BOLD_OVER = 1.15         # x the page's own median ink weight = a bold line
# (measured against every line of this deck read by eye: at 1.15 ten of its
# thirteen bold lines are caught and one regular line is wrongly bolded. The
# three missed are slide 7's headings, which sit at 1.11-1.13 among regular
# lines reading 1.10-1.13 -- no threshold separates those.)
BULLET_REACH = 2.5       # x line height to the left of a block to look in
BULLET_MIN = 0.15        # x line height; anything smaller is a grid line
BULLET_MAX = 0.7         # x line height; anything wider is artwork, not a
                         # bullet. Judged against the TIGHT ink box heights,
                         # which sit well under the OCR boxes: at 0.55 page 9's
                         # dots measured 0.59 of their lines and stayed baked.


def ink_box(arr, box, ink):
    """Shrink a line's box to where its glyphs really are.

    The OCR box is a pixel or two generous on every edge. That did not matter
    while the WIDTH decided the type size, but the height decides it now, and
    at a 40px cap two spare pixels size the line 5% large.

    The edge taken is where a pixel stops being closer to the ink than to the
    background, which is the 50%-coverage boundary an antialiased glyph is
    actually drawn to: the solid ink alone would sit inside the real outline.
    """
    x0, y0, x1, y1 = box
    px = arr[y0:y1, x0:x1].astype(int)
    if px.size == 0:
        return box
    m = (np.abs(px - np.array(ink)).sum(axis=2)
         < np.abs(px - ring_median(arr, box)).sum(axis=2))
    rows, cols = np.argwhere(m.any(1)).ravel(), np.argwhere(m.any(0)).ravel()
    if not len(rows) or not len(cols):
        return box
    return (x0 + cols[0], y0 + rows[0], x0 + cols[-1] + 1, y0 + rows[-1] + 1)


def ink_coverage(arr, box, ink):
    """How much ink a line carries, each pixel counted by how far it sits from
    the background.

    Sub-pixel, unlike a run length, and that is what the job needs: at 16-25pt
    every line on a page measures a 2px median stroke run whether it is bold or
    regular, so the runs cannot separate a bold heading from its own body text.
    """
    x0, y0, x1, y1 = box
    px = arr[y0:y1, x0:x1].astype(float)
    bg = ring_median(arr, box)
    denom = float(np.abs(np.array(ink, float) - bg).sum())
    if px.size == 0 or denom < 30:
        return 0.0
    return float(np.clip(np.abs(px - bg).sum(axis=2) / denom, 0, 1).sum())


_ref_cache = {}


REF_SIZE = 100           # px; draw the reference once this big and scale it


def ref_coverage(text, family, size):
    """Ink in this exact line drawn REGULAR, so what the line SAYS cannot move
    the verdict: an "m" carries more ink than an "l" in any weight.

    Drawn at one fixed size and scaled by the square of the ratio, because a
    font can only be rendered at a whole number of points: rounding a 17.4pt
    line to 17 costs 5% of its width and so 10% of its ink, which is the size
    of the bold signal itself.
    """
    path, key = font_file(family), (family, text[:40])
    if not path:
        return 0.0
    if key not in _ref_cache:
        from PIL import ImageFont, ImageDraw
        f = ImageFont.truetype(path, REF_SIZE)
        b = f.getbbox(text)
        im = Image.new("L", (b[2] - b[0] + 8, b[3] - b[1] + 8), 255)
        ImageDraw.Draw(im).text((4 - b[0], 4 - b[1]), text, font=f, fill=0)
        _ref_cache[key] = float((1.0 - np.array(im) / 255.0).sum())
    return _ref_cache[key] * (size / REF_SIZE) ** 2


def mark_bold(items):
    """Flag every line on a page drawn heavier than others of its size, and how
    many of its leading WORDS are the heavy ones.

    Weight is the line's own ink over the ink of that same line re-drawn
    regular, which cancels both the typeface and what the line says. Compared
    against lines of a similar SIZE, because the ratio still drifts with size
    (small type renders relatively heavier); a headline with nothing its own
    size on the page falls back to the rest of its own family, which is a fair
    baseline for exactly the same reason.

    These decks bold a LEAD-IN phrase and run the sentence on in regular
    ("Indexed Chunks: Each data chunk carries..."), so a per-line verdict types
    the whole first line bold. The lead-in always ends at a colon, and that is
    a far steadier signal than the words' own weights.
    """
    for i in items:
        i["bold"], i["bold_n"] = False, 0
    # a draft-stamp footer is never set bold, and its weight never joins a
    # peer pool: its box rides the panel frame and the frame's ink inflates
    # the weight, both as a verdict and as another line's baseline
    real = [i for i in items if i["ink"] > 0 and i["size"] > 0
            and not STAMP_RE.search(i["text"])]
    for i in real:
        i["weight"] = i["ink"] / max(ref_coverage(i["text"], i["font"],
                                                  i["size"]), 1.0)
    for i in real:
        peers = [p["weight"] for p in real
                 if abs(p["size"] - i["size"]) <= 0.3 * i["size"]]
        if len(peers) < 3:
            peers = [p["weight"] for p in real if p["font"] == i["font"]]
        if len(peers) < 3:
            continue
        i["bold"] = i["weight"] > BOLD_OVER * float(np.median(peers))
        if not i["bold"]:
            continue
        words = i["words"]
        i["bold_n"] = next((k for k, w in enumerate(words, 1)
                            if w["text"].endswith(":") and k < len(words)),
                           len(words))


def bullet_dots(arr, blk):
    """The drawn bullet in front of each line of a block, or None.

    A list is drawn as a coloured dot per line plus indented text, so it comes
    back as a plain paragraph with the dots stuck in the background picture.
    Found here, they become real bullets that survive editing a line.
    """
    x0 = min(i["box"][0] for i in blk)
    h = int(np.median([i["box"][3] - i["box"][1] for i in blk]))
    lo = max(x0 - int(BULLET_REACH * h), 0)
    if len(blk) < 2 or lo >= x0:
        return None
    dots = []
    for i in blk:
        y0, y1 = i["box"][1], i["box"][3]
        strip = arr[y0:y1, lo:x0].astype(int)
        if strip.size == 0:
            return None
        # the dot is small, so the strip's own median IS its background. Judge
        # it against the strongest mark in the strip, not a fixed distance:
        # these pages carry a faint drawn grid that clears any fixed threshold
        bg = np.median(strip.reshape(-1, 3), axis=0)
        d = np.abs(strip - bg).sum(axis=2)
        if d.max() < 2 * COLOR_DIST:
            return None
        mask = d > max(2 * COLOR_DIST, 0.4 * d.max())
        cols = np.argwhere(mask.any(axis=0)).ravel()
        if not len(cols):
            return None
        # several marks can share the strip (a grid tick, a sliver of the
        # text's own first letter). The bullet is the rightmost mark SHAPED
        # like one: a small round dot. A panel's border also sits left of
        # every line, but it is a full-height hairline and never passes.
        hit = None
        for run in reversed(np.split(cols,
                                     np.flatnonzero(np.diff(cols) > 2) + 1)):
            dw = run[-1] - run[0] + 1
            sub = mask[:, run[0]:run[-1] + 1]
            rows = np.argwhere(sub.any(axis=1)).ravel()
            dh = rows[-1] - rows[0] + 1
            solid = float(sub[rows[0]:rows[-1] + 1].mean())
            # a solid dot is small; a drawn CHECKBOX is nearly line-height
            # and hollow, and is a bullet too (typed as an outline square)
            if not ((BULLET_MIN * h <= dw <= BULLET_MAX * h)
                    or (dw <= 1.2 * h and solid < 0.55)):
                continue
            if (not BULLET_MIN * h <= dh <= 1.2 * h
                    or not 0.5 <= dw / dh <= 2.0):
                continue
            hit = ((lo + run[0], y0 + rows[0],
                    lo + run[-1] + 1, y0 + rows[-1] + 1),
                   strip[:, run[0]:run[-1] + 1][sub], solid)
            break
        if hit is None:
            return None
        dots.append(hit)
    colour = tuple(int(v) for v in
                   np.median(np.concatenate([d[1] for d in dots]), axis=0))
    char = "□" if np.median([d[2] for d in dots]) < 0.55 else "•"
    return [d[0] for d in dots], colour, char


ELEM_DILATE = 3          # px of slack that joins the parts of one drawing
# these decks are drawn on a faint blueprint grid with tick marks in the
# margins. At COLOR_DIST that furniture is "ink", it touches every drawing on
# the page, and the whole slide comes back as ONE picture element you cannot
# pull apart -- which is the flattening the tool exists to undo. Judged at
# twice the distance the grid drops out and the drawings separate.
ELEM_DIST = 2 * COLOR_DIST
SPARSE = 0.05            # ink / bbox area below which a component is not one
                         # drawing but several linked by a hairline
NECK_INK = 4             # ink px a cut may cross: a wire or two, no more
NECK_RUN = 15            # px; a shorter neck is the gap between two strokes of
                         # one shape, not the join between two drawings
NECK_EDGE = 0.1          # of the cut line; ink in BOTH end bands means the cut
                         # is crossing the shape's own outline


def _tight(mask):
    """Bounding box of the ink in mask, or None."""
    rows, cols = np.flatnonzero(mask.any(1)), np.flatnonzero(mask.any(0))
    if not len(rows):
        return None
    return (int(cols[0]), int(rows[0]), int(cols[-1]) + 1, int(rows[-1]) + 1)


def _necks(sub, axis):
    """Where a drawing is crossed by only a wire or two: axis 0 for columns.

    A hollow shape is the trap. Every column through the middle of an outlined
    panel crosses just its top and bottom rule, which is as thin as any wire,
    and cutting there saws the panel in half. Such a column is told apart by
    WHERE its ink sits: a join sits somewhere in the middle of the cut line, a
    shape's own outline sits at both ends of it.
    """
    profile = sub.sum(axis=axis)
    low = np.flatnonzero(profile <= NECK_INK)
    if not len(low):
        return []
    edge = max(int(NECK_EDGE * sub.shape[axis]), 3)
    cuts = []
    for r in np.split(low, np.flatnonzero(np.diff(low) > 1) + 1):
        if len(r) < NECK_RUN or r[0] == 0 or r[-1] == len(profile) - 1:
            continue
        band = sub[:, r[0]:r[-1] + 1] if axis == 0 else sub[r[0]:r[-1] + 1, :]
        along = band.any(axis=1 - axis)
        if along[:edge].any() and along[-edge:].any():
            continue
        cuts.append(int(r.mean()))
    return cuts


def cut_necks(mask, box):
    """Cut a drawing where only a couple of thin wires run through it.

    Connectivity alone makes slide 2's whole schematic one picture: the dial,
    the maze, the cross and the verify box are all real wired together, so
    nothing can be moved on its own. But a JOIN looks different from a shape:
    a column through the wires between the maze and the cross carries a dozen
    ink pixels, a column through either of them carries hundreds. Cut at the
    thin columns, then at the thin rows of each piece.

    A shape with a genuine waist (page 7's funnel and its stem) is not cut,
    because every column through the waist still crosses the stem's own two
    sides plus the bowl above it, which is nowhere near this thin.
    """
    x0, y0, x1, y1 = box
    sub = mask[y0:y1, x0:x1]
    for axis in (0, 1):                        # thin columns first, then rows
        cuts = _necks(sub, axis)
        if not cuts:
            continue
        pieces = []
        for a, b in zip([0] + cuts, cuts + [sub.shape[1 - axis]]):
            part = np.zeros_like(mask)
            if axis == 0:
                part[y0:y1, x0 + a:x0 + b] = sub[:, a:b]
            else:
                part[y0 + a:y0 + b, x0:x1] = sub[a:b, :]
            bb = _tight(part)
            if bb and (bb[2] - bb[0]) * (bb[3] - bb[1]) >= ELEM_MIN_AREA:
                pieces += cut_necks(part, bb)
        if pieces:
            return pieces
    return [box]


def _ink_mask(arr, keep_out, dist=ELEM_DIST):
    """Ink on the page, with every keep_out text box blanked out."""
    bg = np.median(arr.reshape(-1, 3), axis=0)
    ink = np.abs(arr.astype(int) - bg).sum(axis=2) > dist
    for x0, y0, x1, y1 in keep_out:
        ink[max(y0 - PAD, 0):y1 + PAD, max(x0 - PAD, 0):x1 + PAD] = False
    return ink


def find_elements(arr, keep_out=()):
    """Individual drawings on the page: connected regions of strong ink.

    The pixel-rule fallback for when no vision model is reachable (the primary
    path is object_boxes + vision_elements). A component whose box is mostly
    empty is not one drawing, it is two joined by a hairline, so it is re-cut
    with less slack, down to touching pixels; what survives is then cut at its
    thin joins. Anything under ELEM_MIN_AREA on the way stays in the background
    image, and so does anything inside keep_out -- the boxes of lines left
    baked as text, which are words, not drawings, and were coming back as a
    scatter of picture shards.
    """
    from scipy import ndimage
    ink = _ink_mask(arr, keep_out)

    def split(mask, dilate):
        grown = ndimage.binary_dilation(mask, iterations=dilate) if dilate else mask
        labels, _ = ndimage.label(grown)
        boxes = []
        for i, (ys, xs) in enumerate(ndimage.find_objects(labels), 1):
            area = (xs.stop - xs.start) * (ys.stop - ys.start)
            if area < ELEM_MIN_AREA:
                continue
            part = mask & (labels == i)
            if dilate and part.sum() < SPARSE * area:
                boxes += split(part, dilate - 1)
            else:
                boxes += cut_necks(part, _tight(part))
        return boxes

    h, w = ink.shape
    return [(max(x0 - ELEM_DILATE, 0), max(y0 - ELEM_DILATE, 0),
             min(x1 + ELEM_DILATE, w), min(y1 + ELEM_DILATE, h))
            for x0, y0, x1, y1 in split(ink, ELEM_DILATE)]



SAME_SIZE = 1.5          # pt; OCR height noise between lines of one paragraph
SAME_SIZE_REL = 0.1      # ...which grows with the type: 1.5pt is tight at 20pt
# (0.1 and no higher: at 0.12 a 17.5pt column heading absorbs its 15.5pt body
# and gets re-typed at the body's size. An inline BOLD lead-in reads 2pt bigger
# than its own paragraph and splits off instead, which costs a box, not a look.)
SAME_COLOR = 24          # per-channel; sampling noise between lines
SAME_LEFT = 6            # px; how far left edges may differ and still be a block
MAX_GAP = 1.9            # line pitch / line height above which the run breaks


def aligned(a, b):
    """Same block? Lines share an edge: flush left, or centred on each other.

    The diagram callouts ("ERROR / BOUNDARY", "TARGET OUTPUT / (SINGLE RUN)")
    are centred, so a left-edge-only test splits every one of them.
    """
    if abs(a[0] - b[0]) <= SAME_LEFT:
        return True
    return abs((a[0] + a[2]) - (b[0] + b[2])) <= 2 * SAME_LEFT


def group_lines(items):
    """Consecutive lines that share font, size, colour and left edge become one
    text box (one paragraph per line), which is how the original deck sets them.
    Anything that differs in any of those gets its own box."""
    rest = sorted(items, key=lambda i: (i["box"][1], i["box"][0]))
    blocks = []
    while rest:
        blk, cur, rest = [rest[0]], rest[0], rest[1:]
        while True:
            h = cur["box"][3] - cur["box"][1]
            nxt = next((i for i in rest
                        if i["font"] == cur["font"]
                        and abs(i["size"] - cur["size"])
                            <= max(SAME_SIZE, SAME_SIZE_REL * cur["size"])
                        and aligned(i["box"], cur["box"])
                        and max(abs(a - b) for a, b in
                                zip(i["color"], cur["color"])) <= SAME_COLOR
                        and 0 < i["box"][1] - cur["box"][1] <= MAX_GAP * h), None)
            if nxt is None:
                break
            blk.append(nxt)
            rest.remove(nxt)
            cur = nxt
        blocks.append(blk)
    return blocks


SAME_PITCH = 0.25        # how far a gap may sit from the paragraph's own pitch
JOIN_SIZE = 0.2          # relative size gap two halves of one paragraph may show


def join_blocks(blocks):
    """Re-join blocks that the size test split out of one paragraph.

    Fitting proportional text in a mono face (or the reverse) makes the fitted
    size wobble line to line, so a paragraph breaks into pieces that then get
    re-typed at DIFFERENT sizes, which is the ragged look. Rhythm is the honest
    signal here: a paragraph keeps one line pitch, and a heading always sits
    further from its body than the body's own lines sit from each other.
    """
    def pitch(blk):
        return blk[1]["box"][1] - blk[0]["box"][1] if len(blk) > 1 else None

    def size(blk):
        return float(np.median([i["size"] for i in blk]))

    out = sorted(blocks, key=lambda b: (b[0]["box"][1], b[0]["box"][0]))
    n = 0
    while n < len(out) - 1:
        a, b = out[n], out[n + 1]
        gap = b[0]["box"][1] - a[-1]["box"][1]
        p = pitch(a) or pitch(b) or 1.15 * (a[-1]["box"][3] - a[-1]["box"][1])
        if (aligned(a[-1]["box"], b[0]["box"])
                and abs(gap - p) <= SAME_PITCH * p
                and abs(size(a) - size(b)) <= JOIN_SIZE * max(size(a), size(b))
                and max(abs(x - y) for x, y in
                        zip(a[-1]["color"], b[0]["color"])) <= SAME_COLOR):
            out[n] = a + b
            del out[n + 1]
        else:
            n += 1
    return out


def block_align(blk):
    """Which edge the original lines were set to.

    Guessing wrong is not cosmetic: type a longer line into a left-aligned box
    that was really centred and the whole paragraph walks sideways.
    """
    if len(blk) < 2:
        return PP_ALIGN.LEFT
    lefts = [i["box"][0] for i in blk]
    rights = [i["box"][2] for i in blk]
    if max(lefts) - min(lefts) <= SAME_LEFT:
        return PP_ALIGN.LEFT
    if max(rights) - min(rights) <= SAME_LEFT:
        return PP_ALIGN.RIGHT
    return PP_ALIGN.CENTER


_A = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
BULLET_XML = (f'<a:buClr {_A}><a:srgbClr val="{{colour}}"/></a:buClr>',
              f'<a:buFont {_A} typeface="Arial"/>',
              f'<a:buChar {_A} char="{{char}}"/>')


def runs(item):
    """(text, bold) pairs for a line: a bold lead-in, then the rest."""
    n = item.get("bold_n", 0)
    words = [w["text"] for w in item.get("words") or []]
    # the line text can differ from its words (fix_roman rewrites it), and the
    # text is what belongs on the slide
    if " ".join(words) != item["text"] or not 0 < n < len(words):
        return [(item["text"], bool(n) or item["bold"])]
    return [(" ".join(words[:n]), True), (" " + " ".join(words[n:]), False)]


def merge_panel_blocks(blocks, bullets, eboxes):
    """Body text sitting inside one PANEL element merges into one textbox.

    The smallest containing element owns the blocks. Same-size stacked blocks
    (a panel's body paragraphs, a two-line footer) merge; a heading set
    clearly bigger than its body stays a separate textbox, the way a table's
    header cell and body cell are separate. Merged items keep their OWN
    sizes (`mixed`) instead of the block median.
    """
    def bbox(blk):
        return (min(i["box"][0] for i in blk), min(i["box"][1] for i in blk),
                max(i["box"][2] for i in blk), max(i["box"][3] for i in blk))
    owners, out = {}, []
    for blk, bu in zip(blocks, bullets):
        b = bbox(blk)
        cands = [e for e in eboxes
                 if e[0] - PAD <= b[0] and e[1] - PAD <= b[1]
                 and e[2] + PAD >= b[2] and e[3] + PAD >= b[3]]
        if bu or not cands:
            out.append((blk, bu))
            continue
        owners.setdefault(
            min(cands, key=lambda e: (e[2] - e[0]) * (e[3] - e[1])),
            []).append(blk)
    def size(blk):
        return float(np.median([i["size"] for i in blk]))

    for blks in owners.values():
        blks.sort(key=lambda b: min(i["box"][1] for i in b))
        # merge only a STACKED, x-overlapping chain of SAME-SIZE, SAME-FONT
        # blocks: a panel's body paragraphs. Its heading is set clearly
        # bigger OR in a different face and stays its own textbox (page 8's
        # panel headings are mono over a proportional body at the same
        # size), and two labels sitting side by side inside one diagram
        # (SOURCE and CITE under page 3's cube) are never one textbox.
        chains = [[blks[0]]]
        for a, b in zip(blks, blks[1:]):
            if (bbox(b)[1] >= bbox(a)[3] - 4
                    and min(bbox(a)[2], bbox(b)[2]) > max(bbox(a)[0], bbox(b)[0])
                    and a[0]["font"] == b[0]["font"]
                    # a white-on-red row and a navy-on-white row stacked in
                    # one table column are NOT one paragraph: a merged box
                    # takes the median colour and one of them vanishes
                    and max(abs(x - y) for x, y in
                            zip(a[-1]["color"], b[0]["color"])) <= SAME_COLOR
                    and abs(size(a) - size(b))
                    <= max(SAME_SIZE, SAME_SIZE_REL * max(size(a), size(b)))):
                chains[-1].append(b)
            else:
                chains.append([b])
        for run in chains:
            if len(run) > 1:
                items = sorted((i for b in run for i in b),
                               key=lambda i: i["box"][1])
                for i in items:
                    i["mixed"] = True
                out.append((items, None))
            else:
                out.append((run[0], None))
    return [b for b, _ in out], [u for _, u in out]


def emit_block(slide, blk, sx, sy, page_w, bullet=None, centred=True):
    x0 = min(i["box"][0] for i in blk)
    x1 = max(i["box"][2] for i in blk)
    y0 = min(i["box"][1] for i in blk)
    y1 = max(i["box"][3] for i in blk)
    align = block_align(blk)
    # a lone line whose ink is centred on the page was SET centred (page 10's
    # title): typed left-aligned, a narrower substitute font walks it sideways
    if centred and len(blk) == 1 and abs((x0 + x1) / 2 * sx - page_w / 2) <= 6:
        align = PP_ALIGN.CENTER
    # one size and colour for the whole block: the per-line jitter is OCR noise
    size = float(np.median([i["size"] for i in blk]))
    col = tuple(int(v) for v in np.median([i["color"] for i in blk], axis=0))

    # hug the text: a box padded by a whole em sprawls across the artwork next
    # to it and is a nuisance to click past
    pad = 0.35 * size
    marL = 0
    if bullet:
        dot_x = min(d[0] for d in bullet[0])
        marL = int(round((x0 - dot_x) * sx * 12700))    # pt -> EMU
        x0 = dot_x
    # the box carries no inset, so its edge IS where the text starts: padding
    # the box out to the left moves the whole paragraph left by that much
    w = (x1 - x0) * sx + 2 * pad
    if align == PP_ALIGN.RIGHT:
        left = x1 * sx - w
    elif align == PP_ALIGN.CENTER:
        left = (x0 + x1) / 2 * sx - w / 2
    else:
        left = x0 * sx
    left = max(0.0, min(left, page_w - w))
    # line 0 is given the same line height as line 1 (see the loop below), which
    # lifts it by however much that height is under the font's natural one
    gap0 = max((blk[1]["box"][1] - blk[0]["box"][1]) * sy, size * 1.02) \
        if len(blk) > 1 else None
    path = font_file(blk[0]["font"])
    if not path:
        drop = size * 0.18
    elif gap0 is None:
        drop = size * (ink_top(path, blk[0]["text"]) - ORIGIN_ABOVE_TOP)
    else:
        drop = GAP_TO_INK * gap0 + (ink_top(path, blk[0]["text"]) - INK_BASE) * size
    tb = slide.shapes.add_textbox(Pt(left), Pt(y0 * sy - drop),
                                  Pt(w), Pt((y1 - y0) * sy + size * 0.4))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for n, item in enumerate(blk):
        p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        p.alignment = align
        if n:
            # each line keeps ITS OWN distance from the line above: one pitch
            # for the whole box stretches a paragraph that sits under a heading
            # out to the heading's gap. Tighter than the glyphs would overlap.
            gap = (item["box"][1] - blk[n - 1]["box"][1]) * sy
            p.line_spacing = Pt(max(gap, size * 1.02))
            if n == 1 and path:
                # ...and so does the FIRST line. A paragraph with no lnSpc keeps
                # the font's natural line height, and the step from that height
                # to the next paragraph's spaced one eats ~2px out of the first
                # gap. Giving line 0 the same spacing removes the step without
                # moving line 0 itself -- but only where `drop` compensated for
                # it. The no-font fallback leaves drop at size*0.18 (the natural
                # height), so spacing line 0 there would shift it uncompensated.
                tf.paragraphs[0].line_spacing = p.line_spacing
        if bullet:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(marL))
            pPr.set("indent", str(-marL))
            for frag in BULLET_XML:
                pPr.append(parse_xml(frag.format(
                    colour="%02X%02X%02X" % bullet[1],
                    char=bullet[2] if len(bullet) > 2 else "•")))
        for text, bold in runs(item):
            r = p.add_run()
            r.text = text
            r.font.name = item["font"]
            # a:sz is in 1/100 pt, so half a point is free: rounding to whole
            # points threw away up to 3% of the size the fit had measured
            r.font.size = Pt(round(item["size"] if item.get("mixed")
                                   else size, 1))
            r.font.bold = bold
            r.font.color.rgb = RGBColor(*col)
            if item["spacing"]:
                # python-pptx has no tracking API; a:rPr/@spc is in 1/100 pt
                r.font._rPr.set("spc", str(int(round(item["spacing"] * 100))))
    return tb


def make_item(arr, l, col, font, font_map, sx, sy):
    """A typed line: fit and place against the real glyph extent, but keep the
    whole OCR box for the erase (generous is what an erase wants)."""
    x0, y0, x1, y1 = tight = ink_box(arr, l["box"], col)
    fixed = next((f for pat, f in font_map.items()
                  if re.search(pat, l["text"], re.I)), None)
    size, spacing = fit_size(l["text"], fixed or font[0], x1 - x0, y1 - y0)
    return {"text": l["text"], "box": tight, "erase": l["box"], "color": col,
            "font": fixed or font[0], "fixed": bool(fixed),
            "size": size * sy, "spacing": spacing * sx, "sure": l.get("sure", False),
            "ink": ink_coverage(arr, tight, col), "words": l["words"]}


def unrotate_box(box, transpose, w, h):
    """Map a box on the turned page back onto the upright w x h page."""
    x0, y0, x1, y1 = box
    if transpose == Image.ROTATE_90:             # the turn was 90 deg CCW
        return (w - y1, x0, w - y0, x1)
    return (y0, h - x1, y1, h - x0)              # ROTATE_270


ROT_LETTERS = 4          # letters a turned line needs, vowel included: artwork
                         # read sideways OCRs as short junk, real labels as words


def rotated_items(img, font, font_map, sx, sy, taken):
    """Vertical text (slide 8's axis label), read by turning the page sideways
    and running the whole OCR pipeline again.

    A turned line is kept only if it reads like a word and lands where no
    upright line already is, which filters the junk a sideways pass makes of
    the artwork. Sizes are scaled with the axes swapped, because the turned
    page's vertical is the real page's horizontal.
    """
    out = []
    for transpose in (Image.ROTATE_90, Image.ROTATE_270):
        rimg = img.transpose(transpose)
        rarr = np.asarray(rimg)
        for l in ocr_lines(rimg):
            t = l["text"]
            if (l["drop"] or len(re.findall(r"[A-Za-z]", t)) < ROT_LETTERS
                    or not re.search(r"[AEIOUaeiou]", t)):
                continue
            obox = unrotate_box(l["box"], transpose, img.width, img.height)
            if any(obox[0] < b[2] and b[0] < obox[2]
                   and obox[1] < b[3] and b[1] < obox[3]
                   for b in taken + [r["obox"] for r in out]):
                continue
            item = make_item(rarr, l, text_color(rarr, l["box"]),
                             font, font_map, sy, sx)
            if item["size"] < MIN_TYPE_PT:
                continue
            item["bold"], item["bold_n"] = False, 0
            out.append({"item": item, "obox": obox, "transpose": transpose})
    return out


def add_picture(slide, pil_img, x, y, w, h):
    buf = io.BytesIO()
    pil_img.save(buf, format="PNG")
    buf.seek(0)
    slide.shapes.add_picture(buf, Pt(x), Pt(y), Pt(w), Pt(h))


def pill_mask(arr, box):
    """The watermark pill's true extent on THIS page.

    The generator composites its pill TRANSLUCENT over the corner, so its
    pixels shift with every page's ground and no cross-page agreement exists
    to find them (measured: zero identical pixels across this deck's ten
    pages). What does hold per page: the slab and its logo sit far from the
    page's background while staying GREY, where the deck's own ink is
    coloured. Grey components overlapping the OCR words are the pill - slab,
    logo, rounded corners. None when nothing there matches."""
    from scipy import ndimage
    h, w = arr.shape[:2]
    x0, y0 = max(box[0] - PILL_PAD, 0), max(box[1] - PILL_PAD, 0)
    x1, y1 = min(box[2] + PILL_PAD, w), min(box[3] + PILL_PAD, h)
    if x1 <= x0 or y1 <= y0:
        return None
    sub = arr[y0:y1, x0:x1].astype(int)
    bg = np.median(arr.reshape(-1, 3), axis=0)
    dist = np.abs(sub - bg).sum(axis=2)
    grey = (np.abs(sub[..., 0] - sub[..., 1])
            + np.abs(sub[..., 1] - sub[..., 2])) < PILL_GREY
    lab, _ = ndimage.label((dist > COLOR_DIST) & grey)
    hit = np.unique(lab[box[1] - y0:box[3] - y0, box[0] - x0:box[2] - x0])
    hit = hit[hit != 0]
    if not len(hit):
        return None
    # the wordmark is the anchor; the slab it sits on only DIMS the ground
    # (a light-on-light corner drops ~20-40 per pixel, well under
    # COLOR_DIST) - reach it by connectivity through dim grey pixels, or
    # the slab survives the patch as a grey box
    core = np.isin(lab, hit)
    win = ndimage.binary_propagation(core, mask=core | ((dist > DIM_DIST)
                                                       & grey))
    m = np.zeros((h, w), bool)
    m[y0:y1, x0:x1] = win
    return ndimage.binary_fill_holes(ndimage.binary_dilation(m, iterations=2))


def rebuild(pdf_path, out_path, font, font_map, cache=None, dump=None):
    font = [f.strip() for f in font.split(",")] if isinstance(font, str) else font
    doc = fitz.open(pdf_path)
    prs = Presentation()
    prs.slide_width = Pt(doc[0].rect.width)
    prs.slide_height = Pt(doc[0].rect.height)
    blank = prs.slide_layouts[6]
    report = []
    pending = []            # slides built but not yet typed (see mark_bold)
    watermark = {}          # image size -> the box the watermark sat in

    imgs = []
    for page in doc:
        xref = page.get_images()[0][0]
        pix = fitz.Pixmap(doc, xref)
        if pix.n > 3:
            pix = fitz.Pixmap(fitz.csRGB, pix)
        imgs.append(Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB"))

    for page, img in zip(doc, imgs):
        sx = page.rect.width / img.width       # px -> pt
        sy = page.rect.height / img.height
        arr = np.array(img)
        arr_pre = arr.copy()   # pristine page, for re-synthesizing the grid
        # the object boxes are needed FIRST: the grid synth must not absorb
        # the artwork's own rules into the "empty page" (see page_grid).
        # Grown, because the model's boxes sit a few pixels off the ink and
        # a box a hair short would leave artwork rules inside the medians.
        objects = object_boxes(img, cache, page.number)

        raw = ocr_lines(img)
        wmb = next((l["box"] for l in raw if l["drop"]),
                   watermark.get(img.size, (None,))[0])
        pill = pill_mask(arr, wmb) if wmb else None
        all_lines = lines = fix_small_text(img, raw, cache, page.number, pill)
        all_lines = lines = proofread_text(img, lines, cache, page.number)
        all_lines = lines = weld_gap_tokens(img, lines, font)
        all_boxes = [l["box"] for l in lines]
        # read BEFORE anything is erased: the turned passes see the same pixels
        rots = rotated_items(img, font, font_map, sx, sy, all_boxes)
        # the synth is built AFTER the OCR so text rows can be claimed: text
        # ink otherwise testifies for any grid line crossing it (page 7's
        # mesh rows read every word they crossed as line evidence, baking
        # word-shaped ghosts into the "empty page", and the patches under the
        # re-typed text then had no honest mesh to paint - flat plateaus)
        claimed0 = np.zeros(arr.shape[:2], bool)
        for x0, y0, x1, y1 in all_boxes + [r["obox"] for r in rots]:
            claimed0[max(y0 - 2, 0):y1 + 2, max(x0 - 2, 0):x1 + 2] = True
        grid = page_grid(arr, [(b[0] - OBJ_GROW, b[1] - OBJ_GROW,
                                b[2] + OBJ_GROW, b[3] + OBJ_GROW)
                               for o in objects for b in [o["box"]]]
                         if objects else (), claimed=claimed0)
        inks = [text_color(arr, l["box"]) for l in lines]
        # text the illustrator drew around stays baked in the background image
        # - except a vision-completed footer ("... // SCALE: 1:1"): its panel
        # is a plain box, not drawing, and baking it loses the completion
        drawn = diagram_lines(img, lines, cache, page.number) \
            - {n for n, l in enumerate(lines)
               if l.get("sure") and "//" in l["text"]
               or STAMP_RE.search(l["text"])}
        # --strip-stamps only: the draft stamps erased, never re-typed
        stamps = [(l["box"], c) for l, c in zip(lines, inks)
                  if not l["drop"] and STAMP_RE.search(l["text"])] \
            if STRIP_STAMPS else []
        pairs = [(l, c) for n, (l, c) in enumerate(zip(lines, inks))
                 if not l["drop"] and n not in drawn
                 and not (STRIP_STAMPS and STAMP_RE.search(l["text"]))]
        dropped = sum(l["drop"] for l in lines) + len(stamps)
        baked = len(lines) - len(pairs) - dropped
        lines = [l for l, _ in pairs]
        colors = [c for _, c in pairs]
        boxes = [l["box"] for l in lines]
        # erase what gets re-typed, plus the watermark, which is simply gone.
        # It sits in the same corner of every page, so a page the OCR misread
        # it on still gets it erased and the deck stays consistent.
        wm = [(l["box"], c) for l, c in zip(all_lines, inks) if l["drop"]]
        if wm:
            watermark[img.size] = wm[0]
        elif img.size in watermark:
            wm = [watermark[img.size]]
            dropped += 1

        # weight ("ink") is read off the ORIGINAL pixels, before the text it
        # belongs to is erased
        items = [make_item(arr, l, col, font, font_map, sx, sy)
                 for l, col in zip(lines, colors)]
        # type this small is what the OCR gets wrong ("SCALE: 1:1" comes back
        # as "smug. 1 •1"), and it is the deck's footers and tick labels, which
        # nobody edits. Left baked, they still READ correctly.
        # a sure line is exempt from the bake floor (it is a footer the vision
        # pass read correctly), but one fitting under 6pt is not text at all:
        # page 1's "3.0" dimension label came through here at 4pt
        small = [i for i in items
                 if i["size"] < MIN_TYPE_PT and not (i["sure"] and i["size"] >= 6)]
        items = [i for i in items
                 if i["size"] >= MIN_TYPE_PT or (i["sure"] and i["size"] >= 6)]
        baked += len(small)
        boxes = [i["erase"] for i in items]
        colors = [i["color"] for i in items]
        blocks = join_blocks(group_lines(items))
        # the family is decided per PARAGRAPH, off all its words at once: one
        # short line is not enough to tell a mono from a proportional face, and
        # a paragraph set in two families is worse than one wrong guess
        for blk in blocks:
            if len(font) < 2 or any(i["fixed"] for i in blk):
                continue
            fam = pick_font([w for i in blk for w in i["words"]], font)
            for i in blk:
                x0, y0, x1, y1 = i["box"]
                size, spacing = fit_size(i["text"], fam, x1 - x0, y1 - y0)
                i["font"], i["size"], i["spacing"] = fam, size * sy, spacing * sx
        mark_bold(items)
        bullets = [bullet_dots(arr, b) for b in blocks]

        for b, c in list(zip(boxes, colors)) + stamps:
            patch(arr, b, others=all_boxes, ink=c, grid=grid)
        # the pill's true extent (padding, corners, logo) runs beyond the OCR
        # words: patch exactly its pixels, or its grey edges survive as
        # smudges and the rectangle guts whatever baked text it overlapped.
        # No pill found -> the old rectangle.
        wmboxes = []
        for b, c in wm:
            pm = pill if pill is not None else pill_mask(arr, b)
            if pm is None:
                patch(arr, b, others=all_boxes, ink=c, grid=grid)
                wmboxes.append(b)
            else:
                patch(arr, _tight(pm), pad=PAD, others=all_boxes, mask=pm,
                      grid=grid)
                wmboxes.append(_tight(pm))
        for bu in bullets:
            if bu:
                for d in bu[0]:
                    # ink given, or the erased dot itself scores the fill and
                    # a tile that copies the dot above wins by reproducing it
                    patch(arr, d, others=all_boxes, ink=bu[1], grid=grid)
        for r in rots:
            patch(arr, r["obox"], others=all_boxes, ink=r["item"]["color"],
                  grid=grid)

        # only ERASED text is blanked out of the element search: text left
        # baked inside the artwork is part of its drawing and must ride with
        # the element, or moving the element leaves its own labels behind.
        # The fallback still blanks every line, because without object boxes
        # a baked footer comes back as a scatter of picture shards.
        keep = [i["erase"] for i in items] + [b for b, _ in stamps] \
            + wmboxes + [r["obox"] for r in rots]
        bakedb = [all_lines[n]["box"] for n in drawn if n < len(all_lines)]
        if objects is None:
            found = find_elements(arr, all_boxes + [r["obox"] for r in rots])
            # drop boxes fully inside a bigger one (already part of that crop)
            found.sort(key=lambda b: (b[2] - b[0]) * (b[3] - b[1]), reverse=True)
            kept = []
            for b in found:
                if not any(b[0] >= k[0] and b[1] >= k[1]
                           and b[2] <= k[2] and b[3] <= k[3] for k in kept):
                    kept.append(b)
            elems = [(b, None) for b in kept]
        else:
            elems = vision_elements(arr, keep, objects, bakedb, grid)
        elems = [[b, m] for b, m in elems]
        eboxes = [b for b, _ in elems]
        # re-synthesize the empty page now that the elements are known: ink
        # an element carries away must not testify for a grid line, or the
        # patch under that element repaints the element's own stroke into
        # the background (page 4's leader row) and dragging the element
        # reveals the copy
        if grid is not None and any(m is not None for _, m in elems):
            from scipy import ndimage as ndi0
            allm = np.zeros(arr.shape[:2], bool)
            for _, m in elems:
                if m is not None:
                    allm |= m
            allm = ndi0.binary_dilation(allm, iterations=FRINGE)
            for x0, y0, x1, y1 in keep:
                allm[max(y0 - PAD, 0):y1 + PAD,
                     max(x0 - PAD, 0):x1 + PAD] = True
            grid = page_grid(arr_pre, [(b[0] - OBJ_GROW, b[1] - OBJ_GROW,
                                        b[2] + OBJ_GROW, b[3] + OBJ_GROW)
                                       for o in (objects or [])
                                       for b in [o["box"]]], claimed=allm)
        # crops are built AFTER the patches, from this pristine copy: the
        # patched page then says which crop pixels are translucent wash
        arr0 = arr.copy()
        # a masked patch is padded so the ring of true background around the
        # element scores the fill candidates; only mask pixels are written, so
        # the pad erases nothing. An unmasked (fallback) patch fills its whole
        # region and must not pad. The ERASE runs FRINGE wider than the claim:
        # the antialiased skirt outside the mask sits under every ink
        # threshold, and left in the background it redraws the element as a
        # ghost outline the moment the element moves (page 4's funnel rim and
        # its OFFICIAL label). The crop keeps the tight mask.
        from scipy import ndimage as ndi_
        texflags = (getattr(vision_elements, "texrode", [])
                    if objects is not None else [])
        painted = np.zeros(arr.shape[:2], bool)
        for n, (b, m) in enumerate(elems):
            if m is None:
                patch(arr, b, pad=0, grid=grid)
                painted[b[1]:b[3], b[0]:b[2]] = True
                continue
            dm = ndi_.binary_dilation(m, iterations=FRINGE)
            if n < len(texflags) and texflags[n]:
                # this element's interior HIDES the page's texture (page
                # 10's disc): the ground under it is unknowable and no
                # ring-scored candidate can say otherwise, so it fills as a
                # smooth membrane from its surroundings instead of carrying
                # a comb of the texture's medians through the hole
                X0, Y0 = max(b[0] - MASK_PAD, 0), max(b[1] - MASK_PAD, 0)
                X1 = min(b[2] + MASK_PAD, arr.shape[1])
                Y1 = min(b[3] + MASK_PAD, arr.shape[0])
                f = _diffuse(arr[Y0:Y1, X0:X1], dm[Y0:Y1, X0:X1])
                if f is not None:
                    mm = dm[Y0:Y1, X0:X1]
                    arr[Y0:Y1, X0:X1][mm] = np.clip(
                        np.round(f), 0, 255).astype(arr.dtype)[mm]
                else:
                    # the membrane could not be built (the dilated mask fills
                    # the whole window, e.g. an element flush to the page edge):
                    # fall back to the masked patch so the element's own ink is
                    # still erased from the bg instead of ghosting when it moves
                    patch(arr, b, pad=MASK_PAD, others=keep + eboxes,
                          mask=dm, grid=grid)
            else:
                patch(arr, b, pad=MASK_PAD, others=keep + eboxes,
                      mask=dm, grid=grid)
            painted |= dm
        # dim leftovers MOSTLY inside an element's extent are that element's
        # own artwork the claims missed (a tint fill bleeding into the page
        # grid, the hatched band whose box was drawn short, a pale shadow):
        # erase them from the page AND donate them to the element, so the
        # fill moves with its drawing instead of vanishing. The page's own
        # grid always runs beyond any one element's extent, and furniture
        # hugs the edge band, so neither can qualify.
        if elems:
            from scipy import ndimage
            # measured against the synthesized empty page where one exists:
            # on a ruled page the grid itself sits past DIM_DIST of the page
            # median, chaining every leftover into one page-wide component
            # that can never sit inside an element's extent
            ref = grid if grid is not None \
                else np.median(arr.reshape(-1, 3), axis=0)
            chartb = [(b[0] - OBJ_GROW, b[1] - OBJ_GROW,
                       b[2] + OBJ_GROW, b[3] + OBJ_GROW)
                      for o in (objects or [])
                      if CHART_RE.search(o["name"]) for b in [o["box"]]]

            def chart_owned(b):
                area = (b[2] - b[0]) * (b[3] - b[1])
                return any(
                    (min(b[2], c[2]) - max(b[0], c[0]))
                    * (min(b[3], c[3]) - max(b[1], c[1])) >= 0.8 * area
                    for c in chartb
                    if min(b[2], c[2]) > max(b[0], c[0])
                    and min(b[3], c[3]) > max(b[1], c[1]))
            res = np.abs(arr.astype(int) - ref).sum(axis=2) > DIM_DIST
            # a pixel an element patch just painted is the tool's own best
            # fill, never a leftover: judging it against the grid reference
            # re-carved page 10's smooth disc fill into comb-line dashes
            res &= ~painted
            for x0, y0, x1, y1 in keep:
                res[max(y0 - PAD, 0):y1 + PAD, max(x0 - PAD, 0):x1 + PAD] = False
            res[:EDGE] = res[-EDGE:] = False
            res[:, :EDGE] = res[:, -EDGE:] = False
            # the page's own repeating texture (page 10's mesh): residue in a
            # colour the page repeats AT SCALE outside the element boxes is
            # the pattern showing through the claims, never the element's own
            # drawing. Donated, it dragged mesh lines into the ring's crop
            # and left flat holes in the background where they were erased.
            pagebg = np.median(arr0.reshape(-1, 3), axis=0)
            outtex = np.abs(arr0.astype(int) - pagebg).sum(axis=2) > PALE_DIST
            for x0, y0, x1, y1 in eboxes + keep:
                outtex[max(y0 - PAD, 0):y1 + PAD,
                       max(x0 - PAD, 0):x1 + PAD] = False
            texcols = []
            tlab, _ = ndimage.label(outtex)
            for ti, tsl in enumerate(ndimage.find_objects(tlab), 1):
                tp = outtex[tsl] & (tlab[tsl] == ti)
                if tp.sum() >= TEX_MIN:
                    tpx = arr0[tsl][tp].reshape(-1, 3).astype(int)
                    td = np.abs(tpx - pagebg).sum(axis=1)
                    texcols.append(np.median(
                        tpx[td >= np.quantile(td, 0.75)], axis=0))
            # the whole DILATED component is one patch region: its true ring
            # is real background. Patched piecewise (a thin line fragments
            # into specks), each speck's ring is the line's own colour and
            # the fill reprints exactly what it should erase.
            lab, _ = ndimage.label(ndimage.binary_dilation(res, iterations=2))
            for i, sl in enumerate(ndimage.find_objects(lab), 1):
                part = lab[sl] == i
                raw_px = int((res[sl] & part).sum())
                if raw_px < 6:
                    continue
                if raw_px < 30:
                    # under the noise floor, but a clearly COLOURED speck is
                    # the element's paint, not antialiasing (page 5's mint
                    # dot sat at 16px through two rounds): it may still go
                    ys0, xs0 = np.nonzero(res[sl] & part)
                    px0 = arr[ys0 + sl[0].start, xs0 + sl[1].start].astype(int)
                    cm0 = np.median(px0, axis=0)
                    if (abs(cm0[0] - cm0[1]) + abs(cm0[1] - cm0[2])
                            <= GREY_SPREAD):
                        continue
                bx = (sl[1].start, sl[0].start, sl[1].stop, sl[0].stop)
                ys, xs = np.nonzero(res[sl] & part)
                ys, xs = ys + sl[0].start, xs + sl[1].start
                # ties break toward the SMALLER element, like every other
                # ownership decision in this file
                frac, _, k = max(
                    (float(((xs >= b[0] - NEAR) & (xs < b[2] + NEAR)
                            & (ys >= b[1] - NEAR)
                            & (ys < b[3] + NEAR)).mean()),
                     -(b[2] - b[0]) * (b[3] - b[1]), n)
                    for n, b in enumerate(eboxes))
                if frac < 0.5 or (1 - frac) * len(ys) > RES_DONATE:
                    continue
                # the fallback path has no mask to donate into: a comp poking
                # outside its rectangular crop would be erased from the page
                # with nowhere to reappear, so every raw pixel must sit
                # inside the crop rectangle itself (no NEAR slack)
                if elems[k][1] is None:
                    b = eboxes[k]
                    if not bool(((xs >= b[0]) & (xs < b[2])
                                 & (ys >= b[1]) & (ys < b[3])).all()):
                        continue
                sub = np.zeros(res.shape, bool)
                sub[sl] = part
                # a GREY comp DARKER than the reference is the page's own
                # achromatic pattern where the synth mispredicts it (page
                # 7's regional grid lines escaping the medians): it stays
                # baked, whole and consistent, like an achromatic divider.
                # Two exceptions are the element's own drawing after all:
                # ink DEEP below the reference (page 5's black tick marks -
                # page pattern is never that dark), and any grey inside a
                # CHART-named object, whose axes, ruler and hatching are
                # legitimately grey (page 8). A grey comp LIGHTER than the
                # reference is an element's paint-over remnant (page 9's
                # white panel interiors) and is erased without donation.
                px = arr[ys, xs].astype(int)
                dp = np.abs(px - np.median(arr.reshape(-1, 3), axis=0)) \
                    .sum(axis=1)
                core = px[dp >= np.quantile(dp, 0.75)] if len(px) > 4 else px
                cm = np.median(core, axis=0)
                refpx = ref[ys, xs] if grid is not None else ref
                grey = abs(cm[0] - cm[1]) + abs(cm[1] - cm[2]) <= GREY_SPREAD
                darker = bool(np.median((px - refpx).sum(axis=1)) <= 0)
                own = (float((np.median(np.atleast_2d(refpx), axis=0)
                              - cm).sum()) >= GREY_INK
                       or chart_owned(eboxes[k]))
                # an OWNED texture band (page 8's fine sub-grid ruler and
                # its ladder rail, round 14): most of its texture sits
                # under DIM_DIST, so donating only the res px ships a
                # peppered band while the page keeps the rest. Take the
                # whole strip: the crop carries the band as it was drawn,
                # and the patch hands the page clean paper. Never over a
                # text box (keep): the original px there hold the ERASED
                # line and would bake a duplicate into the crop.
                bh, bw = sl[0].stop - sl[0].start, sl[1].stop - sl[1].start
                # only take the whole strip when it will actually be DONATED
                # (mirror the donation gate below): a grey band lighter than
                # the ref is a paint-over remnant, erased not donated, so
                # widening its erase to the strip would delete it whole
                # (round 14 review). own is always true in this block.
                if own and (not grey or darker) and elems[k][1] is not None \
                        and max(bw, bh) >= 6 * min(bw, bh) \
                        and min(bw, bh) <= 40 and raw_px >= 0.1 * bw * bh:
                    band_m = np.zeros(res.shape, bool)
                    band_m[sl] = True
                    # stops clamped >=0 like page_grid/panel-sweep: an off-page
                    # keep box gives a negative stop and band_m[..:-n] would
                    # blank most of the strip (round 14 review)
                    for x0, y0, x1, y1 in keep:
                        band_m[max(int(y0) - PAD, 0):max(int(y1) + PAD, 0),
                               max(int(x0) - PAD, 0):max(int(x1) + PAD, 0)] \
                            = False
                    # never claim a sibling's own pixels into this band: they
                    # would double-paint on drag and get erased from the page
                    # here (round 14 review)
                    for kj, (_, mj) in enumerate(elems):
                        if kj != k and mj is not None:
                            band_m &= ~mj
                    # nor another residue component sitting in the strip's
                    # rectangle: it has its OWN owner and iteration, so
                    # swallowing it here double-donates the same pixels
                    band_m &= ~((lab != i) & (lab > 0))
                    sub |= band_m
                    part = sub[sl]
                    ys, xs = np.nonzero(sub)
                # a texture may only veto its OWN kind: page 5's mint mosaic
                # square sat 34 away from a grey texture colour and the flat
                # distance test alone baked it into the page for two rounds
                if not own and any(
                        np.abs(cm - tc).sum() <= 3 * TINT_MATCH
                        and (abs(tc[0] - tc[1]) + abs(tc[1] - tc[2])
                             <= GREY_SPREAD) == grey
                        for tc in texcols):
                    continue           # the page's own pattern: stays baked
                if grey and darker and not own:
                    continue
                patch(arr, bx, pad=MASK_PAD, others=keep + eboxes, mask=sub,
                      grid=grid)
                # donate only the RAW residue pixels: the dilation ring is
                # near-background by construction, and shipped in the crop it
                # rings every donated mark with page fringe and drags the
                # translucent fit toward the background
                if (not grey or (darker and own)) \
                        and elems[k][1] is not None:
                    elems[k][1][ys, xs] = True
            # a donated band can poke past the model's box: re-tighten
            for e in elems:
                if e[1] is not None:
                    e[0] = _tight(e[1]) or e[0]
            eboxes = [b for b, _ in elems]
        # a stamp's panel box is watermark furniture with its stamp: its ink
        # was already patched out above, so dropping it here removes the
        # panel from the deck instead of shipping it as a movable element.
        # Size-capped: a page-wide diagram merely OVERLAPPING a stamp box is
        # not the stamp's panel, and dropping it would delete the artwork.
        panels = [e[0] for e in elems
                  if any(min(e[0][2], s[2]) - max(e[0][0], s[0]) > 0
                         and min(e[0][3], s[3]) - max(e[0][1], s[1]) > 0
                         and (min(e[0][2], s[2]) - max(e[0][0], s[0]))
                         * (min(e[0][3], s[3]) - max(e[0][1], s[1]))
                         >= 0.5 * (s[2] - s[0]) * (s[3] - s[1])
                         and (e[0][2] - e[0][0]) * (e[0][3] - e[0][1])
                         <= STAMP_PANEL_MAX * (s[2] - s[0]) * (s[3] - s[1])
                         for s, _ in stamps)]
        elems = [e for e in elems if e[0] not in panels]
        # the panel's mask can run a few px short of its frame (page 7 left a
        # dark corner speck beside the erased panel): sweep what the drop
        # left behind inside the panel's grown box
        if panels:
            from scipy import ndimage as ndi1
            ref1 = grid if grid is not None \
                else np.median(arr.reshape(-1, 3), axis=0)
            cur_eboxes = [e[0] for e in elems]
            for px0, py0, px1, py1 in panels:
                gx0, gy0 = max(px0 - PANEL_SWEEP, 0), max(py0 - PANEL_SWEEP, 0)
                gx1 = min(px1 + PANEL_SWEEP, arr.shape[1])
                gy1 = min(py1 + PANEL_SWEEP, arr.shape[0])
                sub = np.abs(arr[gy0:gy1, gx0:gx1].astype(int)
                             - (ref1[gy0:gy1, gx0:gx1]
                                if grid is not None else ref1)).sum(axis=2) \
                    > DIM_DIST
                # pixels a patch just painted are the tool's own best fill,
                # never panel debris (same guard as the residue pass)
                sub &= ~painted[gy0:gy1, gx0:gx1]
                # wmboxes stay OUT of the shield: the pill's rectangle is a
                # MASK patch, so unmasked pixels inside it are raw page, and
                # page 7's panel speck hid exactly there
                for x0, y0, x1, y1 in [b for b in keep if b not in wmboxes]:
                    # stops clamped to >=0: a keep box above/left of the sweep
                    # window gives a negative stop, and sub[..:-n] would blank
                    # most of the window instead of nothing
                    sub[max(y0 - PAD - gy0, 0):max(y1 + PAD - gy0, 0),
                        max(x0 - PAD - gx0, 0):max(x1 + PAD - gx0, 0)] = False
                lab1, _ = ndi1.label(sub)
                for i1, sl1 in enumerate(ndi1.find_objects(lab1), 1):
                    part1 = lab1[sl1] == i1
                    if not 3 <= part1.sum() <= PANEL_SPECK:
                        continue
                    m1 = np.zeros(arr.shape[:2], bool)
                    m1[gy0 + sl1[0].start:gy0 + sl1[0].stop,
                       gx0 + sl1[1].start:gx0 + sl1[1].stop] = part1
                    patch(arr, _tight(m1), pad=MASK_PAD,
                          others=keep + cur_eboxes, mask=m1, grid=grid)
        # an artwork element boxed almost wholly inside another is one drawing
        # split in two (page 4's REJECT scribble sits over its own funnel):
        # shipped separately, dragging the diagram leaves the scribble
        # floating. Thin rules are exempt (page 10's underline lies inside
        # the disc's box but is its own furniture). Runs AFTER the stamp-panel
        # drop, so nothing can merge INTO a panel and be deleted with it.
        merged = True
        while merged:
            merged = False
            for ai in range(len(elems)):
                for bi in range(len(elems)):
                    A, B = elems[ai][0], elems[bi][0]
                    bw, bh = B[2] - B[0], B[3] - B[1]
                    if (ai == bi or elems[ai][1] is None
                            or elems[bi][1] is None
                            or min(bw, bh) <= RULE_THIN
                            or (A[2] - A[0]) * (A[3] - A[1]) <= bw * bh):
                        continue
                    ix = max(0, min(A[2], B[2]) - max(A[0], B[0]))
                    iy = max(0, min(A[3], B[3]) - max(A[1], B[1]))
                    if ix * iy >= CONTAIN * bw * bh:
                        elems[ai][1] |= elems[bi][1]
                        elems[ai][0] = _tight(elems[ai][1]) or A
                        del elems[bi]
                        merged = True
                        break
                if merged:
                    break
        eboxes = [b for b, _ in elems]
        crops = []
        for ei, ((x0, y0, x1, y1), m) in enumerate(elems):
            tile = arr0[y0:y1, x0:x1].copy()
            if m is None:
                crops.append(Image.fromarray(tile))
                continue
            mm = m[y0:y1, x0:x1]
            a8 = np.where(mm, 255, 0).astype(np.uint8)
            # what solidify may never fill over or restore to the page:
            # sibling elements' claims and every text line's box -
            # re-typed, baked in-artwork, AND the sub-8pt lines that stay
            # baked (a pale tick label's counters read as furniture cells
            # otherwise; reviewer catch, round 13)
            avoid = np.zeros(mm.shape, bool)
            for ej, (_, m2) in enumerate(elems):
                if ej != ei and m2 is not None:
                    avoid |= m2[y0:y1, x0:x1]
            for bx0, by0, bx1, by1 in keep + bakedb \
                    + [i["erase"] for i in small]:
                ay0 = max(int(by0) - y0, 0)
                ay1 = min(int(by1) + 1 - y0, y1 - y0)
                ax0 = max(int(bx0) - x0, 0)
                ax1 = min(int(bx1) + 1 - x0, x1 - x0)
                if ay0 < ay1 and ax0 < ax1:
                    avoid[ay0:ay1, ax0:ax1] = True
            translucent(tile, a8, mm, arr[y0:y1, x0:x1])
            solidify(tile, a8, mm, arr[y0:y1, x0:x1], avoid)
            crops.append(Image.fromarray(np.dstack([tile, a8]), "RGBA"))

        if dump:                       # diagnostic layers, nothing else uses them
            os.makedirs(dump, exist_ok=True)
            Image.fromarray(arr).save(f"{dump}/p{page.number + 1}_bg.png")
            for n, c in enumerate(crops):
                c.save(f"{dump}/p{page.number + 1}_e{n}.png")

        slide = prs.slides.add_slide(blank)
        add_picture(slide, Image.fromarray(arr), 0, 0,
                    page.rect.width, page.rect.height)
        for ((x0, y0, x1, y1), _), crop in zip(elems, crops):
            add_picture(slide, crop, x0 * sx, y0 * sy,
                        (x1 - x0) * sx, (y1 - y0) * sy)

        blocks, bullets = merge_panel_blocks(blocks, bullets, eboxes)
        pending.append((slide, blocks, bullets, sx, sy,
                        page.rect.width, page.rect.height, rots))
        report.append({"page": page.number + 1,
                       "text_lines": len(items) + len(rots),
                       "blocks": len(blocks) + len(rots),
                       "elements": len(elems), "baked": baked})

    for slide, blocks, bullets, sx, sy, page_w, page_h, rots in pending:
        for blk, bu in zip(blocks, bullets):
            emit_block(slide, blk, sx, sy, page_w, bu)
        for r in rots:
            # emitted in the turned page's frame (axes swapped), then swung
            # into place: pptx rotates a shape about its centre, so the box
            # keeps its horizontal size and only the centre needs mapping.
            tb = emit_block(slide, [r["item"]], sy, sx, page_h, centred=False)
            cx = (tb.left + tb.width / 2) / 12700.0        # EMU -> pt
            cy = (tb.top + tb.height / 2) / 12700.0
            if r["transpose"] == Image.ROTATE_270:
                ux, uy, tb.rotation = cy, page_h - cx, 270.0
            else:
                ux, uy, tb.rotation = page_w - cy, cx, 90.0
            tb.left = int(round(ux * 12700 - tb.width / 2))
            tb.top = int(round(uy * 12700 - tb.height / 2))

    prs.save(out_path)
    return report


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pdf")
    ap.add_argument("out")
    ap.add_argument("--font", default="Segoe UI",
                    help="font family, or a comma-separated list the deck mixes "
                         "(e.g. \"Roboto Mono,Roboto\"): each line is matched to "
                         "whichever one it was actually set in")
    ap.add_argument("--font-map", help="JSON file: {line-text regex: font name}")
    ap.add_argument("--classify-cache", default=None,
                    help="JSON file caching which lines are drawn inside the "
                         "artwork, so a rebuild does not re-pay for the call "
                         "(default: <out>.labels.json)")
    ap.add_argument("--dump", default=None,
                    help="directory to save per-page background and element "
                         "layers into, for checking the cut by eye")
    ap.add_argument("--strip-stamps", action="store_true",
                    help="erase the generator's draft-stamp footers instead "
                         "of keeping them as editable deck content")
    a = ap.parse_args()
    global STRIP_STAMPS
    STRIP_STAMPS = a.strip_stamps
    fm = json.load(open(a.font_map, encoding="utf-8")) if a.font_map else {}
    cache = a.classify_cache or os.path.splitext(a.out)[0] + ".labels.json"
    report = rebuild(a.pdf, a.out, a.font, fm, cache, a.dump)
    for r in report:
        print(f"page {r['page']}: {r['text_lines']} lines in {r['blocks']} "
              f"text boxes, {r['elements']} image elements"
              + (f", {r['baked']} left baked in artwork" if r["baked"] else ""))
    assert any(r["text_lines"] for r in report), "no text extracted at all"
    print("saved:", a.out)
    if MODEL_MISSES:
        print("\nWARNING: vision model unavailable on "
              f"{len(MODEL_MISSES)} pass(es); the deck is OCR-ONLY (garbled "
              "~ % | glyphs, ghost lines, wrong bake). Fix ANTHROPIC_API_KEY "
              "(.env) and rerun.\n  " + "\n  ".join(MODEL_MISSES[:6]),
              file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
