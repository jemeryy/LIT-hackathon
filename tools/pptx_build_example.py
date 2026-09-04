# Builds Udemy Lecture 1.1 as a fully editable PPTX: every text box, shape,
# and connector is an individual native PowerPoint element.
# Slide content comes from the [SLIDE 1.1-NN] visual: markers in
# projects/udemy-course/scripts/section-01.md.
# ponytail: doodle icons are simple outlined shapes as placeholders; swap for
# real doodle art in the animation pass.
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN

BG = RGBColor(0x14, 0x16, 0x1A)
WHITE = RGBColor(0xF5, 0xF5, 0xF5)
MINT = RGBColor(0x6E, 0xE7, 0xB7)
CORAL = RGBColor(0xF8, 0x71, 0x71)
YELLOW = RGBColor(0xFD, 0xE0, 0x47)
GREY = RGBColor(0x8A, 0x91, 0x99)
FONT = "Segoe Print"  # hand-drawn feel, ships with Windows

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s


def text(s, x, y, w, h, txt, size=24, color=WHITE, align=PP_ALIGN.CENTER,
         underline=False, strike=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = txt
    f = r.font
    f.name, f.size, f.color.rgb, f.underline = FONT, Pt(size), color, underline
    if strike:
        f._rPr.set("strike", "sngStrike")
    return tb


def doodle(s, shape_type, x, y, w, h, label=None, color=WHITE, label_size=16):
    sh = s.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.background()
    sh.line.color.rgb = color
    sh.line.width = Pt(2.25)
    sh.shadow.inherit = False
    if label is not None:
        tf = sh.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(label_size), color
    return sh


def arrow(s, x1, y1, x2, y2, color=WHITE):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                               Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(2.25)
    return c


# SLIDE 1.1-01: chat bubble centre-left, "language model"
s = slide()
doodle(s, MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT, 2.2, 2.6, 3.2, 1.8)
text(s, 2.2, 4.7, 3.2, 0.6, "language model", 24)

# SLIDE 1.1-02: two columns, great at / weak at
s = slide()
text(s, 1.2, 0.7, 4.5, 0.8, "great at", 32, MINT)
text(s, 7.6, 0.7, 4.5, 0.8, "weak at", 32, CORAL)
left = [("drafting", MSO_SHAPE.ISOSCELES_TRIANGLE),      # pencil
        ("reshaping", MSO_SHAPE.LEFT_RIGHT_ARROW),        # two arrows
        ("spotting patterns", MSO_SHAPE.OVAL)]            # magnifying glass
right = [("exact maths", MSO_SHAPE.RECTANGLE),            # calculator
         ("facts with no source", MSO_SHAPE.ROUNDED_RECTANGLE),  # fact tag "?"
         ("same answer every time", MSO_SHAPE.CUBE)]      # dice
for i, (label, shp) in enumerate(left):
    y = 1.9 + i * 1.6
    doodle(s, shp, 1.6, y, 0.9, 0.9, color=MINT)
    text(s, 2.8, y + 0.15, 3.3, 0.6, label, 20, WHITE, PP_ALIGN.LEFT)
for i, (label, shp) in enumerate(right):
    y = 1.9 + i * 1.6
    doodle(s, shp, 8.0, y, 0.9, 0.9, color=CORAL)
    text(s, 9.2, y + 0.15, 3.5, 0.6, label, 20, WHITE, PP_ALIGN.LEFT)

# SLIDE 1.1-03: chatbot answers -> agent does the work
s = slide()
doodle(s, MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT, 1.5, 2.7, 2.8, 1.6)
text(s, 1.2, 4.6, 3.4, 0.6, "chatbot: answers", 22)
arrow(s, 5.0, 3.5, 7.4, 3.5)
doodle(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.2, 2.5, 2.2, 1.6)   # robot head
doodle(s, MSO_SHAPE.OVAL, 8.6, 2.9, 0.35, 0.35)              # eye
doodle(s, MSO_SHAPE.OVAL, 9.6, 2.9, 0.35, 0.35)              # eye
doodle(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.7, 3.0, 0.4, 0.9)   # hand
doodle(s, MSO_SHAPE.ROUNDED_RECTANGLE, 10.5, 3.0, 0.4, 0.9)  # hand
text(s, 7.7, 4.6, 3.4, 0.6, "agent: does the work", 22)

# SLIDE 1.1-04: laptop centre, five icons around
s = slide()
doodle(s, MSO_SHAPE.ROUNDED_RECTANGLE, 5.4, 3.0, 2.5, 1.6, "laptop", WHITE, 18)
around = [("folder", MSO_SHAPE.FOLDED_CORNER, 2.0, 1.2),
          ("web page", MSO_SHAPE.RECTANGLE, 9.8, 1.2),
          ("spreadsheet", MSO_SHAPE.RECTANGLE, 1.5, 4.6),
          ("slide deck", MSO_SHAPE.RECTANGLE, 10.3, 4.6),
          ("envelope", MSO_SHAPE.HEXAGON, 5.9, 5.9)]
for label, shp, x, y in around:
    doodle(s, shp, x, y, 1.5, 1.0, label, WHITE, 14)
    cx, cy = x + 0.75, y + 0.5
    arrow(s, 6.65 + (0.9 if cx > 6.65 else -0.9), 3.8 + (0.55 if cy > 3.8 else -0.55),
          cx + (-0.4 if cx > 6.65 else 0.4), cy + (-0.25 if cy > 3.8 else 0.25))

# SLIDE 1.1-05: execution struck through; judgement + strategy underlined
s = slide()
text(s, 3.7, 2.0, 6.0, 1.0, "execution", 44, GREY, strike=True)
text(s, 3.7, 3.4, 6.0, 1.0, "judgement + strategy", 44, YELLOW, underline=True)
text(s, 3.7, 5.6, 6.0, 0.6, "your edge = what you know", 20, WHITE)

OUT = "projects/udemy-course/slides-trial/lecture-1.1-editable.pptx"
prs.save(OUT)

# self-check: every slide has editable text, no slide is a single baked image
from pptx import Presentation as P
chk = P(OUT)
assert len(chk.slides) == 5
for sl in chk.slides:
    runs = [r.text for sh in sl.shapes if sh.has_text_frame
            for p in sh.text_frame.paragraphs for r in p.runs]
    assert runs, "slide with no editable text"
print("OK:", OUT, "| shapes per slide:", [len(sl.shapes) for sl in chk.slides])
