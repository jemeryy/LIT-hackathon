"""Build the pitch deck. Run: python deck/make_deck.py  ->  deck/case-helper-deck.pptx"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import pathlib

ROOT = pathlib.Path(__file__).resolve().parents[1]
SCREENS = ROOT / "notes" / "screens"
INK = RGBColor(0x11, 0x11, 0x11)
GREY = RGBColor(0x55, 0x55, 0x55)
GREEN = RGBColor(0x15, 0x80, 0x3d)
RED = RGBColor(0xb9, 0x1c, 0x1c)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, lines, size=20, bold=False, color=INK, align=PP_ALIGN.LEFT, gap=6):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for line in lines if isinstance(lines, list) else [lines]:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        b = bold
        if isinstance(line, tuple):   # a coloured line is a heading: bold it
            line, c = line
            b = True
        else:
            c = color
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), b, FONT
        r.font.color.rgb = c
        p.alignment = align
        p.space_after = Pt(gap)
    return tb


def slide(title, bullets=None, image=None, note=None, size=22, img_w=7.4):
    s = prs.slides.add_slide(BLANK)
    text(s, Inches(0.6), Inches(0.35), W - Inches(1.2), Inches(1), title, size=34, bold=True)
    line = s.shapes.add_shape(1, Inches(0.6), Inches(1.2), W - Inches(1.2), Pt(2))
    line.fill.solid(); line.fill.fore_color.rgb = INK; line.line.fill.background()
    body_w = W - Inches(1.2) - (Inches(img_w + 0.3) if image else 0)
    if bullets:
        text(s, Inches(0.6), Inches(1.5), body_w, H - Inches(2.3), bullets, size=size, gap=10)
    if image:
        p = SCREENS / image
        pic = s.shapes.add_picture(str(p), W - Inches(0.6 + img_w), Inches(1.5), width=Inches(img_w))
        if pic.height > H - Inches(2.2):   # tall page grab: cap the height, keep the ratio
            ratio = (H - Inches(2.2)) / pic.height
            pic.height, pic.width = int(pic.height * ratio), int(pic.width * ratio)
            pic.left = W - Inches(0.6) - pic.width
        pic.line.color.rgb = INK; pic.line.width = Pt(1.5)
    if note:
        text(s, Inches(0.6), H - Inches(0.75), W - Inches(1.2), Inches(0.5), note, size=14, color=GREY)
    return s


# 1 title
s = prs.slides.add_slide(BLANK)
text(s, Inches(0.8), Inches(2.0), W - Inches(1.6), Inches(1.2), "Case Helper", size=60, bold=True)
text(s, Inches(0.8), Inches(3.2), W - Inches(1.6), Inches(1), "Build your Small Claims case from your evidence.", size=30, color=GREY)
text(s, Inches(0.8), Inches(5.3), W - Inches(1.6), Inches(1.5), [
    "SMU LIT Legal-Tech Hackathon 2026. Problem statement 4, MinLaw: people at the Small Claims Tribunals with no lawyer.",
    "Team: [TEAM NAME]",
    "Live: case-helper-production.up.railway.app    Code: github.com/jemeryy/LIT-hackathon"], size=18, color=GREY)

# 2 problem
slide("The problem", [
    "At the Small Claims Tribunals you cannot bring a lawyer. You prepare the case yourself.",
    "So people ask ChatGPT. It agrees with them, makes up cases, and does not know the court's steps.",
    "They turn up with a story and no proof, or the wrong proof, and lose on the facts.",
    "The Courts' guide on generative AI says the person is responsible for what they file. Nothing helps them meet that duty.",
    ("What they need is not legal advice. It is a way to sort their own evidence and see what the other side will say.", GREEN),
], size=24)

# 3 what we built
slide("What we built: five steps, one page", [
    "1. Tell us what happened. A chat in your own words. The model pulls out the facts and asks for what is missing.",
    "2. Can the tribunal hear it? Four rule checks, each quoting the section of the Act.",
    "3. Your evidence. Add files, see what else to gather, answer what the other side may have, then rank both sides.",
    "4. Your timeline. Past events from your files, then the court's steps ahead and the time bar date.",
    "5. What to do next. In the court's own order, with a claim pack ready to paste into the filing portal.",
], image="deck-1-chat.png", size=17, note="Works for a real case today: tenancy deposit, goods, services, damage to property. Not only the demo.")

# 4 gate
slide("Step 2: rules, not a model", [
    "Type of claim, amount, time bar, other side in Singapore.",
    "Each check quotes the section from Singapore Statutes Online, with the date we fetched it.",
    "Same answer every time. Fail a check and it stops and points to the right place.",
    "Work claims, loans and neighbour noise are turned away here, before any file is read.",
], image="deck-2-gate.png", size=20)

# 5 evidence
slide("Step 3: your evidence, ranked from the files", [
    "Every fact carries an exact quote. We match the quote back into the file. No match, the fact is dropped.",
    "Press a row and the file opens at that page, with the line highlighted.",
    "A file whose names, dates or amounts do not fit your account is marked weak and kept out of the story.",
    "What else to gather: five headings for your kind of claim, what you have and what is still missing.",
], image="deck-3-evidence.png", size=19, img_w=7.4)

# 6 other side
slide("The other side's evidence: the confirmation bias fix", [
    "Seven yes / no / not sure questions about what the other side could bring.",
    "Their list, ranked for them, with why it is strong and how to answer it from your own files.",
    "We never argue their case. We list what they could show so the person is not surprised at the first hearing.",
    "A file you have of theirs goes in their column, never in yours.",
], image="deck-3b-ranked.png", size=20, img_w=7.4)

# 7 timeline + next
s = slide("Steps 4 and 5: the road ahead", None, size=18)
s.shapes.add_picture(str(SCREENS / "deck-4-timeline.png"), Inches(0.6), Inches(1.5), width=Inches(6.0))
s.shapes.add_picture(str(SCREENS / "deck-5-next.png"), Inches(6.75), Inches(1.5), width=Inches(6.0))
text(s, Inches(0.6), Inches(5.4), W - Inches(1.2), Inches(1.5), [
    "Timeline: events from the files, today, then the court's steps and the 2 year time bar. Maps onto the court's Submission for Hearing form.",
    "Next steps in the court's order: written request, pre-filing assessment, the six part claim form with a 500 character summary, $10 fee, serve within 7 working days.",
    "Downloads: an evidence sheet (xlsx), a written request, and a claim pack with the form text and every file as a PDF under 5 MB."], size=16)

# 8 model vs rules
s = slide("Where the model stops and the rules start", None)
text(s, Inches(0.6), Inches(1.5), Inches(4.0), Inches(5), [
    ("The model does", GREEN),
    "Reads what you type and fills the form",
    "Reads each file and returns facts with exact quotes",
    "Writes the story, the 500 character summary and the request letter",
], size=19)
text(s, Inches(4.8), Inches(1.5), Inches(4.0), Inches(5), [
    ("Rules do", INK),
    "The four gate checks",
    "Strength and rank of every fact",
    "What else to gather, the other side's list",
    "Timeline, fee, next steps",
], size=19)
text(s, Inches(9.0), Inches(1.5), Inches(3.8), Inches(5), [
    ("Nobody does", RED),
    "No case names, ever",
    "No guess at who wins",
    "No legal advice",
    "No invented evidence",
    "No file leaves your case",
], size=19)
text(s, Inches(0.6), Inches(5.6), W - Inches(1.2), Inches(1.2), [
    "Every quote is matched into the file or dropped. Legal conclusions and advice words from the model are rejected by the server. File text, names and chat are treated as data, never as instructions.",
    "The person is told on every page: this tool does not give legal advice, you are responsible for what you file, check every rule on Singapore Statutes Online."], size=15, color=GREY)

# 9 real vs stubbed
slide("What is real and what is stubbed", [
    ("Real", GREEN),
    "Chat intake, file reading (PDF, screenshots, photos, one video frame), quotes matched and highlighted, the gate, both rankings, gather list, timeline, fee, next steps, evidence sheet, request letter, claim pack. Live now on Railway.",
    ("Stubbed", RED),
    "Filing on the court's portal (CJTS). There is no public API, so the claim pack is what you paste and upload yourself.",
    "Statute text is only the sections the gate quotes, plus the court's guide. Not all of Singapore law.",
    ("Data", GREY),
    "The demo files are made up by us. No real documents were used at any point.",
], size=19)

# 10 AI disclosure
slide("AI tools used (disclosure)", [
    ("To build it", INK),
    "Claude Code (Anthropic), with Claude Opus and Fable models, for planning, code, browser testing and this deck.",
    "OpenAI Codex, for a second opinion on each plan and a review of each change.",
    "Playwright, driven by Claude Code, for end to end tests in a real browser.",
    ("Inside the product", INK),
    "Claude Sonnet 5 through OpenRouter (anthropic/claude-sonnet-5), with the hackathon credits.",
    ("Design decisions", INK),
    "Every design call was the team's, recorded in the plan file and the commit history in the repo. We can explain every line.",
], size=19)

# 11 references
slide("References", [
    ("Law and court sources", INK),
    "Small Claims Tribunals Act 1984, s 2, s 5, s 23 and the Schedule. Singapore Statutes Online, sso.agc.gov.sg/Act/SCTA1984, read 5 Sep 2026.",
    "State Courts, A Guide to Small Claims (PDF), judiciary.gov.sg. Fees, service, e-Negotiation, appeal.",
    "Judiciary pages: File a small claim; How to file and serve a small claim. Community Justice and Tribunals System, cjts.judiciary.gov.sg.",
    "Registrar's Circular No. 1 of 2024, Guide on the Use of Generative AI Tools by Court Users, in force 1 Oct 2024.",
    ("Libraries and tools", INK),
    "FastAPI, Uvicorn, python-multipart, httpx, Anthropic Python SDK, pdfplumber, PyMuPDF, Pillow, pillow-heif, pytesseract with Tesseract OCR, openpyxl, reportlab, ffmpeg, python-pptx.",
    ("Data", INK),
    "Synthetic sample pack made by the team (builds/case-builder/sample). No organiser or real data.",
], size=16)

# 12 next
slide("What we would build next", [
    "Keep cases and files across restarts, so a person can come back a week later.",
    "Own content sets for services and damage to property, then the other claim types in the Schedule.",
    "Prepare the person for e-Negotiation on CJTS: the five rounds of offers.",
    "Test it with Community Justice Centre volunteers on real, consented cases.",
    ("Ask us anything. We built every part and can show the code.", GREEN),
], size=22)

out = ROOT / "deck" / "case-helper-deck.pptx"
prs.save(out)
print("saved", out, len(prs.slides), "slides")
